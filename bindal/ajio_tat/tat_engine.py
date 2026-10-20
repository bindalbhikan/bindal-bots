"""
AJIO Order Dispatch TAT engine  (10:10 rule)

Rule: orders received from yesterday 10:00 AM to today 10:00 AM must be
processed AND dispatched (picked up by courier) the same day (today).
  -> order time <  10:00  => due date = order date
  -> order time >= 10:00  => due date = order date + 1

Ported from the original main.py script. The rules are unchanged; the differences are:
  * inputs arrive as DataFrames (scraped from OMSGuru, tracked live) instead of
    the latest files in ./input
  * the master carries an extra "Order ID" column (AJIO channel order id) so an
    order that is not invoiced yet can be stored and later replaced by its
    invoiced row
  * write_excel() is kept for the optional local .xlsx copy of the master
"""
import re
from datetime import time, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CUTOFF = time(10, 0)          # 10:10 rule cut-off
SKIP_SUNDAY = False           # True => if due date falls on Sunday, push to Monday

OMS_DATE_FMT = "%B %d, %Y %I:%M:%S %p"
PENDING_PREFIX = "PENDING-"   # key used for orders that have no invoice number yet

# Courier event statuses that mean the parcel was physically picked up.
PICKUP_STATUSES = {
    "shipment picked up", "pickdone", "picked up", "pickup done", "picked",
}
# Events that happen BEFORE pickup (ignored when falling back to first movement).
PRE_PICKUP = re.compile(
    r"data received|manifest|pickup employee|out for pick|pickup has been|"
    r"pickup re-scheduled|auto pickup|p/u|pending|registered|order cancelled|pickup scheduled",
    re.I,
)
# Courier scans that mean a pickup was booked / an agent is on the way (pickup attempt).
PICKUP_ATTEMPT = r"pickup employee|out for pick|pickup has been|auto pickup|pickup scheduled"

# An order processed (RTS) today at or after this time is picked up by the courier tomorrow:
# its Pickup Status is "Next Day Pickup", not "Pending Pickup". Next run it becomes
# "Pending Pickup" again if still not picked.
NEXT_DAY_PICKUP = "Next Day Pickup"
NEXT_DAY_AFTER = time(12, 0)
OPEN_PICKUP = ("Pending Pickup", NEXT_DAY_PICKUP)
CANCEL_ITEM_STATUSES = re.compile(r"cancel", re.I)
# OMSGuru item statuses that mean the courier took the parcel. Old orders (a rebuild from months
# back) have no courier history left and OMSGuru logs no pickup event, so these say "picked".
PICKED_ITEM_STATUSES = re.compile(r"in transit|delivered|return|rto|rtv", re.I)
# Courier summary statuses meaning the parcel has not left the warehouse yet.
# PICKUP_FAILED: the courier came but did not take it. RETURN_PENDING: a scan with
# return wording before pickup. Neither means the parcel left the warehouse.
NOT_PICKED_COURIER_STATUSES = {"UNKNOWN", "NOT_FOUND", "PENDING", "OUT_FOR_PICKUP", "PICKUP_FAILED",
                               "RETURN_PENDING", "nan", "None", ""}

REPORT_COLS = ["Invoice No", "Order ID", "AWB", "Courier", "Order Time", "Dispatch Deadline",
               "Processed Time", "RTS Time", "Pickup Attempt Time", "Pickup Time",
               "Process Status", "Pickup Status", "TAT Result", "OMS Item Status",
               "Courier Status", "Label Printed", "ERP Label Time", "ERP Label By", "Last Updated"]
# Raw fields kept in the master (_data sheet) so statuses can be recomputed every run.
DATA_COLS = ["Invoice No", "Order ID", "AWB", "Courier", "Order Time", "Dispatch Deadline",
             "Processed Time", "RTS Time", "RTS Exact", "Pickup Attempt Time", "Pickup Time",
             "OMS Item Status", "Courier Status", "ERP Label Time", "ERP Label By", "Last Updated"]
TIME_COLS = ["Order Time", "Dispatch Deadline", "Processed Time", "RTS Time",
             "Pickup Attempt Time", "Pickup Time", "ERP Label Time", "Last Updated"]
STATUS_COLS = ["Process Status", "Pickup Status", "TAT Result"]


def parse_oms(s):
    return pd.to_datetime(s, format=OMS_DATE_FMT, errors="coerce")


def due_date_for(order_dt):
    if pd.isna(order_dt):
        return pd.NaT
    d = order_dt.date() if order_dt.time() < CUTOFF else order_dt.date() + timedelta(days=1)
    if SKIP_SUNDAY and d.weekday() == 6:
        d += timedelta(days=1)
    return pd.Timestamp(d)


def awb_of(s):
    m = re.match(r"\s*([A-Za-z0-9]+)", str(s))
    return m.group(1) if m else ""


def prepare_tracking(summary_rows, event_rows):
    """Same shape the original main.py got from the tracking_*.xlsx Summary + Events sheets."""
    summ = pd.DataFrame(summary_rows, columns=["AWB", "PARTNER", "STATUS", "STATUS_RAW",
                                               "CURRENT_LOCATION", "LAST_UPDATE", "DELIVERED_ON"])
    ev = pd.DataFrame(event_rows, columns=["AWB", "SEQ", "TIMESTAMP", "STATUS", "LOCATION"])
    summ["AWB"] = summ["AWB"].astype(str)
    ev["AWB"] = ev["AWB"].astype(str)
    ev["ts"] = pd.to_datetime(ev["TIMESTAMP"], errors="coerce")
    ev = ev.dropna(subset=["ts"]).sort_values(["AWB", "ts", "SEQ"])
    return summ, ev


def pickup_times(ev):
    """Returns DataFrame AWB -> pickup_ts, pickup_event, pickup_source, first_attempt_ts."""
    rows = []
    for awb, g in ev.groupby("AWB", sort=False):
        st = g["STATUS"].astype(str).str.strip()
        exact = g[st.str.lower().isin(PICKUP_STATUSES)]
        attempt = g[st.str.contains(PICKUP_ATTEMPT, case=False, regex=True)]
        first_attempt = attempt["ts"].iloc[0] if len(attempt) else pd.NaT
        if len(exact):
            r = exact.iloc[0]
            before = attempt[attempt["ts"] <= r["ts"]]  # attempt scans logged after pickup are not attempts
            first_attempt = before["ts"].iloc[0] if len(before) else pd.NaT
            rows.append((awb, r["ts"], r["STATUS"], "Pickup scan", first_attempt))
            continue
        # History is only trustworthy when it starts from a pre-pickup event.
        # Shadowfax / Delhivery files carry only the latest event -> no pickup data.
        if not PRE_PICKUP.search(st.iloc[0]):
            rows.append((awb, pd.NaT, "", "No pickup data (courier history missing)", first_attempt))
            continue
        moved = g[~st.str.contains(PRE_PICKUP)]
        if len(moved):
            r = moved.iloc[0]
            rows.append((awb, r["ts"], r["STATUS"], "First movement scan", first_attempt))
        else:
            rows.append((awb, pd.NaT, "", "No pickup scan", first_attempt))
    return pd.DataFrame(rows, columns=["AWB", "Pickup Time", "Pickup Event",
                                       "Pickup Source", "Pickup Attempt Time"])


def build_report(o, hist, summ, ev):
    o = o.copy()
    o["AWB"] = o["AWB / Shipment Tracker"].map(awb_of)

    # ---- one row per order (AWB) ----
    agg = o.groupby(["Input Identifier", "AWB"], sort=False).agg(
        internal_id=("Internal Order ID", "first"),
        channel_order=("Channel Order ID", "first"),
        skus=("SKU Code", lambda s: ", ".join(dict.fromkeys(s.astype(str)))),
        qty=("Qty", "sum"),
        item_status=("Item Status", lambda s: ", ".join(dict.fromkeys(s.astype(str)))),
        order_date=("Order Date", "first"),
        invoice_date=("Invoice Date", "first"),
        sla=("SLA Date", "first"),
        dispatch_log=("Dispatch Log", "first"),
        courier=("Shipping Company", "first"),
        payment=("Payment Mode", "first"),
    ).reset_index()

    agg["Order Time"] = parse_oms(agg["order_date"])
    agg["Invoice Time"] = parse_oms(agg["invoice_date"])
    agg["SLA Time"] = parse_oms(agg["sla"])

    # Status history (only available for some orders)
    def first_ts(status):
        return hist[hist["Status"] == status].groupby("Internal Order ID")["ts"].min()

    h_proc, h_rts = first_ts("Processing"), first_ts("Ready to ship")
    agg["hist_proc"] = agg["internal_id"].map(h_proc)
    agg["hist_rts"] = agg["internal_id"].map(h_rts)

    # Processed time: history 'Processing' else Invoice generation time
    agg["Processed Time"] = agg["hist_proc"].fillna(agg["Invoice Time"])

    # Ready to ship: history 'Ready to ship' else Dispatch Log date (date only)
    dlog_date = pd.to_datetime(agg["dispatch_log"].astype(str).str[:10], errors="coerce")
    agg["RTS Time"] = agg["hist_rts"].fillna(dlog_date)
    agg["RTS Source"] = agg["hist_rts"].notna().map(
        {True: "Status History", False: "Dispatch Log (date only)"})
    # Dispatch Log blank -> invoice time (label is printed on invoicing)
    no_rts = agg["RTS Time"].isna() & agg["Invoice Time"].notna()
    agg.loc[no_rts, "RTS Time"] = agg.loc[no_rts, "Invoice Time"]
    agg.loc[no_rts, "RTS Source"] = "Invoice Date"
    agg.loc[agg["RTS Time"].isna(), "RTS Source"] = "Not available"

    # Courier pickup
    pk = pickup_times(ev)
    agg = agg.merge(pk, on="AWB", how="left")
    agg = agg.merge(summ.rename(columns={"STATUS": "Courier Status"})[["AWB", "Courier Status"]]
                    .drop_duplicates("AWB"), on="AWB", how="left")

    agg["Dispatch Deadline"] = agg["Order Time"].map(due_date_for) + pd.Timedelta(hours=23, minutes=59, seconds=59)
    agg["RTS Exact"] = (agg["RTS Source"] != "Dispatch Log (date only)").astype(object)
    agg.loc[agg["RTS Time"].isna(), "RTS Exact"] = None
    agg["Last Updated"] = pd.Timestamp.now().normalize()

    out = agg.rename(columns={"Input Identifier": "Invoice No", "courier": "Courier",
                              "item_status": "OMS Item Status", "channel_order": "Order ID"})
    as_of = ev["ts"].max() if len(ev) else pd.Timestamp.now()
    return out.reindex(columns=DATA_COLS).reset_index(drop=True), as_of   # ERP columns come later


def normalize_master(m):
    """Coerce a master loaded from Sheets / xlsx into DATA_COLS with proper dtypes."""
    m = m.reindex(columns=DATA_COLS).copy()
    for c in TIME_COLS:
        m[c] = pd.to_datetime(m[c], errors="coerce")
    for c in ("Invoice No", "Order ID", "AWB"):
        m[c] = m[c].astype(str).replace({"nan": None, "None": None, "": None, "NaT": None})
        m[c] = m[c].str.replace(r"\.0$", "", regex=True)
    m["RTS Exact"] = m["RTS Exact"].map(
        lambda v: None if pd.isna(v) or str(v).strip() == "" else str(v).strip().upper() in ("TRUE", "1", "YES"))
    return m.dropna(subset=["Invoice No"])


def merge_master(old, new):
    """New data wins; a blank in the new data never erases an old value.
    Orders missing from today's input stay in the master untouched.
    A PENDING-<order id> row is dropped once that order shows up with an invoice."""
    old, new = old.reindex(columns=DATA_COLS), new.reindex(columns=DATA_COLS)
    for c in DATA_COLS:                  # same dtype on both sides (an empty / newly added column is float)
        if c in TIME_COLS:
            old[c], new[c] = pd.to_datetime(old[c], errors="coerce"), pd.to_datetime(new[c], errors="coerce")
        else:
            old[c], new[c] = old[c].astype(object), new[c].astype(object)
    old = old.drop_duplicates("Invoice No", keep="last").set_index("Invoice No")
    new = new.drop_duplicates("Invoice No", keep="last").set_index("Invoice No")
    both = new.index.intersection(old.index)
    merged = new.combine_first(old)
    # keep RTS Time and its exactness flag together
    take_old_rts = both[new.loc[both, "RTS Time"].isna() & old.loc[both, "RTS Time"].notna()]
    merged.loc[take_old_rts, ["RTS Time", "RTS Exact"]] = old.loc[take_old_rts, ["RTS Time", "RTS Exact"]]

    invoiced_orders = set(merged.loc[~merged.index.str.startswith(PENDING_PREFIX), "Order ID"].dropna()
                          .astype(str).str.replace(r"\.0$", "", regex=True))
    stale = [k for k in merged.index
             if k.startswith(PENDING_PREFIX) and k[len(PENDING_PREFIX):] in invoiced_orders]
    merged = merged.drop(index=stale)

    # one row per AJIO order: if an order id still shows up twice, keep the invoiced / newest row
    out = merged.reset_index().reindex(columns=DATA_COLS)
    oid = out["Order ID"].astype(str).str.replace(r"\.0$", "", regex=True)
    has_oid = out["Order ID"].notna() & ~oid.isin(["", "nan", "None"])
    rank = out.assign(_oid=oid, _pending=out["Invoice No"].astype(str).str.startswith(PENDING_PREFIX))
    rank = rank[has_oid].sort_values(["_pending", "Last Updated"], ascending=[True, False])
    dup_idx = rank.index[rank.duplicated("_oid", keep="first")]
    out = out.drop(index=dup_idx).reset_index(drop=True)

    stats = {"new": len(new.index.difference(old.index)), "updated": len(both),
             "kept (not in input)": len(old.index.difference(new.index)),
             "pending rows replaced by invoice": len(stale), "duplicates removed": len(dup_idx)}
    return out, stats


def apply_tracking(m, summ, ev):
    """Fill pickup scans / courier status from this run's tracking into master rows,
    including orders whose OMS page was not re-read. Known pickup times are never changed."""
    if m.empty or summ.empty:
        return m
    m = m.copy()
    awb = m["AWB"].astype(str)
    pk = pickup_times(ev).drop_duplicates("AWB").set_index("AWB")
    for col in ("Pickup Time", "Pickup Attempt Time"):
        m[col] = m[col].fillna(awb.map(pk[col]))
    status = awb.map(summ.drop_duplicates("AWB").set_index("AWB")["STATUS"].astype(str))
    known = status.notna() & ~status.isin(["UNKNOWN", "NOT_FOUND", "", "nan"])
    m.loc[known, "Courier Status"] = status[known]
    return m


def needs_oms_refresh(df):
    """True for orders whose OMS data can still change: not invoiced, not processed, or not picked up."""
    return (df["Invoice No"].astype(str).str.startswith(PENDING_PREFIX)
            | df["RTS Time"].isna()
            | df["Pickup Status"].isin(OPEN_PICKUP))


def ready_time(r):
    """When the parcel was ready for the courier: RTS time if it has a clock time,
    else the processed (invoice) time. Dispatch Log RTS values are date-only."""
    exact = str(r.get("RTS Exact")).strip().lower() not in ("false", "0", "nan", "none", "")
    if pd.notna(r["RTS Time"]) and exact:
        return r["RTS Time"]
    return r["Processed Time"]


def apply_status(df, as_of):
    df = df.copy()
    dl = df["Dispatch Deadline"]
    today = pd.Timestamp(as_of).normalize()

    def proc_flag(r, dl):
        if pd.isna(r["RTS Time"]):
            if pd.isna(r["Pickup Time"]):
                return "Not Processed"
            return "On Time" if r["Pickup Time"] <= dl else "Late"  # picked => was processed
        return "On Time" if r["RTS Time"] <= dl else "Late"

    def pickup_flag(r, dl):
        if pd.notna(r["Pickup Time"]):
            return "On Time" if r["Pickup Time"] <= dl else "Late"
        if CANCEL_ITEM_STATUSES.search(str(r["OMS Item Status"])) or str(r["Courier Status"]) == "CANCELLED":
            return "Cancelled"
        # No pickup scan, but courier shows movement => it was picked, time unknown
        if (str(r["Courier Status"]) not in NOT_PICKED_COURIER_STATUSES
                or PICKED_ITEM_STATUSES.search(str(r["OMS Item Status"]))):
            return "Picked (Time N/A)"
        # Processed in today's second half: the courier collects it on tomorrow's pickup.
        ready = ready_time(r)
        if pd.notna(ready) and ready.normalize() == today and ready.time() >= NEXT_DAY_AFTER:
            return NEXT_DAY_PICKUP
        return "Pending Pickup"

    def overall(p, k, dl):
        if k == "Cancelled":
            return "CANCELLED"
        if k in OPEN_PICKUP and as_of <= dl:
            return "PENDING"
        if p != "On Time" or k in ("Late",) + OPEN_PICKUP:
            return "NOT CORRECT"
        if k == "Picked (Time N/A)":
            return "CHECK"
        return "CORRECT"

    rows = list(df.iterrows())
    df["Process Status"] = [proc_flag(r, d) for (_, r), d in zip(rows, dl)]
    df["Pickup Status"] = [pickup_flag(r, d) for (_, r), d in zip(rows, dl)]
    df["TAT Result"] = [overall(p, k, d) for p, k, d in zip(df["Process Status"], df["Pickup Status"], dl)]
    df["Label Printed"] = df.reindex(columns=["ERP Label Time"])["ERP Label Time"].notna().map({True: "YES", False: "NO"})
    return df.sort_values("Order Time").reset_index(drop=True)


def apply_erp_labels(df, labels):
    """ERP shipping-label prints into the master: labels = {AWB: (print time, printed by)}.
    A label time already in the master is kept (the first print counts)."""
    df = df.copy()
    awb = df["AWB"].astype(str).str.strip().str.upper()
    hit = awb.map(lambda a: labels.get(a))
    todo = df["ERP Label Time"].isna() & hit.notna()
    df.loc[todo, "ERP Label Time"] = pd.to_datetime([h[0] for h in hit[todo]], errors="coerce")
    df.loc[todo, "ERP Label By"] = [h[1] for h in hit[todo]]
    return df, int(todo.sum())


# ----------------------------------------------------------------- summaries

TAT_RESULTS = ["CORRECT", "NOT CORRECT", "CHECK", "PENDING", "CANCELLED"]
PROCESS_RESULTS = ["On Time", "Late", "Not Processed"]
PICKUP_RESULTS = ["On Time", "Late", "Pending Pickup", NEXT_DAY_PICKUP, "Picked (Time N/A)", "Cancelled"]


def pct(a, b):
    return a / b if b else 0.0


def summarize(df):
    """Counts used by the Sheet Summary tab and the Telegram message."""
    res = df["TAT Result"].value_counts()
    proc = df["Process Status"].value_counts()
    pick = df["Pickup Status"].value_counts()
    couriers = []
    for courier, g in df.groupby(df["Courier"].fillna("UNKNOWN")):
        pk = g["Pickup Status"].value_counts()
        on, late, pend = pk.get("On Time", 0), pk.get("Late", 0), pk.get("Pending Pickup", 0)
        couriers.append({"Courier": courier, "Orders": len(g), "Pickup On Time": on,
                         "Pickup Late": late, "Pending Pickup": pend,
                         "Picked (Time N/A)": pk.get("Picked (Time N/A)", 0),
                         "On Time %": pct(on, on + late + pend)})
    return {
        "total": len(df),
        "tat": {k: int(res.get(k, 0)) for k in TAT_RESULTS},
        "process": {k: int(proc.get(k, 0)) for k in PROCESS_RESULTS},
        "pickup": {k: int(pick.get(k, 0)) for k in PICKUP_RESULTS},
        "tat_pct": pct(res.get("CORRECT", 0), res.get("CORRECT", 0) + res.get("NOT CORRECT", 0)),
        "pickup_pct": pct(pick.get("On Time", 0),
                          pick.get("On Time", 0) + pick.get("Late", 0) + pick.get("Pending Pickup", 0)),
        "couriers": couriers,
    }


# ------------------------------------------------------------ local xlsx copy

def write_excel(data, path, source_note, as_of):
    df = data.reindex(columns=REPORT_COLS).copy()
    exact = data["RTS Exact"].fillna(True).astype(bool)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="Summary", index=False)
        df.to_excel(xw, sheet_name="TAT Report", index=False)
        late = df["TAT Result"] == "NOT CORRECT"
        df[late].to_excel(xw, sheet_name="Late Orders", index=False)
        data.reindex(columns=DATA_COLS).to_excel(xw, sheet_name="_data", index=False)
        pd.DataFrame({"as_of": [as_of]}).to_excel(xw, sheet_name="_meta", index=False)
        late_exact = exact[late].tolist()

    wb = load_workbook(path)
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    grey = PatternFill("solid", fgColor="E7E6E6")
    fills = {"On Time": green, "CORRECT": green,
             "Late": red, "Pending Pickup": red, "Not Processed": red, "NOT CORRECT": red,
             "Picked (Time N/A)": amber, "CHECK": amber, "Pending": amber, "PENDING": amber,
             NEXT_DAY_PICKUP: amber,
             "Cancelled": grey, "CANCELLED": grey, "YES": green, "NO": red}
    widths = {"Invoice No": 11, "Order ID": 13, "AWB": 18, "Courier": 12, "Process Status": 14,
              "Pickup Status": 13, "TAT Result": 14}
    time_fmt = "DD-MMM-YY hh:mm AM/PM"
    col = {c: i + 1 for i, c in enumerate(df.columns)}
    L = {c: get_column_letter(i) for c, i in col.items()}

    for name, rts_exact in (("TAT Report", exact.tolist()), ("Late Orders", late_exact)):
        ws = wb[name]
        for cell in ws[1]:
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for c, i in col.items():
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 20)
        for r, is_exact in zip(range(2, ws.max_row + 1), rts_exact):
            for c in ("Order Time", "Dispatch Deadline", "Processed Time", "RTS Time", "Pickup Attempt Time", "Pickup Time",
                      "ERP Label Time"):
                ws.cell(r, col[c]).number_format = time_fmt
            if not is_exact:  # Dispatch Log gives date only
                ws.cell(r, col["RTS Time"]).number_format = "DD-MMM-YY"
            ws.cell(r, col["Last Updated"]).number_format = "DD-MMM-YY"
            for row in ws.iter_rows(min_row=r, max_row=r):
                for cell in row:
                    cell.font = font
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for c in STATUS_COLS + ["Label Printed"]:
            rng = f"{L[c]}2:{L[c]}{max(ws.max_row, 2)}"
            for txt, fill in fills.items():
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))

    ws = wb["Summary"]
    for r, row in enumerate(summary_rows(summarize(data), as_of, source_note), start=1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(r, j, v)
            c.font = font
            if isinstance(v, float) and 0 <= v <= 1 and j > 1:
                c.number_format = "0.0%"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    for j, w in enumerate((18, 10, 16, 10, 16, 10, 12), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    wb["_data"].sheet_state = "hidden"
    wb["_meta"].sheet_state = "hidden"
    wb.active = 0
    wb.save(path)


def summary_rows(s, as_of, source_note):
    """The Summary tab as a plain grid (shared by the Sheet and the xlsx)."""
    rows = [
        ["AJIO Dispatch TAT Summary"],
        ["Rule (10:10): order before 10 AM -> dispatch same day; 10 AM or later -> next day (by 11:59 PM)."],
        [f"Tracking as of {as_of:%d-%b-%Y %I:%M %p}  |  {source_note}" if pd.notna(as_of)
         else f"Tracking time not known  |  {source_note}"],
        [],
        ["TAT Result", "Orders", "Process", "Orders", "Pickup", "Orders"],
    ]
    for i in range(max(len(TAT_RESULTS), len(PICKUP_RESULTS))):
        row = []
        for lst, counts in ((TAT_RESULTS, s["tat"]), (PROCESS_RESULTS, s["process"]),
                            (PICKUP_RESULTS, s["pickup"])):
            row += [lst[i], counts[lst[i]]] if i < len(lst) else ["", ""]
        rows.append(row)
    rows += [
        [],
        ["Total", s["total"]],
        ["TAT %", s["tat_pct"], "Pickup On Time %", s["pickup_pct"]],
        [],
        ["Courier", "Orders", "Pickup On Time", "Pickup Late", "Pending Pickup", "Picked (Time N/A)", "On Time %"],
    ]
    for c in s["couriers"]:
        rows.append([c["Courier"], c["Orders"], c["Pickup On Time"], c["Pickup Late"],
                     c["Pending Pickup"], c["Picked (Time N/A)"], c["On Time %"]])
    rows += [[], ["Picked (Time N/A) / CHECK = courier shows shipment moved or delivered but has no pickup scan "
                  "(Shadowfax, Delhivery). RTS shown as date only = time not available in OMS."],
             [f"{NEXT_DAY_PICKUP} = processed today after {NEXT_DAY_AFTER:%I:%M %p}, courier collects it "
              "tomorrow; not counted as Pending Pickup."]]
    return rows
