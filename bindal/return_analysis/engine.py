"""
Build return-analysis templates and final reports for Myntra (MAIN / DBCC / FMCC), Ajio and Flipkart.

The rules are those of the hand-run tool BINDAL EXPORTS/RETURN ANALYSIS/main.py (kept unchanged);
the bot fetches the inputs itself and calls build_all() (bindal/return_analysis/pipeline.py).

Expected layout (relative to this file):
    Input/Myntra/<ACCOUNT>/*Seller_Orders_Report*.csv  Myntra order (seller orders) report
    Input/Myntra/<ACCOUNT>/*Seller_Returns_Report*.csv Myntra returns reports (any number)
    Input/Ajio/DropShipOrderReports*.xlsx             Ajio order report
    Input/Ajio/DropShipRtvReports*.xlsx               Ajio return (RTV) reports (any number)
    Input/Flipkart/*.xlsx                             Flipkart sales report (sheet "Sales Report")
                                                      and returns reports (sheet "Returns", any number)
    Supporting/MasterSheet.xlsx                       SKU master
    Supporting/Cost Sheet.xlsx                        style-level cost (FINAL COST), optional

Output:
    Template/<PORTAL> <MON-YY> TEMPLATE.csv
    Output/<MON-YY>/<PORTAL> SALE RETURN ANALYSIS <MON-YY>.xlsx

Rules (with exchange: exchanges / replacements count as Return)
    Myntra: keep only order status C (drops F = cancelled, RTO, SH = in transit).
            Return creation date set -> Return, otherwise Delivered.
            Return reason comes from the returns reports, joined on seller_order_id.
            SALE PRICE = final amount.
    Ajio:   drop Cancelled orders, Delivered + Shipped count as sale.
            Drop every order line that has an RTO return record.
            Customer returns (RVP / REPL / SELF_SHIP*) with Return QTY > 0 -> Return.
            Return requests that only ever had Return QTY = 0 stay Delivered.
            SALE PRICE = Selling Price (per unit).
    Flipkart: sale qty per order item = Sale events - Cancellation events (cancel + re-invoice nets out).
            MONTH = buyer invoice month.
            Drop every order item with an open courier_return (RTO), cancelled returns are ignored.
            Open customer_return (refund / exchange / replace) -> Return, reason = return_sub_reason.
            SALE PRICE = final invoice amount per unit.
    COST:   FINAL COST from Cost Sheet, matched on MY SKU without the size suffix (blank if no Cost Sheet).

Usage:
    python main.py
    python main.py --input "path/to/Input" --supporting "path/to/Supporting" --out "path/to/root"
"""
import argparse
import glob
import os
import re
import warnings

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

BASE = os.path.dirname(os.path.abspath(__file__))
MYNTRA_ACCOUNTS = ['MAIN', 'DBCC', 'FMCC']
COLUMNS = ['MONTH', 'PORTAL', 'BRAND', 'ORDER ID', 'SELLER SKU CODE', 'MY SKU', 'QTY', 'CATEGORY',
           'ORDER STATUS', 'RETURN/EXCHANGE REASON', 'SALE PRICE', 'COST']
AJIO_DATE_FMT = '%a %b %d %H:%M:%S %Y'
# Different brand names that are the same brand, merged in the combined reports (lower case: lower case)
BRAND_ALIASES = {'fashfun clothing': 'fashfun'}
FLIPKART_BRANDS = ['all about you', 'Dressberry', 'FashFun', 'Serein', 'Murcia', 'Leia', 'Fyre Rose']
# Master "Channel Name" values whose PORTAL LISTING DATE counts as the launch date on each portal
PORTAL_CHANNELS = {
    'MYNTRA MAIN': ['MYNTRA', 'MYNTRA PPMP'], 'MYNTRA DBCC': ['MYNTRA DBCC'], 'MYNTRA FMCC': ['MYNTRA FMCC'],
    'AJIO': ['AJIO', 'AJIO API'], 'FLIPKART': ['FLIPKART'],
}

ZONES = ['NORTH', 'WEST', 'EAST', 'SOUTH']
STATE_ZONE = {
    'Delhi': 'NORTH', 'Haryana': 'NORTH', 'Punjab': 'NORTH', 'Himachal Pradesh': 'NORTH',
    'Jammu and Kashmir': 'NORTH', 'Ladakh': 'NORTH', 'Chandigarh': 'NORTH', 'Uttar Pradesh': 'NORTH',
    'Uttarakhand': 'NORTH', 'Rajasthan': 'NORTH',
    'Maharashtra': 'WEST', 'Gujarat': 'WEST', 'Goa': 'WEST', 'Madhya Pradesh': 'WEST',
    'Dadra and Nagar Haveli and Daman and Diu': 'WEST',
    'West Bengal': 'EAST', 'Bihar': 'EAST', 'Jharkhand': 'EAST', 'Odisha': 'EAST', 'Chhattisgarh': 'EAST',
    'Assam': 'EAST', 'Arunachal Pradesh': 'EAST', 'Manipur': 'EAST', 'Meghalaya': 'EAST', 'Mizoram': 'EAST',
    'Nagaland': 'EAST', 'Tripura': 'EAST', 'Sikkim': 'EAST', 'Andaman and Nicobar Islands': 'EAST',
    'Karnataka': 'SOUTH', 'Kerala': 'SOUTH', 'Tamil Nadu': 'SOUTH', 'Telangana': 'SOUTH',
    'Andhra Pradesh': 'SOUTH', 'Puducherry': 'SOUTH', 'Lakshadweep': 'SOUTH',
}
# State codes and spellings seen in portal reports -> name used in STATE_ZONE
STATE_ALIASES = {
    'DL': 'Delhi', 'HR': 'Haryana', 'PB': 'Punjab', 'HP': 'Himachal Pradesh', 'JK': 'Jammu and Kashmir',
    'JAMMU & KASHMIR': 'Jammu and Kashmir', 'LA': 'Ladakh', 'CH': 'Chandigarh', 'UP': 'Uttar Pradesh',
    'UT': 'Uttarakhand', 'UK': 'Uttarakhand', 'RJ': 'Rajasthan', 'MH': 'Maharashtra', 'GJ': 'Gujarat',
    'GA': 'Goa', 'MP': 'Madhya Pradesh', 'DN': 'Dadra and Nagar Haveli and Daman and Diu',
    'DD': 'Dadra and Nagar Haveli and Daman and Diu',
    'DADRA AND NAGAR HAVELI': 'Dadra and Nagar Haveli and Daman and Diu',
    'DAMAN AND DIU': 'Dadra and Nagar Haveli and Daman and Diu',
    'WB': 'West Bengal', 'BR': 'Bihar', 'JH': 'Jharkhand', 'OR': 'Odisha', 'OD': 'Odisha', 'ORISSA': 'Odisha',
    'CT': 'Chhattisgarh', 'CG': 'Chhattisgarh', 'AS': 'Assam', 'AR': 'Arunachal Pradesh', 'MN': 'Manipur',
    'ML': 'Meghalaya', 'MZ': 'Mizoram', 'NL': 'Nagaland', 'TR': 'Tripura', 'SK': 'Sikkim',
    'AN': 'Andaman and Nicobar Islands', 'KA': 'Karnataka', 'KL': 'Kerala', 'KERALAM': 'Kerala',
    'TN': 'Tamil Nadu', 'TG': 'Telangana', 'TS': 'Telangana', 'AP': 'Andhra Pradesh', 'PY': 'Puducherry',
    'PONDICHERRY': 'Puducherry', 'LD': 'Lakshadweep',
}

# Return reasons are merged into these groups in the report. Rules run on the reason in lower case with "_"
# turned into spaces, in this order: the first regex that matches wins (so "wrong size" is checked before "size").
REASON_GROUPS = [
    ('Wrong Size Delivered', r'wrong size|different size|size is different'),
    ('Size Too Big', r'size.*(larg|big)'),
    ('Size Too Small', r'size.*small'),
    ('Fit Issue', r'\bfit\b'),
    ('Damaged / Defective / Missing Item', r'damag|defect|torn|faded|missing|broken'),
    ('Used / Old / Dirty Product', r'used product|looked old|dirty|stain'),
    ('Poor Quality', r'quality'),
    ('Wrong Colour Delivered', r'wrong colou?r|different colou?r received'),
    ('Colour Not As Expected', r'colou?r'),
    ('Wrong Product Delivered', r'different product|completely different|different brand'),
    ('Product Looks Different From Image / Website', r'image|website|displayed'),
    ('Not Needed / Ordered By Mistake', r'mistake|wrong product ordered|not need|not required|does not want|better product'),
    ('Better Price Elsewhere', r'price'),
    ('Delivery Delayed', r'delay|late delivery'),
    ('Did Not Like The Product', r'not look good|not like'),
    ('Others', r'other|generic|thanos|^blank$|^$'),
]

SKU_PREFIX_CATEGORY = {
    'COORD': 'Co-Ords', 'SET': 'Co-Ords', 'DRESS': 'Dresses', 'TOP': 'Tops', 'TOPS': 'Tops',
    'SHIRT': 'Shirts', 'SHRUG': 'Shrug', 'JUMPSUIT': 'Jumpsuit', 'KURTA': 'Kurta Sets',
    'TROUSER': 'Trousers', 'WAISTCOAT': 'Waistcoat', 'SKIRT': 'Skirts',
}
# Checked in order: the first keyword found in the description wins
DESCRIPTION_CATEGORY = [
    ('JUMPSUIT', 'Jumpsuit'), ('SHRUG', 'Shrug'), ('CO-ORD', 'Co-Ords'), ('COORD', 'Co-Ords'),
    ('WAISTCOAT', 'Co-Ords'), ('DRESS', 'Dresses'), ('KAFTAN', 'Dresses'), ('SHIRT', 'Shirts'),
    ('KURTA', 'Kurta Sets'), ('TROUSER', 'Trousers'), ('TOP', 'Tops'), ('TUNIC', 'Tops'),
]


class MissingInput(RuntimeError):
    """A portal's input file is not there; that portal is skipped."""


def norm(s):
    return s.astype(str).str.strip().str.upper()


def find_one(pattern):
    files = sorted(glob.glob(pattern))
    if not files:
        raise MissingInput(f'no file matches {pattern}')
    if len(files) > 1:
        print(f'WARNING: {len(files)} files match {pattern}, using {os.path.basename(files[0])}')
    return files[0]


def load_master(supporting_dir):
    path = os.path.join(supporting_dir, 'MasterSheet.xlsx')
    print(f'Loading master: {path}')
    m = pd.read_excel(path, sheet_name='Master', dtype=str,
                      usecols=['Channel Listing SKU Code', 'Channel Name', 'MYNTRA SKU CODE', 'FINAL SKU',
                               'DESIGN-NO', 'PORTAL LISTING DATE'])
    m = m[m['FINAL SKU'].notna()]
    design_map = m.dropna(subset=['DESIGN-NO']).groupby('FINAL SKU')['DESIGN-NO'].agg(lambda s: s.mode()[0])
    listing = m.dropna(subset=['PORTAL LISTING DATE']).assign(
        date=lambda d: pd.to_datetime(d['PORTAL LISTING DATE'], errors='coerce'))
    launch = listing.groupby(['Channel Name', 'FINAL SKU'])['date'].min()
    keys = pd.concat([
        m[['Channel Listing SKU Code', 'FINAL SKU']].rename(columns={'Channel Listing SKU Code': 'k'}),
        m[['MYNTRA SKU CODE', 'FINAL SKU']].rename(columns={'MYNTRA SKU CODE': 'k'}),
    ]).dropna()
    keys['k'] = norm(keys['k'])
    keys = keys.drop_duplicates()
    conflicts = keys[keys['k'].duplicated(keep=False)]
    sku_map = keys.drop_duplicates('k').set_index('k')['FINAL SKU']

    p = pd.read_excel(path, sheet_name='Portal Combined', dtype=str, usecols=['Seller SKU', 'SIZE SKU']).dropna()
    p['k'] = norm(p['Seller SKU'])
    portal_map = p.drop_duplicates('k').set_index('k')['SIZE SKU']
    return sku_map, portal_map, conflicts, design_map, launch


def launch_dates(launch, portals):
    # Earliest listing date of each FINAL SKU across the channels of one portal (or a list of portals)
    portals = [portals] if isinstance(portals, str) else portals
    channels = [c for p in portals for c in PORTAL_CHANNELS.get(p, []) if c in launch.index.get_level_values(0)]
    if not channels:
        return pd.Series(dtype='datetime64[ns]')
    return pd.concat([launch.loc[c] for c in channels]).groupby(level=0).min()


def load_cost(supporting_dir):
    path = os.path.join(supporting_dir, 'Cost Sheet.xlsx')
    if not os.path.exists(path):
        print(f'WARNING: {path} not found, COST will be blank')
        return pd.Series(dtype=float)
    print(f'Loading cost: {path}')
    c = pd.read_excel(path, dtype={'SKU': str}, usecols=['SKU', 'FINAL COST']).dropna()
    c['k'] = norm(c['SKU'])
    return c.drop_duplicates('k').set_index('k')['FINAL COST']


def cost_for(my_sku, cost_map):
    # Cost Sheet is per style, so try the SKU without its size first, then the full SKU
    sku = my_sku.fillna('').str.strip().str.upper()
    return sku.str.rsplit('-', n=1).str[0].map(cost_map).fillna(sku.map(cost_map))


def category_from_sku(my_sku):
    return my_sku.str.split('-').str[0].str.upper().map(SKU_PREFIX_CATEGORY)


def category_from_description(desc):
    desc = desc.fillna('').str.upper()
    out = pd.Series(pd.NA, index=desc.index, dtype=object)
    for kw, cat in reversed(DESCRIPTION_CATEGORY):
        out = out.mask(desc.str.contains(kw, regex=False), cat)
    return out


def clean_state(s):
    key = s.fillna('').astype(str).str.strip().str.upper()
    by_name = {name.upper(): name for name in STATE_ZONE}
    return key.map(STATE_ALIASES).fillna(key.map(by_name))


def merge_reason(reasons):
    # Unmatched reasons are kept as they are, and main() warns about them
    def one(reason):
        text = ' '.join(str(reason).replace('_', ' ').lower().split())
        for group, pattern in REASON_GROUPS:
            if re.search(pattern, text):
                return group
        return reason
    return reasons.map({r: one(r) for r in reasons.dropna().unique()})


def finish(df):
    return df[COLUMNS].sort_values(['MY SKU', 'BRAND', 'ORDER STATUS'], na_position='last')


def split_partial_returns(df, ret_qty, is_ret):
    # Split partially returned multi-qty lines into a Return row and a Delivered row
    rq = ret_qty.clip(upper=df['QTY']).astype(int)
    part = is_ret & (rq < df['QTY'])
    if part.any():
        rest = df[part].copy()
        rest['QTY'] = df.loc[part, 'QTY'] - rq[part]
        rest['ORDER STATUS'] = 'Delivered'
        rest['RETURN/EXCHANGE REASON'] = 'blank'
        df.loc[part, 'QTY'] = rq[part]
        df = pd.concat([df, rest])
    return df, int(part.sum())


# ----------------------------------------------------------------------------- Myntra

def build_myntra(input_dir, account, sku_map, conflicts, cost_map):
    folder = os.path.join(input_dir, 'Myntra', account)
    o = pd.read_csv(find_one(os.path.join(folder, '*Seller_Orders_Report*.csv')), dtype=str, low_memory=False)
    stats = {'raw': len(o)}
    for status, label in [('F', 'removed_cancelled'), ('RTO', 'removed_rto'), ('SH', 'removed_in_transit')]:
        stats[label] = int((o['order status'] == status).sum())
    o = o[o['order status'] == 'C']
    stats['removed_other'] = stats['raw'] - len(o) - stats['removed_cancelled'] - stats['removed_rto'] - stats['removed_in_transit']

    ret_files = glob.glob(os.path.join(folder, '*Seller_Returns_Report*.csv'))
    if ret_files:
        r = pd.concat([pd.read_csv(f, dtype=str) for f in ret_files], ignore_index=True)
        r = r[r['type'] == 'Return'].drop_duplicates('seller_order_id').set_index('seller_order_id')
        reasons = r['return_reason']
    else:
        print(f'WARNING: no returns report for MYNTRA {account}, reasons will be blank')
        reasons = pd.Series(dtype=str)

    is_ret = o['return creation date'].notna()
    my_sku = norm(o['seller sku code']).map(sku_map)
    df = pd.DataFrame({
        'MONTH': pd.to_datetime(o['created on'].str[:10]).dt.strftime('%m-%Y'),
        'PORTAL': f'MYNTRA {account}',
        'BRAND': o['brand'],
        'ORDER ID': o['seller order id'],
        'SELLER SKU CODE': o['seller sku code'],
        'MY SKU': my_sku,
        'QTY': 1,
        'CATEGORY': o['article type'],
        'ORDER STATUS': is_ret.map({True: 'Return', False: 'Delivered'}),
        'RETURN/EXCHANGE REASON': o['seller order id'].map(reasons).where(is_ret, 'blank').fillna('blank'),
        'SALE PRICE': pd.to_numeric(o['final amount'], errors='coerce'),
        'COST': cost_for(my_sku, cost_map),
        'STATE': o['state'],
    })

    used = set(norm(o['seller sku code']))
    bad = conflicts[conflicts['k'].isin(used)]
    if len(bad):
        print(f'WARNING: MYNTRA {account} SKU codes with more than one FINAL SKU in master: '
              f'{bad.groupby("k")["FINAL SKU"].apply(list).to_dict()}')
    stats['return_without_reason'] = int((is_ret & df['RETURN/EXCHANGE REASON'].eq('blank')).sum())
    return df, stats


# ----------------------------------------------------------------------------- Ajio

def parse_ajio_date(s):
    return pd.to_datetime(s.str.replace(' IST', '', regex=False), format=AJIO_DATE_FMT, errors='coerce')


def build_ajio(input_dir, sku_map, portal_map, cost_map):
    folder = os.path.join(input_dir, 'Ajio')
    o = pd.read_excel(find_one(os.path.join(folder, 'DropShipOrderReports*.xlsx')), dtype=str).reset_index(drop=True)
    rtv_files = sorted(glob.glob(os.path.join(folder, 'DropShipRtvReports*.xlsx')))
    r = pd.concat([pd.read_excel(f, dtype=str) for f in rtv_files], ignore_index=True)

    r['key'] = r['Cust Order No'].str.strip() + '|' + r['SELLER SKU'].str.strip()
    r['qty'] = pd.to_numeric(r['Return QTY'], errors='coerce').fillna(0)
    r['is_rto'] = r['Return Type'] == 'RTO'
    r['created'] = parse_ajio_date(r['Return Created Date'])
    # One record per order line: RTO first, then a record with qty > 0, then the latest
    r = r.sort_values(['key', 'is_rto', 'qty', 'created'], ascending=[True, False, False, False])
    any_rto = r.groupby('key')['is_rto'].any()
    max_qty = r.groupby('key')['qty'].max()
    best = r.drop_duplicates('key').set_index('key')

    stats = {'raw': len(o)}
    o = o[o['Status'] != 'Cancelled'].copy()
    stats['removed_cancelled'] = stats['raw'] - len(o)
    o['key'] = o['Cust Order No'].str.strip() + '|' + o['Seller SKU'].str.strip()
    rto = o['key'].map(any_rto).fillna(False).astype(bool)
    stats['removed_rto'] = int(rto.sum())
    o = o[~rto]

    has_ret = o['key'].isin(best.index)
    ret_qty = o['key'].map(max_qty).fillna(0)
    is_ret = has_ret & (ret_qty > 0)
    stats['zero_qty_requests_kept_delivered'] = int((has_ret & (ret_qty == 0)).sum())

    sk = norm(o['Seller SKU'])
    my_sku = sk.map(sku_map).fillna(sk.map(portal_map))
    category = category_from_sku(my_sku).fillna(category_from_description(o['Description']))

    prefix = o['Seller SKU'].str.split('-').str[0]
    brand_by_prefix = (o.assign(p=prefix).dropna(subset=['Brand'])
                       .drop_duplicates('p').set_index('p')['Brand'])
    brand = o['Brand'].fillna(prefix.map(brand_by_prefix))

    qty = pd.to_numeric(o['Shipped QTY'], errors='coerce').fillna(pd.to_numeric(o['Order Qty'], errors='coerce')).fillna(1).astype(int)
    df = pd.DataFrame({
        'MONTH': parse_ajio_date(o['Cust Order Date']).dt.strftime('%m-%Y'),
        'PORTAL': 'AJIO',
        'BRAND': brand,
        'ORDER ID': o['Cust Order No'],
        'SELLER SKU CODE': o['Seller SKU'],
        'MY SKU': my_sku,
        'QTY': qty,
        'CATEGORY': category,
        'ORDER STATUS': is_ret.map({True: 'Return', False: 'Delivered'}),
        'RETURN/EXCHANGE REASON': o['key'].map(best['Cust Return Reason']).where(is_ret, 'blank').fillna('blank'),
        'SALE PRICE': pd.to_numeric(o['Selling Price'], errors='coerce'),
        'COST': cost_for(my_sku, cost_map),
        'STATE': pd.NA,  # Ajio order / RTV reports carry no customer state
    })
    df, stats['split_partial_returns'] = split_partial_returns(df, ret_qty, is_ret)
    return df, stats


# ----------------------------------------------------------------------------- Flipkart

def unquote(s):
    return s.astype(str).str.strip().str.strip('"').str.strip()


def flipkart_files(folder):
    sales, returns = [], []
    for f in sorted(glob.glob(os.path.join(folder, '*.xlsx'))):
        if os.path.basename(f).startswith('~$'):
            continue
        sheets = pd.ExcelFile(f).sheet_names
        if 'Sales Report' in sheets:
            sales.append(f)
        elif 'Returns' in sheets:
            returns.append(f)
        else:
            print(f'WARNING: {os.path.basename(f)} has neither a "Sales Report" nor a "Returns" sheet, skipped')
    if not sales:
        raise MissingInput(f'no Flipkart sales report (sheet "Sales Report") in {folder}')
    return sales, returns


def flipkart_brand(title):
    t = title.str.lower()
    out = title.str.split().str[0]
    for b in FLIPKART_BRANDS:
        out = out.mask(t.str.startswith(b.lower()), b)
    return out


def build_flipkart(input_dir, sku_map, portal_map, cost_map, category_by_sku):
    folder = os.path.join(input_dir, 'Flipkart')
    sales_files, ret_files = flipkart_files(folder)
    s = pd.concat([pd.read_excel(f, sheet_name='Sales Report', dtype=str) for f in sales_files], ignore_index=True)
    s['item'] = s['Order Item ID'].str.strip()
    s['qty'] = pd.to_numeric(s['Item Quantity'], errors='coerce').fillna(1)
    stats = {'raw_sale_events': int((s['Event Type'] == 'Sale').sum())}

    # Net qty per order item: a cancelled invoice is usually re-invoiced, so Sale - Cancellation
    sold = s[s['Event Type'] == 'Sale']
    cancelled = s[s['Event Sub Type'] == 'Cancellation'].groupby('item')['qty'].sum()
    o = sold.sort_values('Buyer Invoice Date').drop_duplicates('item', keep='last').set_index('item')
    o['QTY'] = (sold.groupby('item')['qty'].sum() - cancelled.reindex(o.index).fillna(0)).astype(int)
    stats['removed_cancelled'] = int((o['QTY'] <= 0).sum())
    o = o[o['QTY'] > 0]

    if ret_files:
        r = pd.concat([pd.read_excel(f, sheet_name='Returns', dtype=str) for f in ret_files], ignore_index=True)
        r = r.drop_duplicates('return_id')
        r['item'] = r['order_item_id'].str.replace('OI:', '', regex=False).str.strip()
        r = r[r['return_status'] != 'cancelled']
        r['qty'] = pd.to_numeric(r['quantity'], errors='coerce').fillna(1)
        r['date'] = pd.to_datetime(r['return_approval_date'].fillna(r['return_requested_date']), errors='coerce')
    else:
        print('WARNING: no Flipkart returns report, every sale will be Delivered')
        r = pd.DataFrame(columns=['item', 'return_type', 'qty', 'date', 'return_reason', 'return_sub_reason'])

    rto_items = set(r.loc[r['return_type'] == 'courier_return', 'item'])
    rto = o.index.isin(rto_items)
    stats['removed_rto'] = int(rto.sum())
    o = o[~rto]

    cr = r[r['return_type'] == 'customer_return'].sort_values('date')
    ret_qty = cr.groupby('item')['qty'].sum().reindex(o.index).fillna(0)
    last = cr.drop_duplicates('item', keep='last').set_index('item')
    reason = last['return_sub_reason'].fillna(last['return_reason'])
    is_ret = ret_qty > 0

    seller_sku = unquote(s.set_index('item')['SKU'].groupby(level=0).first().reindex(o.index)).str.replace('SKU:', '', regex=False)
    title = unquote(o['Product Title/Description'])
    sk = norm(seller_sku)
    my_sku = sk.map(sku_map).fillna(sk.map(portal_map))
    category = (category_from_sku(my_sku).fillna(my_sku.map(category_by_sku))
                .fillna(category_from_description(title)))
    df = pd.DataFrame({
        'MONTH': pd.to_datetime(o['Buyer Invoice Date'].str[:10]).dt.strftime('%m-%Y'),
        'PORTAL': 'FLIPKART',
        'BRAND': flipkart_brand(title),
        'ORDER ID': o['Order ID'],
        'SELLER SKU CODE': seller_sku,
        'MY SKU': my_sku,
        'QTY': o['QTY'],
        'CATEGORY': category,
        'ORDER STATUS': is_ret.map({True: 'Return', False: 'Delivered'}),
        'RETURN/EXCHANGE REASON': pd.Series(o.index, index=o.index).map(reason).where(is_ret, 'blank').fillna('blank'),
        'SALE PRICE': pd.to_numeric(o['Final Invoice Amount (Price after discount+Shipping Charges)'], errors='coerce') / o['QTY'],
        'COST': cost_for(my_sku, cost_map),
        'STATE': o["Customer's Delivery State"],
    })
    df, stats['split_partial_returns'] = split_partial_returns(df, ret_qty, is_ret)
    return df.reset_index(drop=True), stats


# ----------------------------------------------------------------------------- report

THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEFT = Alignment(horizontal='left')
RIGHT = Alignment(horizontal='right')
CENTER = Alignment(horizontal='center')


def fill(rgb):
    return PatternFill('solid', fgColor=rgb)


SUMMARY_TITLE = fill('FFFAFD88')
SUMMARY_HEAD = fill('FFE7FD88')
BLUE_TITLE = fill('FF4F81BD')
BLUE_SECTION = fill('FFD9E2F3')
BLUE_HEAD = fill('FFEEF4FB')


def pct(part, whole, digits=2):
    return f'{(part / whole * 100 if whole else 0):.{digits}f}%'


def agg(df, by):
    g = df.groupby(by).agg(sales=('QTY', 'sum'), returns=('_ret', 'sum'))
    g['rate'] = (g['returns'] / g['sales'] * 100).fillna(0)
    return g


def top(g, col, n=None):
    # Highest first, ties in A-Z order
    g = g.sort_index(kind='stable').sort_values(col, ascending=False, kind='stable')
    return g if n is None else g.head(n)


def reason_counts(df):
    ret = df[df['ORDER STATUS'] == 'Return']
    return top(ret.groupby('RETURN/EXCHANGE REASON')['QTY'].sum().to_frame('qty'), 'qty')['qty']


class Sheet:
    def __init__(self, ws):
        self.ws = ws

    def cell(self, row, col, value, font=None, fill_=None, align=None, fmt=None):
        c = self.ws.cell(row=row, column=col, value=value)
        c.border = BORDER
        if font:
            c.font = font
        if fill_:
            c.fill = fill_
        if align:
            c.alignment = align
        if fmt:
            c.number_format = fmt
        return c

    def title(self, row, col, text, span, font, fill_, align):
        for i in range(span):
            self.cell(row, col + i, text if i == 0 else None, font, fill_, align)
        if span > 1:
            self.ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)


def write_summary(ws, df):
    sh = Sheet(ws)
    bold = Font(bold=True)

    def section(row, text, span, headers):
        sh.title(row, 1, text, span, bold, SUMMARY_TITLE, LEFT)
        for i, h in enumerate(headers):
            sh.cell(row + 1, 1 + i, h, bold, SUMMARY_HEAD, LEFT)
        return row + 2

    def line(row, values):
        for i, v in enumerate(values):
            sh.cell(row, 1 + i, v, align=LEFT)

    sales, returns = int(df['QTY'].sum()), int(df['_ret'].sum())
    sh.title(1, 1, 'OVERALL METRICS', 2, bold, SUMMARY_TITLE, LEFT)
    line(2, ['Total Sales Units', sales])
    line(3, ['Total Return Units', returns])
    line(4, ['Overall Return Rate (%)', pct(returns, sales)])

    row = section(6, 'BRAND WISE SUMMARY', 4, ['BRAND', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for brand, g in agg(df, 'BRAND').iterrows():
        line(row, [brand, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'CATEGORY WISE SUMMARY', 4, ['CATEGORY', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for cat, g in top(agg(df, 'CATEGORY'), 'sales').iterrows():
        line(row, [cat, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'TOP RETURN REASONS', 3, ['RETURN REASON', 'TOTAL QUANTITY', 'PERCENTAGE'])
    for reason, qty in reason_counts(df).head(20).items():
        line(row, [reason, int(qty), pct(qty, returns)])
        row += 1

    skus = agg(df, 'MY SKU')
    row = section(row + 1, 'TOP SELLING SKUs', 4, ['MY SKU', 'SALES QUANTITY', 'RETURN QUANTITY', 'RETURN RATE (%)'])
    for sku, g in top(skus, 'sales', 20).iterrows():
        line(row, [sku, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'TOP RETURNED SKUs', 4, ['MY SKU', 'RETURN QUANTITY', 'SALES QUANTITY', 'RETURN RATE (%)'])
    for sku, g in top(skus, 'returns', 20).iterrows():
        line(row, [sku, int(g.returns), int(g.sales), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'MONTH WISE SUMMARY', 4, ['MONTH', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for month, g in agg(df, '_month').iterrows():
        line(row, [month.strftime('%Y-%m'), int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 18.14
    ws.column_dimensions['D'].width = 16.57


def write_month_block(ws, df, col, heading):
    sh = Sheet(ws)
    bold = Font(bold=True)

    def section(row, text, headers):
        sh.title(row, col, text, 4, bold, BLUE_SECTION, CENTER)
        for i, h in enumerate(headers):
            sh.cell(row + 1, col + i, h, bold, BLUE_HEAD, CENTER)
        return row + 2

    def line(row, values):
        for i, v in enumerate(values):
            if isinstance(v, str) and v.endswith('%'):
                sh.cell(row, col + i, v, align=RIGHT)
            elif isinstance(v, str):
                sh.cell(row, col + i, v, align=LEFT)
            else:
                sh.cell(row, col + i, v, align=RIGHT, fmt='0')

    sh.title(1, col, heading, 4, Font(bold=True, size=12, color='FFFFFFFF'), BLUE_TITLE, CENTER)
    sales, returns = int(df['QTY'].sum()), int(df['_ret'].sum())
    sh.title(3, col, 'OVERALL METRICS', 4, bold, BLUE_SECTION, CENTER)
    line(4, ['Total Sales Units', sales])
    line(5, ['Total Return Units', returns])
    line(6, ['Overall Return Rate (%)', pct(returns, sales)])

    row = section(8, 'STATUS WISE SUMMARY', ['STATUS', 'QTY', '%'])
    for status, qty in [('TOTAL', sales), ('Delivered', sales - returns), ('RETURN', returns), ('EXCHANGE', 0), ('RTO', 0)]:
        line(row, [status, qty, pct(qty, sales, 1)])
        row += 1

    row = section(row + 1, 'BRAND WISE SUMMARY', ['BRAND', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for brand, g in top(agg(df, 'BRAND'), 'sales').iterrows():
        line(row, [brand, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'CATEGORY WISE SUMMARY', ['CATEGORY', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for cat, g in top(agg(df, 'CATEGORY'), 'sales').iterrows():
        line(row, [cat, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    skus = agg(df, 'MY SKU')
    row = section(row + 1, 'TOP SELLING SKUs', ['MY SKU', 'SALES QTY', 'RETURN QTY', 'RETURN RATE (%)'])
    for sku, g in top(skus, 'sales', 10).iterrows():
        line(row, [sku, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'TOP RETURNED SKUs', ['MY SKU', 'RETURN QTY', 'SALES QTY', 'RETURN RATE (%)'])
    for sku, g in top(skus, 'returns', 10).iterrows():
        line(row, [sku, int(g.returns), int(g.sales), f'{g.rate:.2f}%'])
        row += 1

    row = section(row + 1, 'TOP RETURN REASONS', ['ORDER STATUS', 'RETURN REASON', 'TOTAL QUANTITY', 'PERCENTAGE'])
    reasons = reason_counts(df)
    for reason, qty in reasons.items():
        line(row, ['RETURN', reason, int(qty), pct(qty, returns)])
        row += 1
    row += 1
    for status in ['EXCHANGE', 'RTO']:
        line(row, [status, 'No reasons', 0, '0.00%'])
        row += 1
    for reason, qty in reasons.items():
        line(row, ['ALL', reason, int(qty), pct(qty, returns)])
        row += 1

    for i, width in enumerate([24.71, 14.71]):
        ws.column_dimensions[get_column_letter(col + i)].width = width


def write_monthly_summary(ws, df):
    months = sorted(df['_month'].unique())
    for i, month in enumerate(months):
        write_month_block(ws, df[df['_month'] == month], 1 + 5 * i, f'MONTH: {month.strftime("%b-%Y").upper()}')
    write_month_block(ws, df, 1 + 5 * len(months), 'OVERALL')
    ws.row_dimensions[1].height = 15.75
    ws.freeze_panes = 'A2'


def write_category_performance(ws, df):
    sh = Sheet(ws)
    bold = Font(bold=True)
    white = Font(bold=True, size=12, color='FFFFFFFF')

    def section(row, text, headers):
        sh.title(row, 1, text, len(headers), white, BLUE_TITLE, CENTER)
        ws.row_dimensions[row].height = 15.75
        for i, h in enumerate(headers):
            sh.cell(row + 1, 1 + i, h, bold, BLUE_SECTION, CENTER)
        return row + 2

    def line(row, values):
        for i, v in enumerate(values):
            if isinstance(v, str) and v.endswith('%'):
                sh.cell(row, 1 + i, v, align=RIGHT)
            elif isinstance(v, str):
                sh.cell(row, 1 + i, v, align=LEFT)
            else:
                sh.cell(row, 1 + i, v, align=RIGHT, fmt='0')

    sales, returns = int(df['QTY'].sum()), int(df['_ret'].sum())
    sh.title(1, 1, 'CATEGORY PERFORMANCE ANALYSIS', 6, white, BLUE_TITLE, CENTER)
    ws.row_dimensions[1].height = 15.75

    row = section(4, 'OVERALL CATEGORY RANKING',
                  ['RANK', 'CATEGORY', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)', 'SALES SHARE (%)'])
    cats = top(agg(df, 'CATEGORY'), 'sales')
    for rank, (cat, g) in enumerate(cats.iterrows(), 1):
        line(row, [rank, cat, int(g.sales), int(g.returns), f'{g.rate:.2f}%', pct(g.sales, sales)])
        row += 1

    row = section(row + 2, 'MONTH-WISE CATEGORY TREND',
                  ['MONTH', 'CATEGORY', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for month in sorted(df['_month'].unique()):
        for cat, g in top(agg(df[df['_month'] == month], 'CATEGORY'), 'sales').iterrows():
            line(row, [month.strftime('%b-%Y').upper(), cat, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
            row += 1

    row = section(row + 2, 'TOP RETURN REASONS BY CATEGORY',
                  ['CATEGORY', 'RETURN REASON', 'TOTAL QUANTITY', 'CATEGORY RETURN %', 'OVERALL RETURN %'])
    for cat, part in df.groupby('CATEGORY'):
        reasons = reason_counts(part)
        if reasons.empty:
            continue
        cat_returns = reasons.sum()
        for reason, qty in reasons.head(5).items():
            line(row, [cat, reason, int(qty), pct(qty, cat_returns), pct(qty, returns)])
            row += 1
        row += 1

    ws.column_dimensions['A'].width = 22.71
    ws.column_dimensions['C'].width = 16.71
    ws.freeze_panes = 'A2'


class BlueTable(Sheet):
    # Same look as Category Performance: blue title bar, light header row, bordered lines
    def section(self, row, text, headers):
        self.title(row, 1, text, len(headers), Font(bold=True, size=12, color='FFFFFFFF'), BLUE_TITLE, CENTER)
        self.ws.row_dimensions[row].height = 15.75
        for i, h in enumerate(headers):
            self.cell(row + 1, 1 + i, h, Font(bold=True), BLUE_SECTION, CENTER)
        return row + 2

    def line(self, row, values, bold=False):
        font = Font(bold=True) if bold else None
        for i, v in enumerate(values):
            if isinstance(v, str) and v.endswith('%'):
                self.cell(row, 1 + i, v, font, align=RIGHT)
            elif isinstance(v, str):
                self.cell(row, 1 + i, v, font, align=LEFT)
            else:
                self.cell(row, 1 + i, v, font, align=RIGHT, fmt='0')


def write_portal_wise(ws, df):
    sh = BlueTable(ws)
    sales, returns = int(df['QTY'].sum()), int(df['_ret'].sum())
    portals = top(agg(df, 'PORTAL'), 'sales')

    row = sh.section(1, 'PORTAL WISE SALE RETURN', ['PORTAL', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)',
                                                      'SALES SHARE (%)', 'RETURN SHARE (%)'])
    for portal, g in portals.iterrows():
        sh.line(row, [portal, int(g.sales), int(g.returns), f'{g.rate:.2f}%', pct(g.sales, sales), pct(g.returns, returns)])
        row += 1
    sh.line(row, ['TOTAL', sales, returns, pct(returns, sales), '100.00%', '100.00%'], bold=True)
    row += 1

    row = sh.section(row + 2, 'PORTAL WISE BRAND', ['PORTAL', 'BRAND', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for portal in portals.index:
        for brand, g in top(agg(df[df['PORTAL'] == portal], 'BRAND'), 'sales').iterrows():
            sh.line(row, [portal, brand, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
            row += 1

    row = sh.section(row + 2, 'PORTAL WISE CATEGORY', ['PORTAL', 'CATEGORY', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for portal in portals.index:
        for cat, g in top(agg(df[df['PORTAL'] == portal], 'CATEGORY'), 'sales').iterrows():
            sh.line(row, [portal, cat, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
            row += 1

    row = sh.section(row + 2, 'TOP RETURN REASONS BY PORTAL',
                     ['PORTAL', 'RETURN REASON', 'TOTAL QUANTITY', 'PORTAL RETURN %', 'OVERALL RETURN %'])
    for portal in portals.index:
        reasons = reason_counts(df[df['PORTAL'] == portal])
        for reason, qty in reasons.head(5).items():
            sh.line(row, [portal, reason, int(qty), pct(qty, reasons.sum()), pct(qty, returns)])
            row += 1
        if not reasons.empty:
            row += 1

    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 40
    for col in 'CDEF':
        ws.column_dimensions[col].width = 17
    ws.freeze_panes = 'A2'


def write_zone_wise(ws, df):
    sh = BlueTable(ws)
    known = df[df['ZONE'].notna()]
    sales, returns = int(known['QTY'].sum()), int(known['_ret'].sum())

    row = sh.section(1, 'ZONE WISE SALE RETURN', ['ZONE', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)',
                                                    'SALES SHARE (%)', 'RETURN SHARE (%)'])
    zones = agg(known, 'ZONE').reindex(ZONES).fillna(0)
    for zone, g in zones.iterrows():
        sh.line(row, [zone, int(g.sales), int(g.returns), f'{g.rate:.2f}%', pct(g.sales, sales), pct(g.returns, returns)])
        row += 1
    sh.line(row, ['TOTAL', sales, returns, pct(returns, sales), '100.00%', '100.00%'], bold=True)
    row += 1
    missing = int(df.loc[df['ZONE'].isna(), 'QTY'].sum())
    if missing:
        ws.cell(row=row + 1, column=1, value=f'State not available in portal report: {missing} sales units not counted above')
        row += 1

    row = sh.section(row + 2, 'ZONE WISE BRAND', ['ZONE', 'BRAND', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for zone in ZONES:
        for brand, g in top(agg(known[known['ZONE'] == zone], 'BRAND'), 'sales').iterrows():
            sh.line(row, [zone, brand, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
            row += 1

    row = sh.section(row + 2, 'ZONE WISE CATEGORY', ['ZONE', 'CATEGORY', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)'])
    for zone in ZONES:
        for cat, g in top(agg(known[known['ZONE'] == zone], 'CATEGORY'), 'sales').iterrows():
            sh.line(row, [zone, cat, int(g.sales), int(g.returns), f'{g.rate:.2f}%'])
            row += 1

    row = sh.section(row + 2, 'TOP RETURN REASONS BY ZONE',
                     ['ZONE', 'RETURN REASON', 'TOTAL QUANTITY', 'ZONE RETURN %', 'OVERALL RETURN %'])
    for zone in ZONES:
        reasons = reason_counts(known[known['ZONE'] == zone])
        if reasons.empty:
            continue
        for reason, qty in reasons.head(5).items():
            sh.line(row, [zone, reason, int(qty), pct(qty, reasons.sum()), pct(qty, returns)])
            row += 1
        row += 1

    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 30
    for col in 'CDEF':
        ws.column_dimensions[col].width = 17
    ws.freeze_panes = 'A2'


def write_state_wise(ws, df):
    sh = BlueTable(ws)
    sales, returns = int(df['QTY'].sum()), int(df['_ret'].sum())
    states = df.assign(STATE=df['STATE'].fillna('NOT AVAILABLE'), ZONE=df['ZONE'].fillna('-'))

    row = sh.section(1, 'STATE WISE SALE RETURN', ['STATE', 'ZONE', 'SALES UNITS', 'RETURN UNITS', 'RETURN RATE (%)',
                                                     'SALES SHARE (%)', 'RETURN SHARE (%)'])
    g = top(agg(states, ['STATE', 'ZONE']), 'sales')
    # Unknown state always at the bottom
    g = pd.concat([g[g.index.get_level_values(0) != 'NOT AVAILABLE'], g[g.index.get_level_values(0) == 'NOT AVAILABLE']])
    for (state, zone), r in g.iterrows():
        sh.line(row, [state, zone, int(r.sales), int(r.returns), f'{r.rate:.2f}%', pct(r.sales, sales), pct(r.returns, returns)])
        row += 1
    sh.line(row, ['TOTAL', '', sales, returns, pct(returns, sales), '100.00%', '100.00%'], bold=True)

    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 10
    for col in 'CDEFG':
        ws.column_dimensions[col].width = 17
    ws.freeze_panes = 'A3'


def write_zone_mapping(ws, df):
    sh = BlueTable(ws)
    aliases = {}
    for alias, state in STATE_ALIASES.items():
        aliases.setdefault(state, []).append(alias)

    row = sh.section(1, 'STATE TO ZONE MAPPING', ['ZONE', 'STATE', 'CODES / OTHER NAMES'])
    for zone in ZONES:
        for state in sorted(s for s, z in STATE_ZONE.items() if z == zone):
            sh.line(row, [zone, state, ', '.join(aliases.get(state, []))])
            row += 1

    # Every state value exactly as the portal report wrote it, and where it landed
    row = sh.section(row + 2, 'STATE VALUES IN THIS REPORT',
                     ['STATE IN PORTAL REPORT', 'MAPPED STATE', 'ZONE', 'SALES UNITS', 'RETURN UNITS'])
    raw = df.assign(_raw_state=df['_raw_state'].fillna('(blank)').astype(str).str.strip())
    g = top(agg(raw, ['_raw_state']), 'sales')
    first = raw.drop_duplicates('_raw_state').set_index('_raw_state')
    for value, r in g.iterrows():
        state = first.at[value, 'STATE']
        zone = first.at[value, 'ZONE']
        sh.line(row, [value, 'NOT MAPPED' if pd.isna(state) else state, '-' if pd.isna(zone) else zone,
                      int(r.sales), int(r.returns)])
        row += 1

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.freeze_panes = 'A3'


def write_reason_mapping(ws, df):
    sh = BlueTable(ws)
    ret = df[df['ORDER STATUS'] == 'Return']
    returns = int(ret['QTY'].sum())

    row = sh.section(1, 'MERGED RETURN REASONS', ['MERGED REASON', 'RETURN QTY', 'PERCENTAGE'])
    for reason, qty in reason_counts(df).items():
        sh.line(row, [reason, int(qty), pct(qty, returns)])
        row += 1

    # Every reason exactly as the portal report wrote it, and the group it was merged into
    row = sh.section(row + 2, 'REASON VALUES IN THIS REPORT', ['MERGED REASON', 'REASON IN PORTAL REPORT', 'RETURN QTY'])
    g = ret.groupby(['RETURN/EXCHANGE REASON', '_raw_reason'])['QTY'].sum().reset_index()
    order = {r: i for i, r in enumerate(reason_counts(df).index)}
    g = g.sort_values(['RETURN/EXCHANGE REASON', 'QTY'], key=lambda s: s.map(order) if s.name != 'QTY' else -s)
    for _, r in g.iterrows():
        sh.line(row, [r['RETURN/EXCHANGE REASON'], r['_raw_reason'], int(r['QTY'])])
        row += 1

    ws.column_dimensions['A'].width = 44
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 14
    ws.freeze_panes = 'A3'


def write_sku_pivot(ws, df, design_map, launch_map):
    ret = df[df['ORDER STATUS'] == 'Return']
    g = df.groupby('MY SKU').agg(sales=('QTY', 'sum'), returns=('_ret', 'sum'))
    reasons = ret.pivot_table(index='MY SKU', columns='RETURN/EXCHANGE REASON', values='QTY',
                              aggfunc='sum', fill_value=0)
    table = g.join(reasons).fillna(0).sort_index()

    headers = ['MY SKU', 'DESIGN NO', 'PORTAL LAUNCH DATE', 'SALE UNIT', 'RETURN TOTAL', 'RETURN %'] + list(reasons.columns)
    ws.append([])  # row 1 is the Grand Total, filled in below once the data range is known
    ws.append(headers)
    for sku, row in table.iterrows():
        rate = round(row['returns'] / row['sales'] * 100, 2) if row['sales'] else 0
        launch = launch_map.get(sku)
        launch = None if pd.isna(launch) else launch.to_pydatetime()
        ws.append([sku, design_map.get(sku), launch, int(row['sales']), int(row['returns']), rate]
                  + [int(row[c]) for c in reasons.columns])
        ws.cell(row=ws.max_row, column=3).number_format = 'DD-MM-YYYY'
        ws.cell(row=ws.max_row, column=6).number_format = '0.00'

    # Grand Total on top with SUBTOTAL(9, ...) so it follows the filter on the header row
    first, last = 3, max(ws.max_row, 3)
    ws.cell(row=1, column=1, value='Grand Total')
    for col in [4, 5] + list(range(7, len(headers) + 1)):
        letter = get_column_letter(col)
        ws.cell(row=1, column=col, value=f'=SUBTOTAL(9,{letter}{first}:{letter}{last})')
    ws.cell(row=1, column=6, value='=IFERROR(ROUND(E1/D1*100,2),0)').number_format = '0.00'
    for cell in ws[1] + ws[2]:
        cell.font = Font(bold=True)

    ws.column_dimensions['A'].width = 13
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['F'].width = 9.14
    ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{last}'
    ws.freeze_panes = 'B3'


def sheet_title(name, used):
    name = ''.join('_' if ch in '[]:*?/\\' else ch for ch in str(name))[:31] or 'Sheet'
    base, n = name, 2
    while name.lower() in used:
        name = f'{base[:28]}_{n}'
        n += 1
    used.add(name.lower())
    return name


def write_report(df, portal, month_tag, out_root, design_map, launch_map):
    df = df.copy()
    df['_ret'] = df['QTY'].where(df['ORDER STATUS'] == 'Return', 0)
    df['_month'] = pd.to_datetime(df['MONTH'], format='%m-%Y')
    for col in ['BRAND', 'CATEGORY', 'MY SKU']:
        df[col] = df[col].fillna('UNMAPPED')
    is_ret = df['ORDER STATUS'] == 'Return'
    df['_raw_reason'] = df['RETURN/EXCHANGE REASON']
    df.loc[is_ret, 'RETURN/EXCHANGE REASON'] = merge_reason(df.loc[is_ret, 'RETURN/EXCHANGE REASON'])
    df['_raw_state'] = df['STATE']
    df['STATE'] = clean_state(df['STATE'])
    df['ZONE'] = df['STATE'].map(STATE_ZONE)

    wb = Workbook()
    write_summary(wb.active, df)
    wb.active.title = 'Summary'
    write_monthly_summary(wb.create_sheet('Monthly Summary'), df)
    write_category_performance(wb.create_sheet('Category Performance'), df)
    used = {'summary', 'monthly summary', 'category performance', 'all brands', 'portal wise'}
    if df['PORTAL'].nunique() > 1:
        write_portal_wise(wb.create_sheet('Portal Wise'), df)
    if df['ZONE'].notna().any():
        write_zone_wise(wb.create_sheet('Zone Wise'), df)
        write_state_wise(wb.create_sheet('State Wise'), df)
        used |= {'zone wise', 'state wise', 'zone mapping'}
    else:
        print(f'WARNING: {portal} report has no customer state, Zone Wise / State Wise sheets skipped')
    write_sku_pivot(wb.create_sheet('All Brands'), df, design_map, launch_map)
    for brand, part in df.groupby('BRAND'):
        write_sku_pivot(wb.create_sheet(sheet_title(brand, used)), part, design_map, launch_map)
    if df['ZONE'].notna().any():
        write_zone_mapping(wb.create_sheet('Zone Mapping'), df)
    write_reason_mapping(wb.create_sheet(sheet_title('Reason Mapping', used)), df)

    folder = os.path.join(out_root, 'Output', month_tag)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f'{portal} SALE RETURN ANALYSIS {month_tag}.xlsx')
    wb.save(path)
    return path


# ----------------------------------------------------------------------------- main

def write_template(df, portal, month_tag, out_root):
    folder = os.path.join(out_root, 'Template')
    os.makedirs(folder, exist_ok=True)
    final = finish(df)
    final.to_csv(os.path.join(folder, f'{portal} {month_tag} TEMPLATE.csv'), index=False)
    return final


def check(df, portal):
    for col in ['MY SKU', 'CATEGORY', 'BRAND', 'MONTH', 'SALE PRICE', 'COST']:
        n = df[col].isna().sum()
        if n:
            sample = df.loc[df[col].isna(), 'SELLER SKU CODE'].unique()[:10].tolist()
            print(f'WARNING: {portal} has {n} rows with empty {col}, e.g. {sample}')


def write_combined(built, out_root, design_map, launch, tag=None):
    """MYNTRA COMBINED (Myntra MAIN + DBCC + FMCC) and ALL PORTAL reports from (portal, df, stats)
    parts; tag names the file (default: the latest month, MON-YY)."""
    results = []
    for name, parts in [('MYNTRA COMBINED', [b for b in built if b[0].startswith('MYNTRA')]), ('ALL PORTAL', built)]:
        if len(parts) < 2:
            continue
        df = pd.concat([part for _, part, _ in parts], ignore_index=True)
        # Same brand is spelt differently per portal (DressBerry / Dressberry, FASHFUN CLOTHING / FashFun):
        # use the most common spelling
        key = df['BRAND'].str.strip().str.lower().replace(BRAND_ALIASES)
        df['BRAND'] = key.map(df.groupby(key)['BRAND'].agg(lambda s: s.mode()[0]))
        month_tag = tag or pd.to_datetime(max(df['MONTH'].dropna()), format='%m-%Y').strftime('%b-%y').upper()
        launch_map = launch_dates(launch, [p for p, _, _ in parts])
        report = write_report(df, name, month_tag, out_root, design_map, launch_map)
        sold = int(df['QTY'].sum())
        returned = int(df.loc[df['ORDER STATUS'] == 'Return', 'QTY'].sum())
        rate = returned / sold * 100 if sold else 0
        print(f'{name} ({month_tag})  portals={", ".join(p for p, _, _ in parts)}')
        print(f'    sale qty={sold:<6} return qty={returned:<6} return %={rate:.2f}')
        print(f'    report: {report}')
        results.append({'portal': name, 'month': month_tag, 'sold': sold, 'returned': returned, 'rate': rate,
                        'report': report, 'template': None, 'stats': {}})
    return results


# Columns of a portal's monthly data kept for the yearly report: the template columns plus STATE
DATA_COLUMNS = COLUMNS + ['STATE']


def read_data(path):
    """A monthly data CSV (DATA_COLUMNS) back as the frame the builders made."""
    df = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[''])
    df['QTY'] = pd.to_numeric(df['QTY']).astype(int)
    for col in ['SALE PRICE', 'COST']:
        df[col] = pd.to_numeric(df[col], errors='coerce')
    return df


def build_period(frames, tag, supporting_dir, out_root):
    """Reports over several months (e.g. a year) from the portals' monthly data:
    frames = {portal: df}. One report per portal plus the combined ones, all named with tag."""
    _, _, _, design_map, launch = load_master(supporting_dir)
    built = [(portal, df, {}) for portal, df in frames.items()]
    results = []
    for portal, df, _ in built:
        report = write_report(df, portal, tag, out_root, design_map, launch_dates(launch, portal))
        sold = int(df['QTY'].sum())
        returned = int(df.loc[df['ORDER STATUS'] == 'Return', 'QTY'].sum())
        rate = returned / sold * 100 if sold else 0
        print(f'{portal} ({tag})  months={sorted(df["MONTH"].dropna().unique())}')
        print(f'    sale qty={sold:<6} return qty={returned:<6} return %={rate:.2f}')
        results.append({'portal': portal, 'month': tag, 'sold': sold, 'returned': returned, 'rate': rate,
                        'report': report, 'template': None, 'stats': {}})
    return results + write_combined(built, out_root, design_map, launch, tag)


def build_all(input_dir, supporting_dir, out_root):
    """Build every portal found in input_dir plus the combined reports. Returns one dict per report:
    portal, month, sold, returned, rate, report (path), template (path or None), stats;
    and a list of (portal, reason) for portals skipped because their input is missing."""
    sku_map, portal_map, conflicts, design_map, launch = load_master(supporting_dir)
    cost_map = load_cost(supporting_dir)

    built, skipped, results = [], [], []
    for acc in MYNTRA_ACCOUNTS:
        if not os.path.isdir(os.path.join(input_dir, 'Myntra', acc)):
            print(f'Skipping MYNTRA {acc}: folder not found')
            skipped.append((f'MYNTRA {acc}', 'no input folder'))
            continue
        try:
            df, stats = build_myntra(input_dir, acc, sku_map, conflicts, cost_map)
        except MissingInput as exc:
            print(f'Skipping MYNTRA {acc}: {exc}')
            skipped.append((f'MYNTRA {acc}', str(exc)))
            continue
        built.append((f'MYNTRA {acc}', df, stats))
    if os.path.isdir(os.path.join(input_dir, 'Ajio')):
        try:
            df, stats = build_ajio(input_dir, sku_map, portal_map, cost_map)
            built.append(('AJIO', df, stats))
        except (MissingInput, ValueError) as exc:   # ValueError: no RTV files to concat
            print(f'Skipping AJIO: {exc}')
            skipped.append(('AJIO', str(exc)))
    if os.path.isdir(os.path.join(input_dir, 'Flipkart')):
        # Flipkart titles are vague, so reuse the category other portals give the same MY SKU
        known = pd.concat([df[['MY SKU', 'CATEGORY']] for _, df, _ in built]).dropna() if built else \
            pd.DataFrame(columns=['MY SKU', 'CATEGORY'])
        category_by_sku = known.drop_duplicates('MY SKU').set_index('MY SKU')['CATEGORY']
        try:
            df, stats = build_flipkart(input_dir, sku_map, portal_map, cost_map, category_by_sku)
            built.append(('FLIPKART', df, stats))
        except MissingInput as exc:
            print(f'Skipping FLIPKART: {exc}')
            skipped.append(('FLIPKART', str(exc)))

    print()
    for portal, df, stats in built:
        months = df['MONTH'].dropna().unique()
        if len(months) != 1:
            print(f'WARNING: {portal} has orders from several months: {sorted(months)}')
        month_tag = pd.to_datetime(sorted(months)[-1], format='%m-%Y').strftime('%b-%y').upper()
        check(df, portal)
        final = write_template(df, portal, month_tag, out_root)
        launch_map = launch_dates(launch, portal)
        skus = final['MY SKU'].dropna().unique()
        for label, mp in [('DESIGN NO', design_map), ('PORTAL LAUNCH DATE', launch_map)]:
            missing = [s for s in skus if pd.isna(mp.get(s))]
            if missing:
                print(f'WARNING: {portal} has {len(missing)} MY SKUs without {label} in master, e.g. {missing[:10]}')
        raw_reason = df.loc[df['ORDER STATUS'] == 'Return', 'RETURN/EXCHANGE REASON'].dropna()
        unmerged = sorted(set(raw_reason[merge_reason(raw_reason) == raw_reason]) - {g for g, _ in REASON_GROUPS})
        if unmerged:
            print(f'WARNING: {portal} return reasons not in any REASON_GROUPS rule (kept as is): {unmerged}')
        raw_state = df['STATE'].dropna()
        unknown = sorted(raw_state[clean_state(raw_state).isna()].unique())
        if unknown:
            print(f'WARNING: {portal} has states not in STATE_ZONE / STATE_ALIASES: {unknown}')
        # The template drops STATE, so the report is built from df to keep zone / state
        report = write_report(df, portal, month_tag, out_root, design_map, launch_map)
        sold = int(final['QTY'].sum())
        returned = int(final.loc[final['ORDER STATUS'] == 'Return', 'QTY'].sum())
        rate = returned / sold * 100 if sold else 0
        print(f'{portal} ({month_tag})  ' + '  '.join(f'{k}={v}' for k, v in stats.items()))
        print(f'    rows={len(final):<6} sale qty={sold:<6} return qty={returned:<6} return %={rate:.2f}')
        print(f'    report: {report}')
        results.append({'portal': portal, 'month': month_tag, 'sold': sold, 'returned': returned, 'rate': rate,
                        'report': report, 'template': os.path.join(out_root, 'Template', f'{portal} {month_tag} TEMPLATE.csv'),
                        'stats': stats, 'df': df})

    # Combined reports: same sheets plus Portal Wise
    results += write_combined(built, out_root, design_map, launch)
    print('\nDone.')
    return results, skipped


def main():
    ap = argparse.ArgumentParser(description='Build Myntra, Ajio and Flipkart return-analysis templates and reports.')
    ap.add_argument('--input', default=os.path.join(BASE, 'Input'))
    ap.add_argument('--supporting', default=os.path.join(BASE, 'Supporting'))
    ap.add_argument('--out', default=BASE)
    args = ap.parse_args()
    build_all(args.input, args.supporting, args.out)


if __name__ == '__main__':
    main()
