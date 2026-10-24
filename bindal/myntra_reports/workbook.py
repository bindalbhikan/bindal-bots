"""MYNTRA DAILY RETURN DELIVERY REPORT.xlsx: the DATA sheet the bot adds a row to per account and day
(DATE, PORTAL, CUSTOMER RETURN, RTO, CANCELLED, TOTAL) and a DASHBOARD sheet of formulas over its
whole columns, so it is up to date whenever the file is opened. C3 picks the day (the latest by
default; type a date over it to see another day)."""
from __future__ import annotations

from datetime import date, datetime

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

DATA, DASH = "DATA", "DASHBOARD"
HEAD = ["DATE", "PORTAL", "CUSTOMER RETURN", "RTO", "CANCELLED", "TOTAL"]
KINDS = ["CUSTOMER RETURN", "RTO", "CANCELLED"]
# DATA columns of each figure
COL = {"CUSTOMER RETURN": "C", "RTO": "D", "CANCELLED": "E", "TOTAL": "F"}

FONT = "Arial"
NAVY, TEAL, ORANGE, RED, GREY = "1F3864", "2E75B6", "ED7D31", "C00000", "F2F2F2"
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
INT = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
DAY = "DD-MMM-YY"
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")


def _font(**kw) -> Font:
    return Font(name=FONT, **kw)


def _header(ws, row: int, col: int, labels: list[str]) -> None:
    for i, label in enumerate(labels):
        c = ws.cell(row, col + i, label)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def _section(ws, row: int, text: str) -> None:
    c = ws.cell(row, 2, text)
    c.font = _font(bold=True, size=12, color=NAVY)


def _sum(col: str, *conds: str) -> str:
    """SUMIFS over a whole DATA column, so rows the bot adds later count too."""
    return f"SUMIFS({DATA}!${col}:${col}," + ",".join(conds) + ")"


def _cell(ws, row: int, col: int, value, fmt: str = INT, bold: bool = False, fill=None):
    c = ws.cell(row, col, value)
    c.font = _font(bold=bold)
    c.number_format = fmt
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    if fill:
        c.fill = fill
    return c


def style_data(ws) -> None:
    """Header, widths, frozen header and filter of the DATA sheet; each data row's fonts and formats."""
    ws.title = DATA
    _header(ws, 1, 1, HEAD)
    for i, w in enumerate([13, 16, 18, 10, 14, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    last = last_row(ws)
    ws.auto_filter.ref = f"A1:F{max(last, 2)}"
    for r in range(2, last + 1):
        for i in range(1, 7):
            c = ws.cell(r, i)
            c.font = _font()
            c.border = BOX
            c.number_format = DAY if i == 1 else ("@" if i == 2 else INT)
            c.alignment = Alignment(horizontal="left" if i == 2 else "center")


def last_row(ws) -> int:
    """The last row with a value (formatting below the data makes max_row larger)."""
    r = ws.max_row
    while r > 1 and all(c.value in (None, "") for c in ws[r]):
        r -= 1
    return r


def put_rows(ws, rows: list[tuple[date, str, dict[str, int]]]) -> None:
    """rows (day, portal, {CUSTOMER RETURN, RTO, CANCELLED, TOTAL}) into DATA, replacing rows of the
    same day and portal, the whole sheet kept in date then portal order."""
    keys = {(d, p) for d, p, _ in rows}
    kept = []
    for r in range(2, last_row(ws) + 1):
        v = [ws.cell(r, i).value for i in range(1, 7)]
        if v[0] in (None, ""):
            continue
        d = v[0].date() if isinstance(v[0], datetime) else v[0]
        if (d, str(v[1] or "").strip()) not in keys:
            kept.append([d, *v[1:]])
    kept += [[d, p, *(s.get(k, 0) for k in HEAD[2:])] for d, p, s in rows]
    kept.sort(key=lambda v: (v[0] if isinstance(v[0], date) else date.min, str(v[1])))
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for i, v in enumerate(kept, 2):
        for j, x in enumerate(v, 1):
            ws.cell(i, j, x)
    style_data(ws)


def _axes(chart) -> None:
    """Show both axes with their labels (newer openpyxl hides them unless told)."""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.number_format = "#,##0"
    chart.y_axis.majorGridlines = None
    chart.legend.position = "t"
    chart.legend.overlay = False


def build_dashboard(wb, portals: list[str]) -> None:
    """The DASHBOARD sheet (first tab), replacing an older one."""
    if DASH in wb.sheetnames:
        del wb[DASH]
    ws = wb.create_sheet(DASH, 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["B"].width = 18

    ws.merge_cells("B1:I1")
    t = ws["B1"]
    t.value = "MYNTRA RETURN DELIVERY DASHBOARD"
    t.font = _font(bold=True, size=16, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws["B3"] = "Date"
    ws["B3"].font = _font(bold=True)
    d = ws["C3"]
    d.value = f"=MAX({DATA}!$A:$A)"
    d.number_format = DAY
    d.font = _font(bold=True, color="0000FF")
    d.fill = PatternFill("solid", fgColor="FFFF00")
    d.border = BOX
    d.alignment = Alignment(horizontal="center")
    ws["D3"] = "latest day by default; type a date here to see another day"
    ws["D3"].font = _font(italic=True, color="808080", size=9)
    ws["H3"] = "Last data"
    ws["H3"].font = _font(bold=True)
    ws["I3"] = f"=MAX({DATA}!$A:$A)"
    ws["I3"].number_format = DAY
    ws["I3"].font = _font()
    on_day = f"{DATA}!$A:$A,$C$3"

    # KPI tiles: value, share of the total, change from the day before
    tiles = [("TOTAL", NAVY, "B"), ("CUSTOMER RETURN", TEAL, "D"), ("RTO", ORANGE, "F"), ("CANCELLED", RED, "H")]
    for name, color, col in tiles:
        c0 = ws[f"{col}5"]
        ws.merge_cells(f"{col}5:{chr(ord(col) + 1)}5")
        ws.merge_cells(f"{col}6:{chr(ord(col) + 1)}6")
        ws.merge_cells(f"{col}7:{chr(ord(col) + 1)}7")
        ws.merge_cells(f"{col}8:{chr(ord(col) + 1)}8")
        c0.value = name
        c0.font = _font(bold=True, color="FFFFFF")
        c0.fill = PatternFill("solid", fgColor=color)
        c0.alignment = Alignment(horizontal="center")
        v = ws[f"{col}6"]
        v.value = "=" + _sum(COL[name], on_day)
        v.font = _font(bold=True, size=20, color=color)
        v.number_format = INT
        v.alignment = Alignment(horizontal="center")
        s = ws[f"{col}7"]
        s.value = "=IFERROR({0}6/$B$6,0)".format(col) if name != "TOTAL" else "of all returns delivered"
        s.number_format = PCT
        s.font = _font(color="808080", size=9)
        s.alignment = Alignment(horizontal="center")
        p = ws[f"{col}8"]
        p.value = (f'="vs day before: "&TEXT({col}6-' + _sum(COL[name], f"{DATA}!$A:$A,$C$3-1") + ',"+#,##0;-#,##0;0")')
        p.font = _font(color="808080", size=9)
        p.alignment = Alignment(horizontal="center")
        for r in range(5, 9):
            for cc in (col, chr(ord(col) + 1)):
                ws[f"{cc}{r}"].border = BOX
                if r > 5:
                    ws[f"{cc}{r}"].fill = PatternFill("solid", fgColor=GREY)

    head = ["PORTAL", *KINDS, "TOTAL", "RTO %", "CANCEL %"]

    def portal_table(top: int, title: str, cond: str) -> None:
        _section(ws, top, title)
        _header(ws, top + 1, 2, head)
        first = top + 2
        for i, portal in enumerate(portals):
            r = first + i
            _cell(ws, r, 2, portal, "@").alignment = Alignment(horizontal="left")
            for j, k in enumerate([*KINDS, "TOTAL"]):
                _cell(ws, r, 3 + j, "=" + _sum(COL[k], cond, f'{DATA}!$B:$B,$B{r}'))
            _cell(ws, r, 7, f"=IFERROR(D{r}/F{r},0)", PCT)
            _cell(ws, r, 8, f"=IFERROR(E{r}/F{r},0)", PCT)
        r = first + len(portals)
        _cell(ws, r, 2, "TOTAL", "@", True, TOTAL_FILL).alignment = Alignment(horizontal="left")
        for j in range(4):
            col = get_column_letter(3 + j)
            _cell(ws, r, 3 + j, f"=SUM({col}{first}:{col}{r - 1})", INT, True, TOTAL_FILL)
        _cell(ws, r, 7, f"=IFERROR(D{r}/F{r},0)", PCT, True, TOTAL_FILL)
        _cell(ws, r, 8, f"=IFERROR(E{r}/F{r},0)", PCT, True, TOTAL_FILL)

    n = len(portals)
    portal_table(10, "BY PORTAL · selected day", on_day)
    month_top = 10 + n + 4
    portal_table(month_top, "BY PORTAL · month to date",
                 f'{DATA}!$A:$A,">="&DATE(YEAR($C$3),MONTH($C$3),1),{DATA}!$A:$A,"<="&$C$3')

    # last 14 days, oldest first
    trend = month_top + n + 4
    _section(ws, trend, "LAST 14 DAYS")
    _header(ws, trend + 1, 2, ["DATE", *KINDS, "TOTAL", *portals])
    t0 = trend + 2
    for i in range(14):
        r = t0 + i
        _cell(ws, r, 2, f"=$C$3-{13 - i}", DAY)
        for j, k in enumerate([*KINDS, "TOTAL"]):
            _cell(ws, r, 3 + j, "=" + _sum(COL[k], f"{DATA}!$A:$A,$B{r}"))
        for j, portal in enumerate(portals):
            _cell(ws, r, 7 + j, "=" + _sum("F", f"{DATA}!$A:$A,$B{r}", f'{DATA}!$B:$B,"{portal}"'))

    # last 6 months, oldest first
    months = t0 + 14 + 2
    _section(ws, months, "MONTHLY")
    _header(ws, months + 1, 2, ["MONTH", *KINDS, "TOTAL", "RTO %", "CANCEL %", "DAYS"])
    m0 = months + 2
    for i in range(6):
        r = m0 + i
        _cell(ws, r, 2, f"=DATE(YEAR($C$3),MONTH($C$3)-{5 - i},1)", "MMM-YY")
        span = f'{DATA}!$A:$A,">="&$B{r},{DATA}!$A:$A,"<="&EOMONTH($B{r},0)'
        for j, k in enumerate([*KINDS, "TOTAL"]):
            _cell(ws, r, 3 + j, "=" + _sum(COL[k], span))
        _cell(ws, r, 7, f"=IFERROR(D{r}/F{r},0)", PCT)
        _cell(ws, r, 8, f"=IFERROR(E{r}/F{r},0)", PCT)
        # days with data: a day has one row per portal
        _cell(ws, r, 9, f'=COUNTIFS({span})/{n}', "0")

    ws.freeze_panes = "A4"

    # charts, right of the tables
    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
    bar.title, bar.height, bar.width = "Return delivery · last 14 days", 8, 18
    bar.add_data(Reference(ws, min_col=3, max_col=5, min_row=t0 - 1, max_row=t0 + 13), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=t0, max_row=t0 + 13))
    for s, color in zip(bar.series, (TEAL, ORANGE, RED)):
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    bar.x_axis.number_format = "DD-MMM"
    _axes(bar)
    ws.add_chart(bar, "K3")

    by_portal = BarChart()
    by_portal.type, by_portal.grouping, by_portal.overlap = "bar", "stacked", 100
    by_portal.title, by_portal.height, by_portal.width = "By portal · selected day", 7, 18
    by_portal.add_data(Reference(ws, min_col=3, max_col=5, min_row=11, max_row=11 + n), titles_from_data=True)
    by_portal.set_categories(Reference(ws, min_col=2, min_row=12, max_row=11 + n))
    for s, color in zip(by_portal.series, (TEAL, ORANGE, RED)):
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    _axes(by_portal)
    ws.add_chart(by_portal, "K20")

    line = LineChart()
    line.title, line.height, line.width = "Monthly return delivery", 7, 18
    line.add_data(Reference(ws, min_col=3, max_col=6, min_row=m0 - 1, max_row=m0 + 5), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=2, min_row=m0, max_row=m0 + 5))
    for s, color in zip(line.series, (TEAL, ORANGE, RED, NAVY)):
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 22000
        s.smooth = False
    line.x_axis.number_format = "MMM-YY"
    _axes(line)
    ws.add_chart(line, "K35")

    wb.active = 0
    # Excel recalculates every formula when the file is opened (the bot writes no cached values)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)


def update(wb, rows: list[tuple[date, str, dict[str, int]]], portals: list[str]) -> None:
    """The rows into DATA; the DASHBOARD is built when the workbook has none yet."""
    ws = wb[DATA] if DATA in wb.sheetnames else wb.worksheets[0]
    put_rows(ws, rows)
    if DASH not in wb.sheetnames:
        build_dashboard(wb, portals)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
