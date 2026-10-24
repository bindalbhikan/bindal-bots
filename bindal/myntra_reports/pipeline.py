"""Myntra daily reports (10 AM): yesterday's Seller_Orders_Report, Seller_Returns_Report and
Return_Delivery_Report, and today's Seller_Listings_Report, of each account (MAIN / DBCC / FMCC).

The four are different reports, so they are asked for together (the 15-minute cooldown is per
report); a job already on the portal for the same dates (or a listings snapshot of today) is reused.
The orders report is asked for the last 30 days (the same job the 7 AM SPF run asks for): the
Orders file keeps only yesterday's lines, the rest tells the return deliveries of type
RTO_or_Cancel apart by order status (RTO = RTO, F = cancelled; an order older than 30 days counts as RTO).
The CSVs go to Drive AGENT WORK/MYNTRA DAILY REPORTS/<yesterday>/ ("MAIN Orders 2026-10-23.csv"),
each account's return deliveries become a row of AGENT WORK/MYNTRA DAILY RETURN DELIVERY REPORT.xlsx
(DATE, PORTAL, CUSTOMER RETURN, RTO, CANCELLED, TOTAL; a day run again replaces its rows), and
Telegram gets one message: the total return delivery, per account the rows of each report and the folder link.
"""
from __future__ import annotations

import argparse
import io
import logging
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd

from ..core import telegram
from ..core.app import run_bot
from ..core.config import csv_list, env, num, output_dir

log = logging.getLogger("bindal.myntra_reports")

ACCOUNTS = ["MAIN", "DBCC", "FMCC"]
AGENT_WORK = "1uO-aOha_QKkI3f19EA7sweYPgkjMukZI"   # Drive BINDAL EXPORTS/AGENT WORK
FOLDER = "MYNTRA DAILY REPORTS"
WORKBOOK = "MYNTRA DAILY RETURN DELIVERY REPORT.xlsx"
PORTALS = {"MAIN": "MYNTRA MAIN", "DBCC": "MYNTRA DBCC", "FMCC": "MYNTRA FMCC"}
DELIVERY_COLUMNS = ["CUSTOMER RETURN", "RTO", "CANCELLED", "TOTAL"]
# label -> report id
REPORTS = {"Orders": 4, "Returns": 12, "Return Delivery": 84, "Listings": 14}
SNAPSHOT = {"Listings"}
ORDER_DAYS = 30


def _job(portal, label: str, rid: int, day: date) -> dict:
    """A job of the report for yesterday (Orders: the 30 days up to yesterday; Listings: a snapshot
    asked for today), reused when the portal already has one."""
    from ..portals import myntra as mp
    if label == "Orders":
        return portal.ensure(rid, day - timedelta(days=ORDER_DAYS - 1), day)
    if label not in SNAPSHOT:
        return portal.ensure(rid, day, day)
    today = date.today()
    for j in portal.jobs(100):
        status = str(j.get("status") or "").upper()
        created = mp._ts(j.get("createdOn"))
        if j.get("reportId") == rid and created == today and not any(f in status for f in mp.FAILED):
            log.info("Report %s of today already on the portal (job %s, %s)", rid, j.get("id"), status)
            return j
    return portal.schedule(None, None, rid)


def _col(df: pd.DataFrame, name: str) -> pd.Series:
    """The column, whatever its case in the CSV; empty when the report has no such column."""
    for c in df.columns:
        if str(c).strip().lower() == name:
            return df[c].astype(str).str.strip()
    return pd.Series(dtype=str, index=df.index)


def summary(label: str, df: pd.DataFrame) -> str:
    """One short line about a report's rows."""
    n = len(df)
    if label == "Orders":
        amount = pd.to_numeric(_col(df, "final amount"), errors="coerce").sum()
        cancelled = int((_col(df, "order status") == "F").sum())
        return f"{n} lines, ₹{amount:,.0f}, {cancelled} cancelled"
    if label in ("Returns", "Return Delivery"):
        kinds = _col(df, "type").value_counts()
        return f"{n} rows" + (" (" + ", ".join(f"{k} {v}" for k, v in kinds.items()) + ")" if len(kinds) else "")
    if label == "Listings":
        active = int((_col(df, "listing status") == "A").sum())
        return f"{n} SKUs, {active} active"
    return f"{n} rows"


def delivery_split(delivery: pd.DataFrame, orders: pd.DataFrame | None) -> dict[str, int]:
    """CUSTOMER RETURN / RTO / CANCELLED / TOTAL of a Return_Delivery_Report. RTO_or_Cancel rows are
    cancelled when the order's status in the orders report is F, else RTO."""
    kind = _col(delivery, "type")
    status = pd.Series(dtype=str)
    if orders is not None and len(orders):
        status = pd.Series(_col(orders, "order status").values,
                           index=_col(orders, "order release id").str.replace(r"\.0$", "", regex=True).values
                           ).groupby(level=0).first()
    rto_or_cancel = kind == "RTO_or_Cancel"
    cancelled = int((rto_or_cancel & (_col(delivery, "order_id").str.replace(r"\.0$", "", regex=True)
                                      .map(status) == "F")).sum())
    return {"CUSTOMER RETURN": int((kind == "Return").sum()), "RTO": int(rto_or_cancel.sum()) - cancelled,
            "CANCELLED": cancelled, "TOTAL": len(delivery)}


def fetch_account(account: str, day: date, dest: Path, wait_minutes: int) -> tuple[dict, dict | None, list[str]]:
    """The account's four reports into dest: ({label: (file, summary)}, its return delivery split
    or None, notes for reports not got)."""
    from ..portals import myntra as mp
    from ..return_analysis.sources import myntra_portal
    portal = myntra_portal(account)
    jobs, notes, got, frames = {}, [], {}, {}
    for label, rid in REPORTS.items():
        try:
            jobs[label] = _job(portal, label, rid, day)
        except Exception as exc:  # noqa: BLE001 - the other reports still come
            if isinstance(exc, mp.SessionExpired):
                raise
            log.warning("%s %s not asked for: %s", account, label, exc)
            notes.append(f"{account} {label}: {str(exc)[:100]}")
    for label, job in jobs.items():
        try:
            _, data = portal.collect(job, wait_minutes)
            df = pd.read_csv(io.BytesIO(data), low_memory=False)
            path = dest / f"{account} {label} {date.today() if label in SNAPSHOT else day}.csv"
            if label == "Orders":
                frames["orders_all"] = df
                df = df[_col(df, "created on").str.startswith(f"{day}")]
                df.to_csv(path, index=False)
            else:
                path.write_bytes(data)
            frames[label] = df
            got[label] = (path, summary(label, df))
        except Exception as exc:  # noqa: BLE001
            if isinstance(exc, mp.SessionExpired):
                raise
            log.warning("%s %s not downloaded: %s", account, label, exc)
            notes.append(f"{account} {label}: {str(exc)[:100]}")
    mp.save_to_worker(env, portal.cookies(), account)
    split = None
    if "Return Delivery" in frames:
        split = delivery_split(frames["Return Delivery"], frames.get("orders_all"))
        if "orders_all" not in frames and split["RTO"]:
            notes.append(f"{account}: no orders report, RTO and cancelled not told apart (all counted as RTO)")
    return got, split, notes


def update_workbook(root, day: date, splits: dict[str, dict], work: Path) -> str:
    """The day's rows (one per account) into AGENT WORK/MYNTRA DAILY RETURN DELIVERY REPORT.xlsx,
    replacing rows of the same date and portal. Returns the workbook's link."""
    import openpyxl
    from ..core.drive import DriveFile
    found = root.children().get(WORKBOOK)
    if not found:
        raise RuntimeError(f"{WORKBOOK} not found in AGENT WORK on Drive")
    drive = DriveFile(found["id"], root.http)
    path = drive.download(work / WORKBOOK, need_edit=True)
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    head = [str(c.value or "").strip().upper() for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    portals = {PORTALS.get(a, a) for a in splits}

    def cell_date(v):
        return v.date() if isinstance(v, datetime) else v
    for r in range(ws.max_row, 1, -1):
        v = [c.value for c in ws[r]]
        if cell_date(v[col["DATE"]]) == day and str(v[col["PORTAL"]] or "").strip() in portals:
            ws.delete_rows(r)
    # the first empty row (a sheet with formatting below the data reports a larger max_row)
    row = ws.max_row + 1
    while row > 2 and all(c.value in (None, "") for c in ws[row - 1]):
        row -= 1
    for a, split in splits.items():
        values = {"DATE": day, "PORTAL": PORTALS.get(a, a), **split}
        for h, i in col.items():
            if h in values:
                ws.cell(row, i + 1, values[h])
        ws.cell(row, col["DATE"] + 1).number_format = "DD-MM-YYYY"
        row += 1
    wb.save(path)
    drive.upload(path)
    return drive.url


def run(args) -> int:
    import tempfile
    day = date.fromisoformat(args.date) if args.date else date.today() - timedelta(days=1)
    accounts = csv_list("MYNTRA_REPORT_ACCOUNTS") or ACCOUNTS
    dest = output_dir("myntra-reports") / f"{day}"
    dest.mkdir(parents=True, exist_ok=True)
    wait = num("MYNTRA_REPORT_WAIT_MINUTES", 40)
    results: dict[str, dict] = {}
    splits: dict[str, dict] = {}
    notes: list[str] = []
    with ThreadPoolExecutor(len(accounts)) as pool:
        futures = {a: pool.submit(fetch_account, a, day, dest, wait) for a in accounts}
        for a, f in futures.items():
            try:
                results[a], split, more = f.result()
                if split:
                    splits[a] = split
                notes += more
            except Exception as exc:  # noqa: BLE001 - the other accounts still go out
                log.exception("%s reports failed", a)
                results[a] = {}
                notes.append(f"{a}: {str(exc)[:150]}")
    files = [p for got in results.values() for p, _ in got.values()]
    if not files:
        raise RuntimeError("no Myntra report downloaded: " + "; ".join(notes))

    link = book = ""
    if args.no_upload:
        notes.append(f"not uploaded to Drive (test run): {dest}")
    else:
        from ..core.drive import DriveFolder
        root = DriveFolder(env("AGENT_WORK_FOLDER_ID", AGENT_WORK))
        folder = root.subfolder(FOLDER, create=True).subfolder(f"{day}", create=True)
        for p in files:
            try:
                folder.put(p, "text/csv")
            except Exception as exc:  # noqa: BLE001 - the other files are still saved
                log.exception("%s not uploaded", p.name)
                notes.append(f"{p.name} not saved to Drive: {str(exc)[:100]}")
        link = f"https://drive.google.com/drive/folders/{folder.folder_id}"
        if splits:
            try:
                book = update_workbook(root, day, splits, Path(tempfile.mkdtemp(prefix="myntra_rd_")))
            except Exception as exc:  # noqa: BLE001 - the message still goes out
                log.exception("%s not updated", WORKBOOK)
                notes.append(f"{WORKBOOK} not updated: {str(exc)[:150]}")

    out = [f"📊 <b>Myntra daily reports · {day:%d %b %Y}</b>"]
    if splits:
        total = {k: sum(s[k] for s in splits.values()) for k in DELIVERY_COLUMNS}
        out += ["", f"📦 <b>Total return delivery: {total['TOTAL']}</b>",
                f"Customer return {total['CUSTOMER RETURN']} · RTO {total['RTO']} · Cancelled {total['CANCELLED']}"]
    for a in accounts:
        out += ["", f"<b>{telegram.esc(a)}</b>"]
        for label, (_, s) in results[a].items():
            if label == "Return Delivery" and a in splits:
                sp = splits[a]
                s = f"{sp['TOTAL']} (customer return {sp['CUSTOMER RETURN']}, RTO {sp['RTO']}, cancelled {sp['CANCELLED']})"
            out.append(f"• {telegram.esc(label)}: {telegram.esc(s)}")
        if not results[a]:
            out.append("• no report")
    if link:
        out += ["", f'📁 <a href="{link}">AGENT WORK / {FOLDER} / {day}</a>']
    if book:
        out.append(f'📗 <a href="{book}">{WORKBOOK}</a>')
    if notes:
        out += [""] + ["⚠️ " + telegram.esc(n) for n in notes]
    msg = "\n".join(out)
    print(msg) if args.no_telegram else telegram.send(msg)
    return 1 if notes and not args.no_upload else 0


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="run.py myntra-reports",
                                description="Myntra orders / returns / return delivery / listings reports, daily")
    p.add_argument("--date", help="report day YYYY-MM-DD (default yesterday)")
    p.add_argument("--no-upload", action="store_true", help="keep the files in output/, not on Drive")
    p.add_argument("--no-telegram", action="store_true", help="print the summary instead of sending it")
    return run_bot("Myntra daily reports", p, run, argv)
