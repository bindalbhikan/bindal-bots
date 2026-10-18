"""AJIO ads / TOT gap, monthly: G4 amount per month against the TOT marketing charge AJIO took.

On the 3rd of every month at 9 AM (worker cron), and again on the 5th when the 3rd's run did
not save the report (--if-missing). The SRM statement of accounts (the AJIO vendor ledger)
from STATEMENT_FROM to yesterday is fetched, the hand-run tool's rules build the single-sheet
report (engine.py, unchanged), and it goes to Drive AJIO/Ajio Ads Gap Working/output/
AJIO_Simple_Monthly_G4_Charges.xlsx (replaced in place), with a Telegram table.
"""
from __future__ import annotations

import argparse
import importlib
import logging
import os
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

from ..core import dashboard, telegram
from ..core.app import run_bot
from ..core.config import env, output_dir
from ..core.drive import DriveFolder

log = logging.getLogger("bindal.ajio_ads")

REPORT_NAME = "AJIO_Simple_Monthly_G4_Charges.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def done_this_month(out: DriveFolder | None, today: date) -> bool:
    """The report on Drive was already saved this month on or after the 3rd."""
    f = (out.children().get(REPORT_NAME) if out else None)
    if not f:
        return False
    saved = datetime.fromisoformat(f["modifiedTime"].replace("Z", "+00:00")).astimezone().date()
    return saved.year == today.year and saved.month == today.month and saved.day >= 3


def fetch_statement(work: Path, start: date, end: date) -> int:
    """The SRM statement as the StatementOfAccount*.xlsx the tool reads; rows written."""
    from ..ajio_spf.inputs import write_xlsx
    from ..portals.srm import PAYMENT_COLUMNS, Srm, payment_rows
    srm = Srm(env("SRM_USERNAME"), env("SRM_PASSWORD"))
    srm.login()
    rows = payment_rows(srm.statement(start, end))
    if not rows:
        raise RuntimeError(f"SRM statement {start}..{end} has no lines")
    write_xlsx(work / "input" / f"StatementOfAccount_{end:%d-%m-%Y}_bot.xlsx",
               {"Sheet1": (list(PAYMENT_COLUMNS), rows)})
    return len(rows)


def report_table(path: Path) -> list[list]:
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]


def summary_text(rows: list[list], start: date, end: date, n: int) -> str:
    """Month | G4 | marketing charge | % - fits a phone."""
    lines = [f"{'Month':<9}{'G4':>9}{'Mkt':>8}{'%':>5}"]
    for r in rows[1:]:
        month, g4, total = r[0], r[1], r[4] if len(r) > 4 else None
        pct = r[5] if len(r) > 5 else None
        if month is None or not isinstance(g4, (int, float)):
            continue
        mkt = f"{total:,.0f}" if isinstance(total, (int, float)) else "-"
        p = f"{pct * 100:.1f}" if isinstance(pct, (int, float)) and pct < 1 else (
            f"{pct:.1f}" if isinstance(pct, (int, float)) else "-")
        label = f"{month:%b-%y}" if hasattr(month, "strftime") else str(month)[:9]
        lines.append(f"{label:<9}{g4:>9,.0f}{mkt:>8}{p:>5}")
    return "\n".join([f"📣 <b>AJIO ads / TOT gap</b>",
                      f"<i>SRM statement {start:%d %b %Y} - {end:%d %b %Y}, {n} lines</i>", "",
                      "<pre>" + telegram.esc("\n".join(lines)) + "</pre>", "",
                      "📁 Drive: AJIO / Ajio Ads Gap Working / output"])


def run(args) -> int:
    if not env("AJIO_ADS_FOLDER_ID"):
        raise RuntimeError("AJIO_ADS_FOLDER_ID (Drive folder AJIO/Ajio Ads Gap Working) not set")
    today = date.today()
    root = DriveFolder(env("AJIO_ADS_FOLDER_ID"))
    out = root.subfolder("output", create=not args.no_upload)
    if args.if_missing and done_this_month(out, today):
        log.info("Report already saved this month: nothing to do")
        return 0
    start = datetime.strptime(env("STATEMENT_FROM", "2026-01-01"), "%Y-%m-%d").date()
    end = today - timedelta(days=1)
    work = Path(tempfile.mkdtemp(prefix="ajio_ads_"))
    n = fetch_statement(work, start, end)
    os.environ["AJIO_ADS_DIR"] = str(work)
    from . import engine
    engine = importlib.reload(engine)             # picks up AJIO_ADS_DIR
    engine.main()
    report = Path(engine.TARGET)
    (output_dir("ajio_ads") / REPORT_NAME).write_bytes(report.read_bytes())
    if not args.no_upload:
        out.put(report, XLSX)
    text = summary_text(report_table(report), start, end, n)
    if args.no_telegram:
        print(text)
    else:
        telegram.send(text)
        telegram.send_document(report, f"<b>{REPORT_NAME}</b>")
    dashboard.publish("ajio_ads", {"from": start, "to": end, "lines": n})
    return 0


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="run.py ajio-ads", description="AJIO ads / TOT gap, monthly")
    p.add_argument("--if-missing", action="store_true",
                   help="do nothing when this month's report is already on Drive (the 5th's retry)")
    p.add_argument("--no-upload", action="store_true", help="keep the report in output/ajio_ads only")
    p.add_argument("--no-telegram", action="store_true", help="print the summary instead of sending it")
    return run_bot("AJIO ads / TOT gap", p, run, argv, key="ajio_ads")
