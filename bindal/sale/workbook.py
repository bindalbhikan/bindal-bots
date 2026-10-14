"""SALE REPORT.xlsx: one sheet per month, one row per day, a block per channel.

Row 2 holds the block names (AZ/FX/FBA, MEESHO, AJIO B2C, FLIPKART, FBF, MYNTRA,
SJIT, NYKAA, SNAPDEAL, Total), row 3 the metrics (GROSS, RETURN, Cancel, NET),
column A the dates and a TOTAL row closes the month. RETURN is entered by hand
and never written. Layout detection and formula repair are ported from the
original main.py.
"""
from __future__ import annotations

import calendar
import copy
from collections import defaultdict
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BLOCK_ALIASES = {
    "AZ/FX/FBA": "AZ/FX/FBA", "AMAZON/FLEX/FBA": "AZ/FX/FBA", "AMAZON": "AZ/FX/FBA", "COCOBLU": "COCOBLU",
    "MEESHO": "MEESHO", "AJIO B2C": "AJIO B2C", "AJIO": "AJIO B2C", "FLIPKART": "FLIPKART",
    "FBF": "FBF", "MYNTRA": "MYNTRA", "SJIT": "SJIT", "NYKAA": "NYKAA", "SNAPDEAL": "SNAPDEAL",
    "TOTAL": "TOTAL",
}
FIRST_ROW = 4


def _day(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


class SaleWorkbook:
    def __init__(self, path):
        self.path = path
        self.wb = load_workbook(path)

    # ---------------------------------------------------------------- sheets
    def _month_sheets(self) -> list[tuple[tuple[int, int], object]]:
        out = []
        for ws in self.wb.worksheets:
            first = _day(ws["A1"].value)
            days = [_day(ws.cell(r, 1).value) for r in range(FIRST_ROW, min(ws.max_row, 40) + 1)]
            days = [d for d in days if d]
            if first:
                out.append(((first.year, first.month), ws))
            elif len(days) >= 20:
                out.append(((days[0].year, days[0].month), ws))
        return out

    def month_sheet(self, year: int, month: int):
        for key, ws in self._month_sheets():
            if key == (year, month):
                return ws
        return self._new_month_sheet(year, month)

    def _day_rows(self, ws) -> list[int]:
        """Every row a day can sit in: from the first date row down to the row before TOTAL
        (a 30-day sheet copied from a 31-day one keeps its 31st row, with no date)."""
        _, rows, total_row = self.layout(ws)
        if not rows:
            return []
        first = min(rows.values())
        last = total_row - 1 if total_row else max(rows.values())
        return list(range(first, last + 1))

    def _new_month_sheet(self, year: int, month: int):
        """Copy the latest month sheet that has a row for every day of this month,
        give it this month's dates and empty values."""
        days = calendar.monthrange(year, month)[1]
        sheets = sorted(self._month_sheets(), key=lambda x: x[0], reverse=True)
        if not sheets:
            raise RuntimeError("No month sheet to copy the layout from")
        tpl = next((ws for _, ws in sheets if len(self._day_rows(ws)) >= days), None)
        if tpl is None:
            raise RuntimeError(f"No month sheet has {days} day rows to copy for {month:02d}-{year}")
        layout, _, _ = self.layout(tpl)
        data_rows = self._day_rows(tpl)
        ws = self.wb.copy_worksheet(tpl)
        ws.title = f"{calendar.month_abbr[month].upper()}-{year % 100:02d}"
        ws["A1"].value = date(year, month, 1)
        value_cols = [c for block, metrics in layout.items() if block != "TOTAL"
                      for m, c in metrics.items() if m in ("GROSS", "RETURN", "CANCEL")]
        for i, r in enumerate(data_rows):
            ws.cell(r, 1).value = date(year, month, i + 1) if i < days else None
            for c in value_cols:
                ws.cell(r, c).value = None
        self.repair_formulas(ws)
        return ws

    # ---------------------------------------------------------------- layout
    @staticmethod
    def layout(ws):
        layout: dict[str, dict[str, int]] = defaultdict(dict)
        group = ""
        for col in range(1, ws.max_column + 1):
            header = str(ws.cell(2, col).value or "").strip()
            if header:
                group = BLOCK_ALIASES.get(header.upper(), header.upper())
            metric = str(ws.cell(3, col).value or "").strip().upper()
            if group and metric in {"GROSS", "RETURN", "CANCEL", "NET"}:
                layout[group][metric] = col
        rows: dict[date, int] = {}
        total_row = None
        for r in range(FIRST_ROW, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if _day(v):
                rows[_day(v)] = r
            elif str(v or "").strip().casefold() == "total":
                total_row = r
        if total_row is None and rows:  # some months have no TOTAL label: the row after the last day
            last = max(rows.values())
            total_row = last + 1 if ws.cell(last + 1, 2).value and str(ws.cell(last + 1, 2).value).startswith("=") else None
        return dict(layout), rows, total_row

    def repair_formulas(self, ws) -> None:
        layout, date_rows, total_row = self.layout(ws)
        blocks = {k: m for k, m in layout.items() if k != "TOTAL" and {"GROSS", "RETURN", "CANCEL", "NET"} <= set(m)}
        all_rows = sorted(set(date_rows.values()) | ({r for r in range(min(date_rows.values()), (total_row or 0))}
                                                    if date_rows and total_row else set()))
        L = get_column_letter
        for r in all_rows:
            for m in blocks.values():
                ws.cell(r, m["NET"]).value = f"={L(m['GROSS'])}{r}-{L(m['RETURN'])}{r}-{L(m['CANCEL'])}{r}"
        if total_row and all_rows:
            first, last = min(all_rows), max(all_rows)
            for m in blocks.values():
                for metric in ("GROSS", "RETURN", "CANCEL"):
                    col = L(m[metric])
                    ws.cell(total_row, m[metric]).value = f"=SUM({col}{first}:{col}{last})"
                ws.cell(total_row, m["NET"]).value = (f"={L(m['GROSS'])}{total_row}-{L(m['RETURN'])}{total_row}"
                                                      f"-{L(m['CANCEL'])}{total_row}")
        tot = layout.get("TOTAL", {})
        if {"GROSS", "RETURN", "CANCEL", "NET"} <= set(tot):
            for r in all_rows + ([total_row] if total_row else []):
                for metric in ("GROSS", "RETURN", "CANCEL", "NET"):
                    refs = ",".join(f"{L(m[metric])}{r}" for m in blocks.values())
                    ws.cell(r, tot[metric]).value = f"=SUM({refs})"

    # ---------------------------------------------------------------- writing
    def write(self, day: date, block: str, metric: str, value: int, only_if_empty: bool = False) -> bool:
        ws = self.month_sheet(day.year, day.month)
        layout, rows, _ = self.layout(ws)
        if block not in layout or metric not in layout[block]:
            raise RuntimeError(f'Sheet {ws.title} has no "{block}" {metric} column')
        if day not in rows:
            raise RuntimeError(f"Sheet {ws.title} has no row for {day}")
        cell = ws.cell(rows[day], layout[block][metric])
        if only_if_empty and cell.value not in (None, ""):
            return False
        cell.value = int(value)
        return True

    def has_block(self, day: date, block: str) -> bool:
        layout, _, _ = self.layout(self.month_sheet(day.year, day.month))
        return block in layout

    def move_blocks(self, ws, new_block: str, after: str, remove: str) -> None:
        """Rebuild a month sheet's columns: drop the `remove` block and add an empty `new_block`
        (GROSS / RETURN / Cancel / NET, styled like `after`) right after `after`. Values,
        styles, widths and merged headers move with their columns; formulas are rewritten."""
        layout, _, _ = self.layout(ws)
        if new_block in layout:
            return
        if after not in layout:
            raise RuntimeError(f"Sheet {ws.title} has no {after} block")
        ins = max(layout[after].values()) + 1                    # first column of the new block
        width = 4
        gone = sorted(layout[remove].values()) if remove in layout else []
        gone = list(range(gone[0], gone[-1] + 1)) if gone else []
        if gone and gone[0] < ins:
            raise RuntimeError(f"{remove} must sit right of {after}")
        max_col, max_row = ws.max_column, ws.max_row

        def new_col(c):                                          # old column -> new column (None = dropped)
            if c in gone:
                return None
            if c < ins:
                return c
            if gone and c > gone[-1]:
                return c + width - len(gone)
            return c + width

        cells = {(r, c): (ws.cell(r, c).value, copy.copy(ws.cell(r, c)._style))
                 for r in range(1, max_row + 1) for c in range(1, max_col + 1)}
        widths = {c: ws.column_dimensions[get_column_letter(c)].width for c in range(1, max_col + 1)}
        merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
        for m in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(m))
        last = max(c for c in (new_col(c) for c in range(1, max_col + 1)) if c)
        for r in range(1, max_row + 1):
            for c in range(1, max(max_col, last) + 1):
                ws.cell(r, c).value = None
        for (r, c), (value, style) in cells.items():
            nc = new_col(c)
            if nc:
                cell = ws.cell(r, nc)
                cell.value, cell._style = value, style
        src = sorted(layout[after].values())                     # the new block copies `after`'s look
        for k in range(width):
            for r in range(1, max_row + 1):
                ws.cell(r, ins + k)._style = copy.copy(cells[(r, src[min(k, len(src) - 1)])][1])
            ws.column_dimensions[get_column_letter(ins + k)].width = widths.get(src[min(k, len(src) - 1)])
        for c, w in widths.items():
            nc = new_col(c)
            if nc and w:
                ws.column_dimensions[get_column_letter(nc)].width = w
        for r0, c0, r1, c1 in merges:
            a, b = new_col(c0), new_col(c1)
            if r0 == 1 and r1 == 1 and c0 == 1:                  # the title row keeps its full width
                b = b or c1
            if a and b:
                ws.merge_cells(start_row=r0, start_column=a, end_row=r1, end_column=b)
        ws.cell(2, ins).value = new_block
        ws.merge_cells(start_row=2, start_column=ins, end_row=2, end_column=ins + width - 1)
        for k, metric in enumerate(("GROSS", "RETURN", "Cancel", "NET")):
            ws.cell(3, ins + k).value = metric
        self.repair_formulas(ws)

    def read(self, day: date, block: str, metric: str) -> int:
        ws = self.month_sheet(day.year, day.month)
        layout, rows, _ = self.layout(ws)
        v = ws.cell(rows[day], layout[block][metric]).value if day in rows and block in layout else None
        return int(v) if isinstance(v, (int, float)) else 0

    def month_gross(self, day: date) -> dict[str, int]:
        """Month-to-date GROSS per block from the values in the sheet."""
        ws = self.month_sheet(day.year, day.month)
        layout, rows, _ = self.layout(ws)
        out = {}
        for block, m in layout.items():
            if block == "TOTAL" or "GROSS" not in m:
                continue
            out[block] = sum(int(ws.cell(r, m["GROSS"]).value or 0) for d, r in rows.items()
                             if d <= day and isinstance(ws.cell(r, m["GROSS"]).value, (int, float)))
        return out

    def save(self, path=None) -> None:
        try:
            self.wb.calculation.fullCalcOnLoad = True
            self.wb.calculation.forceFullCalc = True
        except AttributeError:
            pass
        self.wb.save(path or self.path)
