"""MYNTRA ORDER ID WISE VMS.xlsx: one sheet per month, one row per SLA date.

Layout of every month sheet (kept exactly as the hand-made ones):
  row 1  month title (merged A1:M1)
  row 2  TLPL | MMRJ | BOAT | TOTAL  (each merged over 3 columns)
  row 3  DATE, then ORDER COUNT / VMS COUNT / % per group
  row 4+ one row per day of the month, then a TOTAL row with SUBTOTALs
A new month's sheet is copied from the latest month sheet, so fonts, borders,
widths and merges carry over. A hidden sheet keeps the order id -> Myntra AWB
pairs read from shipping labels, so a label is decoded only once.
"""
from __future__ import annotations

import calendar
import copy
import re
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CODES = ["TLPL", "MMRJ", "BOAT"]
GROUP_COL = {"TLPL": 2, "MMRJ": 5, "BOAT": 8}  # ORDER COUNT column; VMS COUNT is the next one
FIRST_ROW = 4
NCOLS = 13
AWB_SHEET = "_myntra_awb"
NO_VMS_SHEET = "NO VMS"
NO_VMS_HEADER = ["DATE", "CHANNEL CODE", "AWB", "REMARK"]
MONTHS = {m.upper(): i for i, m in enumerate(calendar.month_abbr) if m}
MONTHS.update({m.upper(): i for i, m in enumerate(calendar.month_name) if m})
MONTHS["SEPT"] = 9


def sheet_month(name: str) -> tuple[int, int] | None:
    m = re.fullmatch(r"([A-Za-z]+)[-_ ]?(\d{2}|\d{4})", name.strip())
    if not m or m.group(1).upper() not in MONTHS:
        return None
    year = int(m.group(2))
    return (year + 2000 if year < 100 else year), MONTHS[m.group(1).upper()]


class VmsWorkbook:
    def __init__(self, path):
        self.path = path
        self.wb = load_workbook(path)

    # ---------------------------------------------------------------- awb cache
    def awb_cache(self) -> dict[str, str]:
        if AWB_SHEET not in self.wb.sheetnames:
            return {}
        ws = self.wb[AWB_SHEET]
        return {str(r[0]): str(r[2]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[2]}

    def save_awb_cache(self, rows: list[tuple[str, str, str]]) -> None:
        if AWB_SHEET in self.wb.sheetnames:
            ws = self.wb[AWB_SHEET]
        else:
            ws = self.wb.create_sheet(AWB_SHEET)
            ws.append(["Order Id", "Courier AWB", "Myntra AWB"])
            ws.sheet_state = "hidden"
        known = {str(r[0]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        for row in rows:
            if row[0] not in known:
                ws.append(list(row))

    # ---------------------------------------------------------------- month sheets
    def month_sheet(self, year: int, month: int):
        for ws in self.wb.worksheets:
            if sheet_month(ws.title) == (year, month):
                return ws
        return self._new_month_sheet(year, month)

    def _template(self):
        months = [(sheet_month(ws.title), ws) for ws in self.wb.worksheets if sheet_month(ws.title)]
        if not months:
            raise RuntimeError("No month sheet (like OCT-26) to copy the layout from")
        return max(months, key=lambda x: x[0])[1]

    def _new_month_sheet(self, year: int, month: int):
        tpl = self._template()
        total_row = next(r for r in range(FIRST_ROW, tpl.max_row + 1)
                         if str(tpl.cell(r, 1).value or "").strip().upper() == "TOTAL")
        day_style = [copy.copy(tpl.cell(FIRST_ROW, c)._style) for c in range(1, NCOLS + 1)]
        total_style = [copy.copy(tpl.cell(total_row, c)._style) for c in range(1, NCOLS + 1)]
        height = tpl.row_dimensions[FIRST_ROW].height

        ws = self.wb.copy_worksheet(tpl)
        ws.title = f"{calendar.month_abbr[month].upper()}-{year % 100:02d}"
        # keep month sheets together, in date order, before CODE and hidden sheets
        self.wb.move_sheet(ws, offset=self.wb.index(tpl) + 1 - self.wb.index(ws))
        ws.delete_rows(FIRST_ROW, ws.max_row)
        ws["A1"].value = date(year, month, 1)

        days = calendar.monthrange(year, month)[1]
        for i in range(days):
            r = FIRST_ROW + i
            self._write_formulas(ws, r)
            ws.cell(r, 1).value = date(year, month, i + 1)
            for c in range(1, NCOLS + 1):
                ws.cell(r, c)._style = copy.copy(day_style[c - 1])
            ws.row_dimensions[r].height = height
        last = FIRST_ROW + days - 1
        t = last + 1
        ws.cell(t, 1).value = "TOTAL"
        for c in (2, 3, 5, 6, 8, 9, 11, 12):
            col = get_column_letter(c)
            ws.cell(t, c).value = f"=SUBTOTAL(9,{col}{FIRST_ROW}:{col}{last})"
        for v, o in ((4, 2), (7, 5), (10, 8), (13, 11)):
            a, b = get_column_letter(o + 1), get_column_letter(o)
            ws.cell(t, v).value = f"=IFERROR({a}{t}/{b}{t},0)"
        for c in range(1, NCOLS + 1):
            ws.cell(t, c)._style = copy.copy(total_style[c - 1])
        return ws

    @staticmethod
    def _write_formulas(ws, r: int) -> None:
        ws.cell(r, 4).value = f"=IFERROR(C{r}/B{r},0)"
        ws.cell(r, 7).value = f"=IFERROR(F{r}/E{r},0)"
        ws.cell(r, 10).value = f"=IFERROR(I{r}/H{r},0)"
        ws.cell(r, 11).value = f"=B{r}+E{r}+H{r}"
        ws.cell(r, 12).value = f"=C{r}+F{r}+I{r}"
        ws.cell(r, 13).value = f"=IFERROR(L{r}/K{r},0)"

    def write_day(self, day: date, counts: dict[str, tuple[int, int]]) -> None:
        """counts: code -> (order count, vms count) for one SLA date."""
        ws = self.month_sheet(day.year, day.month)
        row = None
        for r in range(FIRST_ROW, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if hasattr(v, "date") and v.date() == day or v == day:
                row = r
                break
        if row is None:
            raise RuntimeError(f"No row for {day} in sheet {ws.title}")
        for code in CODES:
            orders, vms = counts.get(code, (0, 0))
            ws.cell(row, GROUP_COL[code]).value = int(orders)
            ws.cell(row, GROUP_COL[code] + 1).value = int(vms)
        self._write_formulas(ws, row)

    # ---------------------------------------------------------------- NO VMS sheet
    def write_no_vms(self, dates: list[date], rows: list[tuple[date, str, str]]) -> None:
        """Replace the NO VMS rows of the given SLA dates with rows (date, code, awb)."""
        tpl = self._template()
        if NO_VMS_SHEET in self.wb.sheetnames:
            ws = self.wb[NO_VMS_SHEET]
            kept = [r for r in ws.iter_rows(min_row=2, values_only=True)
                    if r[0] is not None and _as_date(r[0]) not in set(dates)]
        else:
            ws = self.wb.create_sheet(NO_VMS_SHEET, index=self.wb.index(tpl) + 1)
            kept = []
        head_style = copy.copy(tpl.cell(3, 1)._style)
        body_style = [copy.copy(tpl.cell(FIRST_ROW, c)._style) for c in (1, 2, 2, 2)]
        all_rows = sorted([(_as_date(r[0]), r[1], r[2], r[3]) for r in kept]
                          + [(d, code, awb, "NO VMS") for d, code, awb in rows],
                          key=lambda r: (r[0], r[1] or "", r[2] or ""))
        ws.delete_rows(1, ws.max_row)
        for c, name in enumerate(NO_VMS_HEADER, start=1):
            cell = ws.cell(1, c, name)
            cell._style = copy.copy(head_style)
        for i, row in enumerate(all_rows, start=2):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(i, c, value)
                cell._style = copy.copy(body_style[c - 1])
        for col, width in zip("ABCD", (12, 16, 20, 12)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{max(len(all_rows) + 1, 2)}"

    def month_days(self, year: int, month: int) -> list[dict]:
        """[{date, codes: {code: [orders, vms]}}] for every filled day of a month sheet."""
        ws = next((w for w in self.wb.worksheets if sheet_month(w.title) == (year, month)), None)
        out = []
        for r in range(FIRST_ROW, (ws.max_row + 1) if ws is not None else FIRST_ROW):
            v = ws.cell(r, 1).value
            d = v.date() if hasattr(v, "date") else v
            if not isinstance(d, date):
                continue
            codes = {c: [_as_int(ws.cell(r, GROUP_COL[c]).value), _as_int(ws.cell(r, GROUP_COL[c] + 1).value)]
                     for c in CODES}
            if any(o for o, _ in codes.values()):
                out.append({"date": d, "codes": codes})
        return out

    def month_totals(self, upto: date) -> dict[str, tuple[int, int]]:
        """Month-to-date (order count, vms count) per code, summed from the month sheet."""
        ws = self.month_sheet(upto.year, upto.month)
        out = {c: [0, 0] for c in CODES}
        for r in range(FIRST_ROW, ws.max_row + 1):
            v = ws.cell(r, 1).value
            d = v.date() if hasattr(v, "date") else v
            if not isinstance(d, date) or d > upto:
                continue
            for c in CODES:
                out[c][0] += _as_int(ws.cell(r, GROUP_COL[c]).value)
                out[c][1] += _as_int(ws.cell(r, GROUP_COL[c] + 1).value)
        return {c: tuple(v) for c, v in out.items()}

    def save(self, path=None) -> None:
        self.wb.save(path or self.path)


def _as_date(v):
    return v.date() if hasattr(v, "date") else v


def _as_int(v) -> int:
    try:
        return int(v or 0)
    except (TypeError, ValueError):
        return 0
