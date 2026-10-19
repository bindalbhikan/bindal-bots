"""Myntra SPF master, every day at 9 AM (reports asked for at 7 AM with --request): returns of Myntra MAIN / DBCC / FMCC with claim aging, tickets,
payouts and ERP receipts, in Drive SPF/MYNTRA/MYNTRA SPF MASTER.xlsx.

The rules are those of the hand-run tool (engine.py, unchanged). The bot builds its input/
folder in a work folder:
  * Myntra partner portal, per account (saved sessions): Seller_Returns_Report and
    Return_Delivery_Report for the last RETURN_DAYS, Seller_Orders_Report for the last
    ORDER_DAYS (one month per job; the same report again only after the 15-minute cooldown)
  * ERP: GR GRN report and the dispatch (packing) report
  * Drive SPF/MYNTRA/input: what is still dropped by hand - each account's Ticket1 and payment
    exports, MANUAL RETURN and LEGACY MASTER
The previous master (from Drive) is the base, so a run only adds and updates rows. The new
master replaces the Drive file in place (Drive keeps the old versions).
"""
from __future__ import annotations

import argparse
import logging
import os
import shutil
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

from ..core import dashboard, telegram
from ..core.app import run_bot
from ..core.config import env, num, output_dir
from ..core.drive import DriveFile, DriveFolder

log = logging.getLogger("bindal.myntra_spf")

MASTER_NAME = "MYNTRA SPF MASTER.xlsx"
ACCOUNTS = ["MAIN", "DBCC", "FMCC"]
DELIVERY_REPORT_ID = 84            # Return_Delivery_Report
COOLDOWN_SECONDS = 15 * 60 + 45


# ------------------------------------------------------------------ Myntra portal

def _windows(today: date) -> dict[int, list[tuple[date, date]]]:
    """report id -> date ranges of at most 30 days, ending yesterday."""
    from ..portals import myntra as mp
    end = today - timedelta(days=1)

    def split(days: int) -> list[tuple[date, date]]:
        out, stop = [], end
        while len(out) * 30 < days:
            start = max(stop - timedelta(days=29), end - timedelta(days=days - 1))
            out.append((start, stop))
            stop = start - timedelta(days=1)
        return list(reversed(out))
    returns = split(num("MYNTRA_SPF_RETURN_DAYS", 30))
    return {mp.RETURNS_REPORT_ID: returns, DELIVERY_REPORT_ID: returns,
            mp.REPORT_ID: split(num("MYNTRA_SPF_ORDER_DAYS", 60))}


def fetch_account(account: str, today: date, dest: Path, notes: list[str],
                  old_orders: list[tuple[date, date]] | None = None) -> None:
    """The account's returns, delivery and order reports into dest. Different reports are asked
    for together; the same report again only after the cooldown. Jobs for the same dates are reused.

    old_orders: older order-report months for returns that still have no ORDER AMOUNT. They never
    hold up the run: a report already on the portal is just downloaded, at most
    MYNTRA_SPF_OLD_ORDERS_NEW (1) new one is asked for per run, after the daily ones, and is
    waited for only MYNTRA_SPF_OLD_WAIT_MINUTES (10); the rest, a refusal or a report not ready
    in time are left to the next runs. Daily reports refused are asked again after the cooldown,
    up to 3 times."""
    from ..return_analysis.sources import myntra_portal
    from ..portals import myntra as mp
    portal = myntra_portal(account)
    dest.mkdir(parents=True, exist_ok=True)
    left: list[str] = []
    old_jobs, ask_old = [], []
    listed = portal.jobs(100) if old_orders else []
    for a, b in old_orders or []:
        job = mp.exact_job(listed, mp.REPORT_ID, a, b)
        status = str((job or {}).get("status") or "").upper()
        if job and status == "COMPLETED" and job.get("successFile"):
            try:
                name, data = portal.download(job)
                (dest / name).write_bytes(data)
                continue
            except Exception as exc:  # noqa: BLE001 - expired link: asked for again
                log.warning("%s old order report %s..%s not downloaded: %s", account, a, b, exc)
        elif job and not any(f in status for f in mp.FAILED):
            old_jobs.append(job)             # still running: collected below
            continue
        ask_old.append((a, b))
    cap = num("MYNTRA_SPF_OLD_ORDERS_NEW", 1)
    left += [f"{a:%b %Y}" for a, _ in ask_old[cap:]]
    queues = [[(rid, a, b, False) for a, b in ranges] for rid, ranges in _windows(today).items()]
    for q in queues:
        if q and q[0][0] == mp.REPORT_ID:
            q.extend((mp.REPORT_ID, a, b, True) for a, b in ask_old[:cap])
    tries: dict[tuple, int] = {}
    jobs = []
    while any(queues):
        asked = False
        for q in queues:
            if not q:
                continue
            rid, a, b, old = q[0]
            job = mp.exact_job(portal.jobs(100), rid, a, b)
            if not job:
                asked = True                     # a refusal too: the next one of this report waits
                try:
                    job = portal.schedule(a, b, rid, cooldown_wait=False)
                except Exception as exc:  # noqa: BLE001 - that report only
                    key = (rid, a, b)
                    tries[key] = tries.get(key, 0) + 1
                    log.warning("%s report %s %s..%s refused (try %d): %s", account, rid, a, b, tries[key], exc)
                    if old:
                        left.append(f"{a:%b %Y}")
                        q.pop(0)
                        continue
                    if tries[key] < 3:
                        continue                 # stays first: asked again after the cooldown
                    notes.append(f"{account} report {rid} {a}..{b}: {str(exc)[:80]}")
                    q.pop(0)
                    continue
            (old_jobs if old else jobs).append(job)
            q.pop(0)
        if asked and any(queues):
            log.info("%s: waiting for the report cooldown", account)
            time.sleep(COOLDOWN_SECONDS)
    for job in jobs:
        try:
            name, data = portal.collect_or_regenerate(job, wait_minutes=num("MYNTRA_SPF_WAIT_MINUTES", 45))
            (dest / name).write_bytes(data)
        except Exception as exc:  # noqa: BLE001
            notes.append(f"{account} report job {job.get('id')}: {str(exc)[:80]}")
    for job in old_jobs:                         # no regenerating: a next run finds the job done
        try:
            name, data = portal.collect(job, wait_minutes=num("MYNTRA_SPF_OLD_WAIT_MINUTES", 10))
            (dest / name).write_bytes(data)
        except Exception as exc:  # noqa: BLE001
            log.warning("%s old order report job %s left to the next run: %s", account, job.get("id"), exc)
            left.append(str((job.get("queryParams") or {}).get("reportStartDate", job.get("id"))))
    if left:
        notes.append(f"{account}: older order reports still to fetch (next runs): {', '.join(left)}")
    mp.save_to_worker(env, portal.cookies(), account)


def request_reports(today: date) -> list[str]:
    """7 AM step: ask every account for its returns and delivery reports, so the 9 AM run finds
    them done. The order reports are left to the 9 AM run: at 7 AM the SJIT request asks MAIN
    for an order report, and the same report twice within 15 minutes is refused."""
    from ..return_analysis.sources import myntra_portal
    from ..portals import myntra as mp
    windows = _windows(today)
    done = []
    for account in ACCOUNTS:
        try:
            portal = myntra_portal(account)
            for rid in (mp.RETURNS_REPORT_ID, DELIVERY_REPORT_ID):
                for a, b in windows[rid][:1]:
                    job = portal.ensure(rid, a, b)
                    done.append(f"{account} report {rid} {a}..{b}: job {job.get('id')} {job.get('status')}")
            mp.save_to_worker(env, portal.cookies(), account)
        except Exception as exc:  # noqa: BLE001 - the 9 AM run asks again
            done.append(f"{account}: {str(exc)[:100]}")
    return done


# ------------------------------------------------------------------ ERP

def fetch_erp(dest: Path) -> list[str]:
    """GR GRN report and dispatch report as the xlsx files the engine recognises by their headers."""
    from ..portals.erp import Erp
    erp = Erp.from_env()
    erp.login()
    dest.mkdir(parents=True, exist_ok=True)
    out = []
    for name, rows in [("GR GRN Report.xlsx", erp.grn()), ("Dispatch.xlsx", erp.dispatch())]:
        pd.DataFrame(rows).to_excel(dest / name, index=False)
        out.append(f"{name}: {len(rows)} rows")
    return out


# ------------------------------------------------------------------ Drive

def hand_inputs(folder: DriveFolder, dest: Path) -> list[str]:
    """From Drive SPF/MYNTRA/input: the files still dropped by hand (tickets, payments, manual
    returns, legacy master). The portal report folders are fetched fresh instead."""
    src = folder.subfolder("input")
    if not src:
        return []
    got = []
    for name, item in src.children().items():
        if item["mimeType"] != DriveFolder.FOLDER:
            continue
        sub = DriveFolder(item["id"], folder.http)
        if name in ACCOUNTS:
            for fname, f in sub.children().items():     # Ticket1 / payment files at the account root
                if f["mimeType"] != DriveFolder.FOLDER and fname.lower().endswith((".xlsx", ".csv")):
                    sub.download(fname, dest / name / fname)
                    got.append(f"{name}/{fname}")
        elif name in ("MANUAL RETURN", "LEGACY MASTER"):
            for p in sub.download_all(dest / name):
                got.append(f"{name}/{p.name}")
    return got


# ------------------------------------------------------------------ OMSGuru (after the ERP, before the manual sheet)

def _num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def final_status(r: dict) -> str:
    """FINAL RETURN STATUS as the master's formula computes it (engine.build_formulas)."""
    if str(r.get("SPF PAID") or "").strip():
        return "SPF Paid"
    oms, man = str(r.get("OMSGURU STATUS") or "").upper(), str(r.get("MANUAL RETURN STATUS") or "").upper()
    if "RECEIVED" in oms or "RECEIVED" in man or "PICKUP NOT DONE" in man:
        return "Return Received"
    if (_num(r.get("ORDER SETTLEMENT")) is not None and not str(r.get("RETURN DEBIT") or "").strip()
            and not str(r.get("RETURN CHARGE") or "").strip()):
        return "Payment Received"
    return ""


def not_worked(r: dict) -> bool:
    """SJIT and FBM orders are not worked for SPF (engine.NOT_WORKED_MODELS)."""
    from .engine import NOT_WORKED_MODELS
    return str(r.get("model") or "").strip().upper() in NOT_WORKED_MODELS


def missing_order_months(path: Path, today: date) -> dict[str, list[tuple[date, date]]]:
    """Account -> calendar months of order reports still needed: for every worked return with no
    ORDER AMOUNT, the months from MYNTRA_SPF_ORDER_LOOKBACK days (default 45) before it was
    created up to its creation. Months the daily order window already covers are left out."""
    from openpyxl import load_workbook
    from .engine import FIRST_DATA_ROW, HEADER_ROW
    lookback = timedelta(days=num("MYNTRA_SPF_ORDER_LOOKBACK", 45))
    daily_from = today - timedelta(days=num("MYNTRA_SPF_ORDER_DAYS", 60))
    yesterday = today - timedelta(days=1)
    wb = load_workbook(path, read_only=True)
    out = {}
    for name in ACCOUNTS:
        if name not in wb.sheetnames:
            continue
        rows = wb[name].iter_rows(min_row=HEADER_ROW, values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(rows)]
        for _ in range(FIRST_DATA_ROW - HEADER_ROW - 1):
            next(rows, None)
        months = set()
        for values in rows:
            r = dict(zip(head, values))
            if not r.get("order_line_id") or not_worked(r) or _num(r.get("ORDER AMOUNT")) is not None:
                continue
            created = r.get("return_created_date")
            if not hasattr(created, "date"):
                continue
            m, last = (created.date() - lookback).replace(day=1), created.date().replace(day=1)
            while m <= last:
                months.add(m)
                m = (m + timedelta(days=32)).replace(day=1)
        ranges = []
        for m in sorted(months):
            if m < daily_from:               # the daily window covers the rest
                end = min((m + timedelta(days=32)).replace(day=1) - timedelta(days=1), yesterday)
                ranges.append((m, end))
        if ranges:
            out[name] = ranges
    return out


def claim_status(r: dict, age) -> str:
    """CLAIM STATUS as the master's formula computes it (SPF POLICY v5)."""
    if str(r.get("TICKET ID") or "").strip():
        return "CLAIM DONE"
    if not_worked(r) or age is None or final_status(r):
        return "NON ELIGIBLE"
    if r.get("deliver_to_seller_date"):
        return "ELIGIBLE" if age <= 7 else "EXPIRED"
    return "UPCOMING CLAIM" if age < 60 else "ELIGIBLE" if age <= 120 else "EXPIRED"


AGE_BUCKETS = ((30, "0-30"), (60, "31-60"), (120, "61-120"), (None, "120+"))


def master_stats(path: Path, today: date, months: int = 6) -> dict:
    """Dashboard numbers of the master (SJIT / FBM rows only counted as not_worked), per account
    and in total: final status, where the
    received ones were found, claim status with order value, open returns by age, and returns
    per month (last `months`). Same rules as the master's own formulas."""
    from openpyxl import load_workbook
    from .engine import FIRST_DATA_ROW, HEADER_ROW
    wb = load_workbook(path, read_only=True)
    first_month = (today.replace(day=1) - timedelta(days=31 * (months - 1))).strftime("%Y-%m")
    out = {}
    for name in ACCOUNTS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        rows = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(rows)]
        for _ in range(FIRST_DATA_ROW - HEADER_ROW - 1):
            next(rows, None)
        a = {"rows": 0, "final": {}, "received_by": {"ERP": 0, "OMSGuru": 0, "Manual": 0},
             "claims": {}, "open_age": {label: {"count": 0, "value": 0.0} for _, label in AGE_BUCKETS},
             "months": {}}
        for values in rows:
            r = dict(zip(head, values))
            if not r.get("order_line_id"):
                continue
            if not_worked(r):
                a["not_worked"] = a.get("not_worked", 0) + 1
                continue
            a["rows"] += 1
            final = final_status(r) or "Open"
            a["final"][final] = a["final"].get(final, 0) + 1
            if final == "Return Received":
                oms = str(r.get("OMSGURU STATUS") or "").upper()
                src = "ERP" if "ERP" in oms else "OMSGuru" if "RECEIVED" in oms else "Manual"
                a["received_by"][src] += 1
            start = r.get("deliver_to_seller_date") or r.get("return_created_date")
            age = max(0, (today - start.date()).days) if hasattr(start, "date") else None
            value = _num(r.get("ORDER AMOUNT")) or 0.0
            claim = claim_status(r, age)
            c = a["claims"].setdefault(claim, {"count": 0, "value": 0.0, "valued": 0})
            c["count"] += 1
            c["value"] += value
            c["valued"] += value > 0          # rows whose ORDER AMOUNT is known
            if final == "Open" and age is not None:
                label = next(lb for lim, lb in AGE_BUCKETS if lim is None or age <= lim)
                a["open_age"][label]["count"] += 1
                a["open_age"][label]["value"] += value
            created = r.get("return_created_date")
            if hasattr(created, "strftime"):
                m = created.strftime("%Y-%m")
                if m >= first_month:
                    mm = a["months"].setdefault(m, {"returns": 0, "open": 0, "received": 0})
                    mm["returns"] += 1
                    mm["open"] += final == "Open"
                    mm["received"] += final == "Return Received"
        for c in a["claims"].values():
            c["value"] = round(c["value"])
        for b in a["open_age"].values():
            b["value"] = round(b["value"])
        a["months"] = [{"month": k, **v} for k, v in sorted(a["months"].items())]
        out[name] = a
    return out


def open_returns(path: Path, since: date) -> dict[tuple, dict]:
    """(order_line_id, return_id) -> lookup terms for every row whose FINAL RETURN STATUS is
    blank, created since `since`. Only these are looked up in OMSGuru."""
    from openpyxl import load_workbook
    from .engine import FIRST_DATA_ROW, HEADER_ROW, norm
    wb = load_workbook(path, read_only=True)
    out = {}
    for name in ACCOUNTS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        rows = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(rows)]
        for _ in range(FIRST_DATA_ROW - HEADER_ROW - 1):
            next(rows, None)
        for values in rows:
            r = dict(zip(head, values))
            if not r.get("order_line_id") or final_status(r) or not_worked(r):
                continue
            created = r.get("return_created_date")
            created = created.date() if hasattr(created, "date") else None
            if created is None or created < since:
                continue
            out[(norm(r.get("order_line_id")), norm(r.get("return_id")))] = {
                "forward_awb": str(r.get("forward_tracking_number") or "").strip(),
                "return_awb": str(r.get("return_tracking_number") or "").strip(),
                "order_id": str(r.get("STORE ORDER ID") or "").strip(), "sku": "",
            }
    return out


def oms_statuses(wanted: dict[tuple, dict], notes: list[str]) -> dict[tuple, str]:
    """OMSGuru status of each open return: "Return Received OMSGuru" when the warehouse logged
    it in OMSGuru, else its current return status. The Myntra master has no SKU, so an order is
    matched on its AWBs / store order id and any returned item of it counts."""
    from ..portals.omsguru_returns import ReturnLookup

    class MyntraLookup(ReturnLookup):
        def lookup(self, ret):
            ids = []
            for term, kind in ((ret["forward_awb"], "AWB"), (ret["return_awb"], "AWB"), (ret["order_id"], "order")):
                if term:
                    ids = self.search(term, kind)
                    if ids:
                        break
            status = ""
            for oid in ids:
                items, order = self.order(oid)
                for item in items:
                    if item["log"]:
                        return {"status": "Return Received OMSGuru", "received_row": True}
                    status = status or item["status"] or order["latest"]
            return {"status": status, "received_row": None} if ids else None

    if not wanted:
        return {}
    client = MyntraLookup.from_env()
    keys = list(wanted)
    found, failed = client.lookup_many([dict(wanted[k], key=k) for k in keys])
    if failed:
        notes.append(f"OMSGuru: {failed} of {len(keys)} lookups failed")
    return {ret["key"]: res["status"] for ret, res in found if res.get("status")}


def run_engine(work: Path, oms: dict | None = None) -> None:
    """The engine's main() inside the work folder, with OMSGuru results when given."""
    import importlib
    os.environ["MYNTRA_SPF_DIR"] = str(work)
    from . import engine
    engine = importlib.reload(engine)             # picks up MYNTRA_SPF_DIR
    engine.backup = lambda: None                  # Drive keeps the versions
    engine.OMS_RESULTS = oms or {}
    argv, sys.argv = sys.argv, ["myntra_spf", "--in-transit", "y"]
    try:
        engine.main()
    finally:
        sys.argv = argv


# ------------------------------------------------------------------ run

def run(args) -> int:
    if args.request:
        for line in request_reports(date.today()):
            log.info(line)
        return 0
    if not env("MYNTRA_SPF_FOLDER_ID"):
        raise RuntimeError("MYNTRA_SPF_FOLDER_ID (Drive folder SPF/MYNTRA) not set")
    folder = DriveFolder(env("MYNTRA_SPF_FOLDER_ID"))
    found = folder.children().get(MASTER_NAME)
    if not found:
        raise RuntimeError(f"{MASTER_NAME} not found in the Drive folder SPF/MYNTRA")
    work = Path(tempfile.mkdtemp(prefix="myntra_spf_"))
    inp = work / "input"
    master = DriveFile(found["id"])
    master.download(work / MASTER_NAME, need_edit=not args.no_upload)
    notes: list[str] = []
    log.info("Hand inputs from Drive: %s", hand_inputs(folder, inp))

    today = date.today()
    try:
        old_orders = missing_order_months(work / MASTER_NAME, today)
        for a, ranges in old_orders.items():
            log.info("%s: returns without ORDER AMOUNT need order reports of %s", a,
                     ", ".join(f"{s:%b %Y}" for s, _ in ranges))
    except Exception:  # noqa: BLE001 - the daily reports still run
        log.exception("Months of missing order amounts not worked out")
        old_orders = {}
    tasks = {a: (fetch_account, (a, today, inp / a, notes, old_orders.get(a))) for a in ACCOUNTS}
    with ThreadPoolExecutor(max_workers=len(tasks) + 1) as pool:
        futures = {name: pool.submit(fn, *a) for name, (fn, a) in tasks.items()}
        erp = pool.submit(fetch_erp, inp)
    for name, fut in futures.items():
        try:
            fut.result()
        except Exception as exc:  # noqa: BLE001 - the master is still rebuilt from the rest
            log.exception("Myntra %s reports failed", name)
            notes.append(f"Myntra {name} reports not fetched: {str(exc)[:100]}")
    try:
        log.info("ERP: %s", erp.result())
    except Exception as exc:  # noqa: BLE001
        log.exception("ERP reports failed")
        notes.append(f"ERP GRN / dispatch not read: {str(exc)[:100]}")

    # 1) ERP (GRN) and the manual sheet: the engine's own run, in the work folder
    run_engine(work)
    # 2) OMSGuru for the rows still without a final status, then the engine again with them
    try:
        since = today - timedelta(days=num("MYNTRA_SPF_OMS_DAYS", 180))
        wanted = open_returns(work / MASTER_NAME, since)
        log.info("OMSGuru: %d returns with a blank final status since %s", len(wanted), since)
        oms = oms_statuses(wanted, notes)
        received = sum(1 for v in oms.values() if "RECEIVED" in v.upper())
        log.info("OMSGuru: %d statuses, %d received", len(oms), received)
        if oms:
            run_engine(work, oms)
        notes.insert(0, f"OMSGuru checked {len(wanted)} open returns: {received} received")
    except Exception as exc:  # noqa: BLE001 - the ERP / manual result still stands
        log.exception("OMSGuru check failed")
        notes.append(f"OMSGuru not checked: {str(exc)[:100]}")
    out = output_dir("myntra_spf")
    shutil.copy2(work / MASTER_NAME, out / MASTER_NAME)
    if args.no_upload:
        notes.append("not uploaded to Drive (test run): output/myntra_spf")
    else:
        master.upload(work / MASTER_NAME)
    text = summary_text(work / MASTER_NAME, notes)
    print(text) if args.no_telegram else telegram.send(text)
    try:
        accounts = master_stats(work / MASTER_NAME, today)
    except Exception:  # noqa: BLE001 - the master is out; the dashboard only misses details
        log.exception("Master numbers for the dashboard failed")
        accounts = {}
    dashboard.publish("myntra_spf", {"day": today, "notes": notes, "accounts": accounts})
    return 0


def summary_text(path: Path, notes: list[str]) -> str:
    """Rows per account sheet of the new master."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    lines = [f"{'Account':<8}{'Rows':>8}"]
    for name in ACCOUNTS:
        if name in wb.sheetnames:
            lines.append(f"{name:<8}{max(wb[name].max_row - 2, 0):>8}")
    out = ["🧾 <b>Myntra SPF master updated</b>", "<i>returns, tickets, payouts, ERP receipts</i>", "",
           "<pre>" + telegram.esc("\n".join(lines)) + "</pre>", "",
           "📁 Drive: SPF / MYNTRA / MYNTRA SPF MASTER.xlsx"]
    if notes:
        out += ["", "⚠️ <b>Notes</b>", *(f"• {telegram.esc(n)}" for n in notes)]
    return "\n".join(out)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="run.py myntra-spf", description="Myntra SPF master")
    p.add_argument("--request", action="store_true",
                   help="only ask Myntra for the returns and delivery reports (the 7 AM run)")
    p.add_argument("--no-upload", action="store_true", help="keep the master in output/myntra_spf only (test run)")
    p.add_argument("--no-telegram", action="store_true", help="print the summary instead of sending it")
    return run_bot("Myntra SPF master", p, run, argv, key="myntra_spf")
