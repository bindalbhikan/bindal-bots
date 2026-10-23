"""Myntra SPF claims raised by the bot, on the policy's timelines (SPF POLICY v5, Annexure 1).

Two kinds, both raised the way the portal's bulk upload does it (Support > Seller Protection
Fund > Create New Claim > "Here"): an xlsx of order line ids per issue category.

  Return Not Received  (RTS_ORDER_WAS_NEVER_DELIVERED)
      Myntra shows the return delivered to us, but it never reached the warehouse. Policy:
      within 7 days of the delivery date. Raised from day MYNTRA_SPF_RNR_AFTER_DAYS (4) after
      delivery, so the GRN / OMSGuru have had time to book a return that did arrive.
  Delayed Returns      (DELAYED_DELIVERED_RETURN)
      Still in transit, or no return tracking at all. Policy: after 60 and up to 120 days from
      the return / RTO creation date. Raised from day MYNTRA_SPF_DELAY_AFTER_DAYS (61).

A row is claimed only when nothing says it is settled: no ticket on it (master TICKET ID or
any ticket on the portal for that order line, any status - the policy allows one claim per
order and treats repeat tickets as fraud), FINAL RETURN STATUS blank (not received in the ERP /
OMSGuru / manual sheet, no SPF payout, no clean settlement), not an SJIT / FBM order, and no
RETURN POD from the courier (a signed return handover).

Every ticket on the portal is also written into the work folder as a ticket report, so the
master's TICKET ID / STATUS columns follow the portal. A claim just raised gets
TICKET ID "BOT JOB <job id>" until its real ticket turns up on the next run.
"""
from __future__ import annotations

import json
import logging
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from ..core.config import num

log = logging.getLogger(__name__)

RNR = "RTS_ORDER_WAS_NEVER_DELIVERED"
DELAYED = "DELAYED_DELIVERED_RETURN"
CATEGORY_NAME = {RNR: "Return Not Received", DELAYED: "Delayed Returns"}
DESCRIPTION = {
    RNR: "RETURN DELIVERED BUT PHYSICAL RETURN PARCEL NOT RECEIVED YET",
    DELAYED: "RETURN SHOWING IN TRANSIT, KINDLY CHECK AND RETURN THE PARCEL ASAP",
}
# The ticket report the engine reads (kind TICKET: ticketId + orderLineId columns)
TICKET_FILE = "SPF Tickets (portal).xlsx"
BOT_FILE = "SPF Tickets (raised by bot).xlsx"
TICKET_HEADER = ["ticketId", "orderLineId", "returnId", "issueCategory", "status", "createdOn", "validationReason"]
# Tickets the bot just raised whose real ticket turned up in the same run (see resolve_raised)
NEW_FILE = "SPF Tickets (new).xlsx"

# The master's TICKETS sheet: every claim of the last MYNTRA_SPF_DETAIL_DAYS (120) days and any older
# one still open, as its portal page shows it, with Seller Support's replies.
DETAIL_SHEET = "TICKETS"
DETAIL_COLS = ["ACCOUNT", "TICKET ID", "ORDER LINE ID", "RETURN ID", "ORDER ID", "CATEGORY", "STATUS",
               "NEEDS OUR REPLY", "CREATED", "LAST UPDATED", "STATUS HISTORY", "REASON", "COMPENSATION",
               "UTR", "PAYMENT INITIATED", "DESCRIPTION", "SUPPORT TICKET", "REPLIES", "LAST REPLY BY",
               "LAST REPLY AT", "LAST REPLY"]
FINAL_STATUSES = {"REJECT", "PAYMENT_COMPLETED", "CLOSED"}
AWAITING_US = "AWAITING_SELLER_RESPONSE"
IST = timedelta(hours=5, minutes=30)


def _day(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def claim_of(r: dict, today: date, ticketed: set[str]) -> str | None:
    """The category to claim this master row under today, or None."""
    from .pipeline import final_status, not_worked
    from .engine import norm, is_delivered_status
    oli = norm(r.get("order_line_id"))
    if not oli or oli in ticketed or not _blank(r.get("TICKET ID")):
        return None
    if not_worked(r) or final_status(r):
        return None
    if str(r.get("RETURN POD") or "").strip().upper() == "YES":
        return None
    delivered = _day(r.get("deliver_to_seller_date"))
    if delivered or is_delivered_status(r.get("status")):
        if not delivered:
            return None                          # delivered without a date: the TAT cannot be told
        age = (today - delivered).days
        lo, hi = num("MYNTRA_SPF_RNR_AFTER_DAYS", 4), num("MYNTRA_SPF_RNR_MAX_DAYS", 7)
        return RNR if lo <= age <= hi else None
    created = _day(r.get("return_created_date"))
    if not created:
        return None
    age = (today - created).days
    lo, hi = num("MYNTRA_SPF_DELAY_AFTER_DAYS", 61), num("MYNTRA_SPF_DELAY_MAX_DAYS", 120)
    return DELAYED if lo <= age <= hi else None


def candidates(path: Path, today: date, ticketed: dict[str, set[str]],
               skip: set[tuple] = frozenset()) -> dict[str, dict[str, list[str]]]:
    """account -> category -> order line ids to claim, from the master's account sheets.
    ticketed: account -> order lines with a portal ticket. skip: (order_line_id, return_id)
    keys known to be received (OMSGuru) since the master was written."""
    from openpyxl import load_workbook
    from .engine import FIRST_DATA_ROW, HEADER_ROW, SHEETS, norm
    wb = load_workbook(path, read_only=True)
    out: dict[str, dict[str, list[str]]] = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        rows = wb[name].iter_rows(min_row=HEADER_ROW, values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(rows)]
        for _ in range(FIRST_DATA_ROW - HEADER_ROW - 1):
            next(rows, None)
        seen: set[str] = set()
        for values in rows:
            r = dict(zip(head, values))
            if (norm(r.get("order_line_id")), norm(r.get("return_id"))) in skip:
                continue
            cat = claim_of(r, today, ticketed.get(name, set()))
            oli = norm(r.get("order_line_id"))
            if cat and oli not in seen:
                seen.add(oli)
                out.setdefault(name, {}).setdefault(cat, []).append(oli)
    wb.close()
    return out


def write_ticket_file(path: Path, rows: list[dict]) -> None:
    """xlsx, not CSV: ticket ids then stay numbers in the master, as from the hand export."""
    from openpyxl import Workbook
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("tickets")
    ws.append(TICKET_HEADER)
    for t in rows:
        ws.append([t.get(c) for c in TICKET_HEADER])
    wb.save(path)


def portal_tickets(portal, account: str, dest: Path, cached: dict | None = None,
                   details: dict | None = None) -> set[str]:
    """Every ticket of the account into dest/<account>/TICKET_FILE; the order lines they cover.
    With `details` (account -> TICKETS rows), the recent / open claims are read in full too
    (ticket_details, reusing `cached` rows of unchanged tickets) and a rejection's reason goes
    into the ticket file, so the master's REJECTION REASON shows it."""
    from .engine import norm
    rows = portal.spf_tickets()
    if details is not None:
        # The full list is 35 pages and the portal's order can shift while they are read (on
        # 23 Oct 2026 one read missed 900 of MAIN's 2,288 recent claims): the recent ones are
        # read again on their own with the claim date filter, a short and steady list.
        since = datetime.now() - timedelta(days=num("MYNTRA_SPF_DETAIL_DAYS", 120))
        seen = set()
        rows = [t for t in rows + portal.spf_tickets(since=since)
                if not (t.get("ticketId") in seen or seen.add(t.get("ticketId")))]
        details[account] = ticket_details(portal, account, rows, cached or {})
        reason = {d["TICKET ID"]: d["REASON"] for d in details[account] if d.get("REASON")}
        for t in rows:
            t["validationReason"] = reason.get(str(t.get("ticketId")))
    write_ticket_file(dest / account / TICKET_FILE, rows)
    log.info("%s: %d SPF tickets on the portal", account, len(rows))
    return {norm(t.get("orderLineId")) for t in rows if t.get("orderLineId")}


# ------------------------------------------------------------------ ticket details

def _utc(v) -> datetime | None:
    try:
        return datetime.strptime(str(v)[:19], "%Y-%m-%dT%H:%M:%S") + IST
    except (TypeError, ValueError):
        return None


def _ist(v) -> datetime | None:
    """The portal's MM-DD-YYYY HH:MM:SS (IST) times."""
    try:
        return datetime.strptime(str(v)[:19], "%m-%d-%Y %H:%M:%S")
    except (TypeError, ValueError):
        return None


def _when(d: datetime | None) -> str:
    return f"{d:%d-%b-%Y %H:%M}" if d else ""


def plain(html) -> str:
    t = re.sub(r"<[^>]+>", " ", str(html or "")).replace("&nbsp;", " ").replace(" ", " ")
    return re.sub(r"\s+", " ", t).strip().lstrip(": ").strip()


def history_text(history) -> str:
    """{"OPEN": "09-30-2026 17:28:03", "REJECT": "10-01-2026 18:37:21"} -> OPEN 30-Sep 17:28 > REJECT 01-Oct 18:37"""
    try:
        h = json.loads(history) if isinstance(history, str) else (history or {})
    except ValueError:
        return str(history or "")
    steps = sorted(((_ist(v), k) for k, v in h.items()), key=lambda s: s[0] or datetime.min)
    return " > ".join(f"{k} {d:%d-%b %H:%M}" if d else k for d, k in steps)


def detail_row(account: str, t: dict, claim: dict, comments: list[dict]) -> dict:
    """One TICKETS row from the ticket list entry, its claim page and its Seller Support replies."""
    last = comments[0] if comments else {}
    status = claim.get("status") or t.get("status")
    reason = claim.get("rejectReasonDescription") or claim.get("rejectReason") or claim.get("validationReason")
    return {
        "ACCOUNT": account, "TICKET ID": str(t.get("ticketId")), "ORDER LINE ID": str(t.get("orderLineId") or ""),
        "RETURN ID": str(t.get("returnId") or claim.get("returnId") or ""),
        "ORDER ID": str(t.get("orderId") or claim.get("orderId") or ""),
        "CATEGORY": CATEGORY_NAME.get(t.get("issueCategory"), t.get("issueCategory")), "STATUS": status,
        "NEEDS OUR REPLY": "YES" if status == AWAITING_US else "",
        "CREATED": _when(_utc(t.get("createdOn"))), "LAST UPDATED": _when(_utc(t.get("lastModifiedOn"))),
        "STATUS HISTORY": history_text(claim.get("history")), "REASON": reason or "",
        "COMPENSATION": claim.get("compensationAmount"), "UTR": claim.get("utr") or "",
        "PAYMENT INITIATED": str(claim.get("paymentInitiatedOn") or ""),
        "DESCRIPTION": claim.get("description") or "", "SUPPORT TICKET": str(claim.get("fwTicketId") or ""),
        "REPLIES": len(comments), "LAST REPLY BY": last.get("user") or "",
        "LAST REPLY AT": _when(_ist(last.get("createdOn"))), "LAST REPLY": plain(last.get("comment"))[:2000],
    }


def ticket_details(portal, account: str, rows: list[dict], cached: dict[str, dict]) -> list[dict]:
    """TICKETS rows of the account: claims created in the last MYNTRA_SPF_DETAIL_DAYS (120) days
    and older ones not settled yet. A ticket whose LAST UPDATED is the cached row's is not read
    again; at most MYNTRA_SPF_DETAIL_MAX (3000) claim pages a run, the rest keep the cached row
    (or the list's own fields) until the next run."""
    since = datetime.now() - timedelta(days=num("MYNTRA_SPF_DETAIL_DAYS", 120))
    keep = [t for t in rows if (_utc(t.get("createdOn")) or datetime.min) >= since
            or t.get("status") not in FINAL_STATUSES]
    todo, out = [], {}
    for t in keep:
        key, c = str(t.get("ticketId")), cached.get(str(t.get("ticketId")))
        if c and c.get("LAST UPDATED") == _when(_utc(t.get("lastModifiedOn"))):
            out[key] = {**c, "ACCOUNT": account}
        else:
            todo.append(t)
    cap = num("MYNTRA_SPF_DETAIL_MAX", 3000)

    def one(t):
        try:
            claim = portal.spf_claim(t["ticketId"])
            fw = claim.get("fwTicketId")
            return detail_row(account, t, claim, portal.spf_comments(fw) if fw else [])
        except Exception as exc:  # noqa: BLE001 - the cached row / the list's fields stand
            log.warning("%s ticket %s not read: %s", account, t.get("ticketId"), str(exc)[:100])
            return cached.get(str(t.get("ticketId"))) or detail_row(account, t, {}, [])

    with ThreadPoolExecutor(max_workers=num("MYNTRA_SPF_DETAIL_WORKERS", 6)) as pool:
        for row in pool.map(one, todo[:cap]):
            out[row["TICKET ID"]] = row
    for t in todo[cap:]:
        out[str(t["ticketId"])] = cached.get(str(t["ticketId"])) or detail_row(account, t, {}, [])
    # a claim already on the sheet stays there, even when this run's list missed it
    for key, c in cached.items():
        if key not in out and c.get("ACCOUNT") == account:
            out[key] = c
    log.info("%s: %d claims on the TICKETS sheet (%d read now, %d unchanged)", account, len(out),
             min(len(todo), cap), len(keep) - len(todo))
    return sorted(out.values(), key=lambda d: d["TICKET ID"], reverse=True)


def read_details(path: Path) -> dict[str, dict]:
    """The previous master's TICKETS rows by ticket id (none when the sheet is not there)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        if DETAIL_SHEET not in wb.sheetnames:
            return {}
        rows = wb[DETAIL_SHEET].iter_rows(values_only=True)
        head = [str(h or "") for h in next(rows, [])]
        out = {}
        for v in rows:
            r = {h: ("" if x is None else x) for h, x in zip(head, v)}
            if r.get("TICKET ID"):
                r["TICKET ID"] = str(r["TICKET ID"])
                out[r["TICKET ID"]] = r
        return out
    finally:
        wb.close()


def resolve_raised(portal, account: str, olis: list[str], since: datetime, cats: dict[str, str]) -> list[dict]:
    """The real tickets of claims just raised (a bulk job creates them a minute or so later):
    the account's tickets created since `since`, asked every 20 s for up to
    MYNTRA_SPF_RESOLVE_MINUTES (2) minutes - they turn up within 2 - until every order line has one."""
    from .engine import norm
    want, found = set(olis), {}
    deadline = time.time() + 60 * num("MYNTRA_SPF_RESOLVE_MINUTES", 2)
    while True:
        try:
            for t in portal.spf_tickets(since=since - timedelta(minutes=10), page_size=200):
                oli = norm(t.get("orderLineId"))
                if oli in want and (oli not in found or str(t.get("issueCategory")) == cats.get(oli)):
                    found[oli] = t
        except Exception as exc:  # noqa: BLE001 - asked again, or the BOT JOB rows stay
            log.warning("%s new tickets not read: %s", account, str(exc)[:100])
        if len(found) == len(want) or time.time() > deadline:
            break
        time.sleep(20)
    log.info("%s: %d of %d claims just raised have their ticket", account, len(found), len(want))
    return list(found.values())


def raise_claims(portal, account: str, wanted: dict[str, list[str]], dest: Path, live: bool,
                 seller: int | None = None, details: list[dict] | None = None) -> list[str]:
    """Raise this account's claims (live) or only list them. The real tickets are waited for
    (resolve_raised) and written as NEW_FILE, so the master shows their ticket ids at once; a
    claim whose ticket has not turned up yet is in the bot's file as "BOT JOB <job id>" until
    the next run. Their TICKETS rows go into `details`. One line per category for the report."""
    from ..portals.myntra import SPF_BULK_MAX
    cap = num("MYNTRA_SPF_TICKET_MAX", SPF_BULK_MAX)
    lines, raised = [], []
    started = datetime.now()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    for cat in (RNR, DELAYED):
        ids = wanted.get(cat) or []
        if not ids:
            continue
        todo, rest = ids[:cap], len(ids) - min(len(ids), cap)
        more = f", {rest} left for the next run" if rest else ""
        if not live:
            lines.append(f"{account} {CATEGORY_NAME[cat]}: {len(todo)} due (not raised: tickets off){more}")
            continue
        jobs = []
        for i in range(0, len(todo), SPF_BULK_MAX):
            part = todo[i:i + SPF_BULK_MAX]
            try:
                job = portal.spf_bulk_claim(part, cat, DESCRIPTION[cat], seller=seller)
            except Exception as exc:  # noqa: BLE001 - the other category / account still go
                log.exception("%s %s claim failed", account, cat)
                lines.append(f"{account} {CATEGORY_NAME[cat]}: {len(part)} NOT raised: {str(exc)[:80]}")
                continue
            jobs.append(job)
            raised += [{"ticketId": f"BOT JOB {job}", "orderLineId": o, "issueCategory": cat,
                        "status": "RAISED BY BOT", "createdOn": now} for o in part]
        if jobs:
            n = sum(1 for t in raised if t["issueCategory"] == cat)
            lines.append(f"{account} {CATEGORY_NAME[cat]}: {n} raised (job {', '.join(map(str, jobs))}){more}")
    if raised:
        real = resolve_raised(portal, account, [t["orderLineId"] for t in raised], started,
                              {t["orderLineId"]: t["issueCategory"] for t in raised})
        from .engine import norm
        got = {norm(t.get("orderLineId")) for t in real}
        write_ticket_file(dest / account / NEW_FILE, real)
        write_ticket_file(dest / account / BOT_FILE, [t for t in raised if t["orderLineId"] not in got])
        if details is not None:
            details += [detail_row(account, t, {}, []) for t in real]
        if len(got) < len(raised):
            lines.append(f"{account}: {len(raised) - len(got)} claims raised but no ticket id yet (BOT JOB, next run)")
    return lines
