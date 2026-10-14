"""Courier pickup check for returns whose final return status is not "Return Received".

Each such return's AWB, and its forward AWB, is tracked on the courier's own
site (the tracker package shared with the AJIO TAT bot) and these columns are
written into the master's Sheet1:

  COURIER PICKUP / PICKUP DATE / COURIER STATUS / COURIER UPDATED   return AWB
  FWD COURIER STATUS / FWD COURIER UPDATED                           forward AWB

A return the courier never picked up (or whose pickup was cancelled / failed)
gets no SPF ticket and no follow-up: there is nothing AJIO owes for it.
Couriers keep only recent history (Delhivery answers "very old package"), so
where the courier has nothing, AJIO's own 3PL Delivery Status decides.
"""
from __future__ import annotations

import json
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..awb_tracker import CourierRegistry, Job, track_all
from ..awb_tracker import models as m
from ..core.config import ROOT, env
from . import claims

log = logging.getLogger(__name__)

COLUMNS = ("COURIER PICKUP", "PICKUP DATE", "COURIER STATUS", "COURIER UPDATED")
FWD_COLUMNS = ("FWD COURIER STATUS", "FWD COURIER UPDATED")
PICKED = "PICKED UP"
NOT_PICKED = "NOT PICKED UP"
CANCELLED = "PICKUP CANCELLED"
FAILED = "PICKUP FAILED"
NO_TRACKING = "NO TRACKING"
NO_PICKUP = {NOT_PICKED, CANCELLED, FAILED}
# Already settled: tracking them again every day only slows the run down.
FINAL = re.compile(r"deliver|dlvd|rto|returned|cancel", re.I)

MOVED = {m.PICKED_UP, m.RETURN_PICKED_UP, m.IN_TRANSIT, m.RETURN_IN_TRANSIT, m.OUT_FOR_DELIVERY,
         m.RETURN_OUT_FOR_DELIVERY, m.DELIVERED, m.RETURN_DELIVERED, m.RTO, m.UNDELIVERED, m.LOST}
WAITING = {m.PENDING, m.RETURN_PENDING, m.OUT_FOR_PICKUP}
# Courier wording the shared classifier does not know.
TEXT_RULES = (
    (re.compile(r"cancel", re.I), CANCELLED),
    (re.compile(r"pick ?up (attempt|fail)|not picked|unable to pick|wrong pincode|customer (not|refus)", re.I), FAILED),
    (re.compile(r"out (to|for) p/?u|out ?for ?pick|pick ?up sched|pending pick", re.I), NOT_PICKED),
    (re.compile(r"picked|in ?transit|dlvd|deliver|received|returned|rto|bag|hub|dispatched|connected|"
                r"prod_replaced|reached", re.I), PICKED),
)
# AJIO's own 3PL Delivery Status, used when the courier has no history.
AJIO_RULES = (
    (re.compile(r"cancel", re.I), CANCELLED),
    (re.compile(r"picked|intransit|in transit|delivered|delayed", re.I), PICKED),
)
# Scans that happen before the courier has the parcel in hand.
PRE_PICKUP = re.compile(r"out ?(for|to) ?p|pick ?up (attempt|fail|sched|cancel|reschedul)|"
                        r"data received|manifest|soft ?data|booked|not picked|cancel", re.I)
CONFIG = ROOT / "tracking_config.json"


def _by_text(rules, *texts) -> str:
    for text in texts:
        for pattern, verdict_ in rules:
            if text and pattern.search(text):
                return verdict_
    return ""


def verdict(result, ajio_status: str = "") -> str:
    if result is not None and result.ok and result.status not in (m.NOT_FOUND, m.UNKNOWN):
        if result.status == m.PICKUP_FAILED:
            return FAILED
        if result.status == m.CANCELLED:
            return CANCELLED
        if result.status in WAITING:
            return NOT_PICKED
        if result.status in MOVED:
            return PICKED
    if result is not None and result.ok:
        found = _by_text(TEXT_RULES, result.status_raw, result.status_detail)
        if found:
            return found
    return _by_text(AJIO_RULES, ajio_status) or NO_TRACKING


def pickup_date(result) -> date | None:
    """The first scan after which the courier had the parcel."""
    for ev in sorted(result.events if result else [], key=m.TrackEvent.sort_key):
        if ev.timestamp and not PRE_PICKUP.search(f"{ev.status} {ev.remark}"):
            return as_date(ev.timestamp)
    return None


def as_date(value: str):
    try:
        return datetime.fromisoformat(str(value)[:10]).date() if value else None
    except ValueError:
        return None


def settled(row: dict) -> bool:
    return (claims.txt(row.get("COURIER PICKUP")) == CANCELLED
            or (claims.txt(row.get("COURIER PICKUP")) == PICKED and bool(FINAL.search(claims.txt(row.get("COURIER STATUS"))))))


def check(path: Path, workers: int = 8) -> dict:
    """Track the return and forward AWBs of every return not received and write
    the courier columns into Sheet1. Returns counts."""
    rows = claims.read_sheet(path)
    jobs, seen = [], set()
    for r in rows:
        if claims.is_received(r):
            continue
        for awb_col, carrier_col, done in (("Return AWB No", "Return Carrier Name", settled(r)),
                                           ("FWD AWB", "FWD Carrier Name", False)):
            awb = claims.txt(r.get(awb_col))
            if awb and awb not in seen and not done:
                seen.add(awb)
                jobs.append(Job(awb, claims.txt(r.get(carrier_col))))
    cfg = json.loads(CONFIG.read_text(encoding="utf-8")) if CONFIG.exists() else {}
    if env("SHADOWFAX_TOKEN"):
        cfg.setdefault("couriers", {}).setdefault("SHADOWFAX", {})["token"] = env("SHADOWFAX_TOKEN")
    results = {r.awb: r for r in track_all(jobs, CourierRegistry(cfg), workers=workers, progress=False)}

    wb = load_workbook(path)
    ws = wb["Sheet1"]
    head = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    for name in COLUMNS + FWD_COLUMNS:
        if name not in head:
            head[name] = ws.max_column + 1
            ws.cell(1, head[name], name).font = copy(ws.cell(1, 1).font)

    def put(row, name, value):
        cell = row[head[name] - 1]
        cell.value = value
        if isinstance(value, date):
            cell.number_format = "dd-mm-yyyy"

    counts = {"tracked": len(results), PICKED: 0, NOT_PICKED: 0, CANCELLED: 0, FAILED: 0, NO_TRACKING: 0}
    for row in ws.iter_rows(min_row=2):
        if claims.is_received({k: row[c - 1].value for k, c in head.items() if k in
                               ("OMSGURU STATUS", "MANUAL RETURN STATUS")}):
            continue
        awb = claims.txt(row[head["Return AWB No"] - 1].value)
        ajio = claims.txt(row[head["3PL Delivery Status"] - 1].value) if "3PL Delivery Status" in head else ""
        res = results.get(awb)
        if awb and (res is not None or not claims.txt(row[head["COURIER PICKUP"] - 1].value)):
            v = verdict(res, ajio)
            put(row, "COURIER PICKUP", v)
            put(row, "PICKUP DATE", pickup_date(res) if v == PICKED else None)
            put(row, "COURIER STATUS", (res.status_raw or res.status or res.error) if res else ajio)
            put(row, "COURIER UPDATED", as_date(res.last_update) if res else None)
        fwd = results.get(claims.txt(row[head["FWD AWB"] - 1].value)) if "FWD AWB" in head else None
        if fwd is not None:
            put(row, "FWD COURIER STATUS", fwd.status_raw or fwd.status or fwd.error)
            put(row, "FWD COURIER UPDATED", as_date(fwd.last_update))
        v = claims.txt(row[head["COURIER PICKUP"] - 1].value)
        if v in counts:
            counts[v] += 1
    wb.save(path)
    log.info("Courier pickup: %s", counts)
    return counts
