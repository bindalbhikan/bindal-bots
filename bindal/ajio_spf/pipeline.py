"""Daily AJIO SPF run.

  fetch     AJIO returns report + SPF tickets, ERP GRN and dispatch, SRM ledger,
            manual register (Drive)
  build     the user's master script (master_engine.py) on those inputs, then
            OMSGuru order lookups for the returns still open, and a rebuild
  raise     SPF tickets for returns that crossed a claim threshold (claims.py)
  follow    reply on / re-open tickets where AJIO answered (followup.py)
  publish   master back to Drive, action list + summary to Telegram

SPF_DRY_RUN=1 (the default) does everything except posting to AJIO: the planned
tickets and drafted replies are sent to Telegram for review instead. Only an
explicit 0 / false / no / off turns posting on.
"""
from __future__ import annotations

import argparse
import contextlib
import importlib
import io
import logging
import os
import shutil
import sys
import time
import traceback
from datetime import date, datetime, timedelta
from pathlib import Path

from ..core import telegram
from ..core.app import run_bot
from ..core.config import env, flag, num, output_dir
from . import claims, followup, inputs
from .inputs import OMS_COLUMNS, write_xlsx

OUTPUT = output_dir("ajio_spf")
WORK = OUTPUT / "work"
MASTER_NAME = "AJIO MASTER.xlsx"
ENGINE = __package__ + ".master_engine"
log = logging.getLogger("bindal.ajio_spf")


def dry_run(cli_value: str | None) -> bool:
    """True unless posting is switched on explicitly (--dry-run 0 or SPF_DRY_RUN=0/false/no/off)."""
    if cli_value is not None:
        return cli_value != "0"
    return flag("SPF_DRY_RUN", default=True)


# ------------------------------------------------------------------ fetch

def rtv_start(today: date) -> date:
    """First day the returns report is downloaded from.

    New data starts where the master's INPUT LOG coverage ends, or at its first
    missing day. The last RTV_REFRESH_DAYS (90) are always downloaded again:
    open returns keep changing (delivered date, credit note) and the claim
    rules depend on those fields.
    """
    refresh = today - timedelta(days=num("RTV_REFRESH_DAYS", 90))
    try:
        from openpyxl import load_workbook
        wb = load_workbook(WORK / MASTER_NAME, read_only=True)
        spans = sorted((claims.as_date(r[4]), claims.as_date(r[5]))
                       for r in list(wb["INPUT LOG"].iter_rows(values_only=True))[1:]
                       if r and claims.as_date(r[4]) and claims.as_date(r[5]))
        wb.close()
    except Exception as exc:  # noqa: BLE001 - no log yet: fall back to the refresh window
        log.warning("INPUT LOG not readable (%s); returns report from %s", exc, refresh)
        return refresh
    if not spans:
        return refresh
    reach = spans[0][1]
    for start, end in spans[1:]:
        if start > reach + timedelta(days=1):
            break                       # first gap in the coverage
        reach = max(reach, end)
    first_missing = reach + timedelta(days=1)
    log.info("Returns report coverage runs to %s; downloading from %s", reach, min(first_missing, refresh))
    return min(first_missing, refresh)


def fetch_all(today: date, ctx: dict) -> None:
    """Fill WORK/input with fresh exports. Any failure is recorded in ctx['errors']."""
    from .sources.ajio import AjioClient
    from .sources.erp import Erp
    from .sources.srm import Srm, payment_rows

    inp = WORK / "input"
    lookups = inp / "Manual Sheet"
    lookups.mkdir(parents=True, exist_ok=True)

    def step(name, fn):
        t0 = time.time()
        try:
            fn()
            log.info("%s done in %.0fs", name, time.time() - t0)
        except Exception as exc:  # noqa: BLE001 - one source failing must not stop the others
            log.error("%s failed: %s", name, exc)
            log.debug(traceback.format_exc())
            ctx["errors"].append(f"{name}: {exc}")

    def ajio():
        client = AjioClient(env("AJIO_USERNAME"), env("AJIO_PASSWORD"))
        client.login()
        ctx["ajio"] = client
        start = rtv_start(today)
        for a, b, blob in client.rtv_reports(start, today):
            (inp / f"{a:%Y-%m-%d} to {b:%Y-%m-%d}.xlsx").write_bytes(blob)
        listed = client.ticket_list(today - timedelta(days=num("TICKET_DAYS", 240)), today)
        spf_ids = [str(t["ticketId"]) for t in listed if str(t.get("ticketType") or "").upper() == "SPF"]
        details = client.ticket_details(spf_ids)
        ctx["tickets"] = details
        write_xlsx(inp / "Tickets.xlsx", inputs.ticket_sheets(details, client.seller_id))
        log.info("AJIO: %d tickets listed, %d SPF tickets read", len(listed), len(details))

    def erp():
        client = Erp(env("ERP_USERNAME"), env("ERP_PASSWORD"), env("ERP_USER_TOKEN"))
        client.login()
        grn = client.grn()
        write_xlsx(inp / "GR GRN Report.xlsx", {"Sheet": (list(grn[0].keys()) if grn else [], grn)})
        dispatch = [r for r in client.dispatch() if "AJIO" in str(r.get("Order Portal Name") or "").upper()]
        write_xlsx(inp / "Dispatch.xlsx", {"Sheet": (list(dispatch[0].keys()) if dispatch else [], dispatch)})

    def srm():
        client = Srm(env("SRM_USERNAME"), env("SRM_PASSWORD"))
        client.login()
        items = client.statement(today - timedelta(days=num("SRM_DAYS", 365)), today)
        rows = payment_rows(items)
        write_xlsx(inp / "PAYMENT.xlsx", {"data": (list(rows[0].keys()) if rows else [], rows)})

    step("AJIO", ajio)
    step("ERP", erp)
    step("SRM", srm)


def fetch_drive(ctx: dict) -> None:
    """Master (editable) plus the manual register and old OMS sheet (read only) from Drive."""
    lookups = WORK / "input" / "Manual Sheet"
    master_id = env("SPF_MASTER_DRIVE_FILE_ID")
    local = env("SPF_MASTER_XLSX_PATH")
    if master_id:
        from ..core.drive import DriveFile
        ctx["master_drive"] = DriveFile(master_id)
        ctx["master_drive"].download(WORK / MASTER_NAME, need_edit=True)
        http = ctx["master_drive"].http
        for var, name in (("SPF_MANUAL_DRIVE_FILE_ID", "AJIO RETURN SHEET MANUAL - DRIVE.xlsx"),
                          ("SPF_OMS_DRIVE_FILE_ID", "AJIO RETURN SHEET OMS GURU - DRIVE.xlsx")):
            if env(var):
                try:
                    DriveFile(env(var), http).download(lookups / name)
                except Exception as exc:  # noqa: BLE001
                    ctx["errors"].append(f"Drive {name}: {exc}")
    elif local:
        shutil.copy2(local, WORK / MASTER_NAME)
        for var, name in (("SPF_MANUAL_XLSX_PATH", "AJIO RETURN SHEET MANUAL - DRIVE.xlsx"),
                          ("SPF_OMS_XLSX_PATH", "AJIO RETURN SHEET OMS GURU - DRIVE.xlsx")):
            if env(var) and Path(env(var)).exists():
                shutil.copy2(env(var), lookups / name)
    else:
        raise RuntimeError("Set SPF_MASTER_DRIVE_FILE_ID (or SPF_MASTER_XLSX_PATH for a local run)")


# ------------------------------------------------------------------ build

def build_master() -> str:
    """Run the master script inside WORK; returns its printed log."""
    sys.modules.pop(ENGINE, None)
    buf = io.StringIO()
    cwd = os.getcwd()
    os.chdir(WORK)
    try:
        engine = importlib.import_module(ENGINE)
        with contextlib.redirect_stdout(buf):
            engine.main()
    except SystemExit as exc:
        raise RuntimeError(f"Master build stopped: {exc}\n{buf.getvalue()[-1500:]}") from None
    finally:
        os.chdir(cwd)
    (OUTPUT / "master_build.log").write_text(buf.getvalue(), encoding="utf-8")
    return buf.getvalue()


def oms_lookup(today: date, ctx: dict) -> int:
    """Look returns up in OMSGuru and write what it knows into the master.

    Looked up: every return not yet received (last OMS_LOOKUP_DAYS, 180) and
    every row whose FWD AWB is blank (the returns report leaves it out).
    Written back to Sheet1: the forward AWB and carrier where blank (the master
    script then matches the ERP dispatch register on it) and the current
    OMSGuru return status of returns not received. Receipts go to a return
    sheet that the master script reads. Returns how many rows changed.
    """
    from openpyxl import load_workbook
    from .sources.omsguru_returns import ReturnLookup
    path = WORK / MASTER_NAME
    rows = claims.read_sheet(path)
    oldest = today - timedelta(days=num("OMS_LOOKUP_DAYS", 180))
    wanted, seen = [], set()
    for r in rows:
        created = claims.as_date(r.get("Return Created Date"))
        open_return = (not claims.is_received(r) and not claims.is_cancelled(r)
                       and created is not None and created >= oldest)
        if not open_return and claims.txt(r.get("FWD AWB")):
            continue
        f = claims.row_facts(r)
        key = (f["return_order"], f["sku"])
        if key in seen:
            continue
        seen.add(key)
        wanted.append({"return_order": f["return_order"], "return_awb": f["return_awb"],
                       "forward_awb": f["forward_awb"], "sku": f["sku"], "invoice": f["invoice"],
                       "order_id": claims.txt(r.get("FWD Seller Order ID"))})
    client = ReturnLookup.from_env()
    found, failed = client.lookup_many(wanted)
    if failed > max(5, len(wanted) // 10):
        raise RuntimeError(f"{failed} of {len(wanted)} OMSGuru lookups failed")

    received = [res["received_row"] for _, res in found if res["received_row"]]
    if received:
        write_xlsx(WORK / "input" / "Manual Sheet" / "AJIO RETURN SHEET OMS GURU - LIVE.xlsx",
                   {"Sheet1": (OMS_COLUMNS, received)})
    by_key = {(ret["return_order"], ret["sku"]): res for ret, res in found}
    wb = load_workbook(path)
    ws = wb["Sheet1"]
    col = {str(c.value).strip(): c.column - 1 for c in ws[1] if c.value}
    need = ("RETURN ORDER NUMBER", "SELLER SKU", "FWD AWB", "FWD Carrier Name", "OMSGURU STATUS", "MANUAL RETURN STATUS")
    if any(h not in col for h in need):
        raise RuntimeError(f"Master Sheet1 lacks one of {need}")
    awbs = statuses = 0
    for row in ws.iter_rows(min_row=2):
        res = by_key.get((claims.txt(row[col["RETURN ORDER NUMBER"]].value), claims.txt(row[col["SELLER SKU"]].value)))
        if not res:
            continue
        if res["forward_awb"] and not claims.txt(row[col["FWD AWB"]].value):
            row[col["FWD AWB"]].value = res["forward_awb"]
            if res["carrier"] and not claims.txt(row[col["FWD Carrier Name"]].value):
                row[col["FWD Carrier Name"]].value = res["carrier"]
            awbs += 1
        current = claims.txt(row[col["OMSGURU STATUS"]].value)
        if (res["status"] and not res["received_row"] and "RECEIVED" not in current.upper()
                and claims.txt(row[col["MANUAL RETURN STATUS"]].value).upper() != "RETURN RECEIVED"
                and current != res["status"]):
            row[col["OMSGURU STATUS"]].value = res["status"]
            statuses += 1
    wb.save(path)
    log.info("OMSGuru: %d forward AWBs and %d return statuses written to the master, %d receipts",
             awbs, statuses, len(received))
    ctx.update(oms_checked=len(wanted), oms_found=len(received), fwd_awbs=awbs)
    return awbs + statuses + len(received)


# ------------------------------------------------------------------ raise

def describe(cand: claims.Candidate) -> str:
    f = claims.row_facts(cand.first)
    if cand.sub_issue == claims.DISPUTE:
        text = (f"Return {f['return_order']} (AWB {f['return_awb']}, {f['return_carrier']}) is marked delivered "
                f"on {f['marked_delivered_on'] or 'the portal'}, but the shipment has not been received at our "
                f"warehouse - there is no inward / GRN entry for it.")
        if f["ic_reversed"]:
            text += " The invoice payment was later reversed."
        return text + " Kindly share the POD or reimburse the invoice amount under SPF."
    return (f"Return {f['return_order']} (AWB {f['return_awb']}, {f['return_carrier']}) was created on "
            f"{f['return_created']} and has still not been received at our warehouse. Kindly check and "
            f"refund the invoice amount under SPF.")


def raise_tickets(cands: list[claims.Candidate], ctx: dict, dry: bool) -> list[dict]:
    from .sources.ajio import spf_payload
    client = ctx.get("ajio")
    done, cap = [], num("MAX_RAISE_PER_RUN", 40)
    for cand in cands:
        f = claims.row_facts(cand.first)
        rec = {"return_order": cand.return_order, "pob": cand.pob_id, "type": cand.sub_issue,
               "reason": cand.reason, "awb": f["return_awb"], "invoice": f["invoice"],
               "amount": f["invoice_amount"], "sku": ", ".join(claims.txt(r.get("SELLER SKU")) for r in cand.rows),
               "ticket": "", "result": ""}
        if len([d for d in done if d["result"] in ("created", "dry-run")]) >= cap:
            rec["result"] = "over daily cap"
            done.append(rec)
            continue
        if dry or client is None:
            rec["result"] = "dry-run" if dry else "skipped (AJIO not logged in)"
            done.append(rec)
            continue
        try:
            existing = client.tickets_for_return(cand.return_order)
            if existing:
                rec.update(result="already has ticket", ticket=", ".join(str(t.get("ticketId")) for t in existing))
                done.append(rec)
                continue
            ro = client.return_order(cand.return_order, cand.pob_id)
            if not ro or not ro.get("returnOrderItems"):
                rec["result"] = "return order not found on portal"
                done.append(rec)
                continue
            sub = claims.DISPUTE if str(ro.get("returnStatus")).upper() == "RETURN_DELIVERED" else cand.sub_issue
            ids = []
            for item in ro["returnOrderItems"]:
                payload = spf_payload(ro, item, sub, describe(cand), cand.pob_id, env("AJIO_USERNAME"))
                ids.append(client.create_spf(payload))
                time.sleep(1.2)
            rec.update(result="created", ticket=", ".join(ids), type=sub)
        except Exception as exc:  # noqa: BLE001
            rec["result"] = f"error: {exc}"[:300]
        done.append(rec)
    return done


# ------------------------------------------------------------------ follow up

def follow_up(rows: list[dict], ctx: dict, dry: bool, now: datetime) -> list[dict]:
    client, tickets = ctx.get("ajio"), ctx.get("tickets") or []
    if client is None:
        return []
    by_rt, by_line = {}, {}
    for r in rows:
        by_rt.setdefault(claims.txt(r.get("RETURN ORDER NUMBER")), []).append(r)
        by_line[(claims.txt(r.get("FWD B2B INVOICE No")), claims.txt(r.get("SELLER SKU")))] = r
    gemini = followup.Gemini(env("GEMINI_API_KEY")) if env("GEMINI_API_KEY") else None
    cap, cooldown = num("MAX_REPLIES_PER_RUN", 30), num("REPLY_COOLDOWN_HOURS", 24)
    out = []
    for t in tickets:
        if len(out) >= cap:
            break
        od = t.get("orderDetails") or {}
        mine = by_rt.get(str(od.get("returnOrderId") or "")) or [
            by_line.get((str(od.get("forwardInvoiceNumber") or ""), str(od.get("skuId") or "")))]
        mine = [r for r in mine if r]
        # Received, or never picked up from the customer: nothing to chase.
        if not mine or any(claims.is_received(r) or claims.not_picked(r) for r in mine):
            continue
        if not followup.needs_us(t, client.seller_id, now, cooldown):
            continue
        facts = claims.row_facts(mine[0])
        plan = followup.plan_reply(t, facts, client.seller_id, gemini)
        if gemini is not None:
            time.sleep(num("GEMINI_PAUSE_SECONDS", 8))   # free tier: stay under the per-minute limit
        if plan is None:
            continue
        rec = {"ticket": plan.ticket_id, "return_order": plan.return_order, "status": plan.status,
               "reopen": "yes" if plan.reopen else "", "why": plan.reason, "by": plan.source,
               "reply": plan.reply, "result": "dry-run"}
        if not dry:
            try:
                client.comment(plan.ticket_id, plan.reply)
                rec["result"] = "sent"
                time.sleep(1.0)
            except Exception as exc:  # noqa: BLE001
                rec["result"] = f"error: {exc}"[:300]
        out.append(rec)
    return out


# ------------------------------------------------------------------ master log

LOG_SHEET = "SPF BOT LOG"
LOG_COLUMNS = ["RUN DATE", "MODE", "ACTION", "RETURN ORDER", "TICKET ID", "TYPE / STATUS", "REASON",
               "AMOUNT", "RESULT", "REPLY TEXT"]


def read_bot_log(path: Path) -> list[list]:
    """The bot's own history sheet of the master; the master script does not keep it."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True)
        if LOG_SHEET not in wb.sheetnames:
            return []
        rows = [list(r) for r in list(wb[LOG_SHEET].iter_rows(values_only=True))[1:] if any(r)]
        wb.close()
        return rows
    except Exception as exc:  # noqa: BLE001
        log.warning("Could not read %s: %s", LOG_SHEET, exc)
        return []


def write_bot_log(path: Path, previous: list[list], raised: list[dict], replies: list[dict],
                  dry: bool, today: date) -> None:
    """Append this run to the SPF BOT LOG sheet and, for tickets really created,
    fill TICKET ID / STATUS on their Sheet1 rows at once (the ticket export
    would only bring them in on the next run)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    mode = "DRY RUN" if dry else "LIVE"
    new = [[today, mode, "RAISE", r["return_order"], r["ticket"], r["type"], r["reason"], r["amount"],
            r["result"], ""] for r in raised]
    new += [[today, mode, "REPLY" + (" + RE-OPEN" if r["reopen"] else ""), r["return_order"], r["ticket"],
             r["status"], r["why"], None, r["result"], r["reply"]] for r in replies]
    wb = load_workbook(path)
    if LOG_SHEET in wb.sheetnames:
        del wb[LOG_SHEET]
    ws = wb.create_sheet(LOG_SHEET)
    ws.append(LOG_COLUMNS)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="FFFFFF00")
    for row in new[::-1] + previous:       # newest run on top
        ws.append(row)
    for col, width in zip("ABCDEFGHIJ", (12, 9, 18, 14, 16, 26, 40, 10, 22, 80)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    for cell in ws["A"][1:]:
        cell.number_format = "dd-mm-yyyy"

    created = {r["return_order"]: r["ticket"] for r in raised if r["result"] == "created"}
    if created:
        sheet = wb["Sheet1"]
        head = {str(c.value).strip(): c.column for c in sheet[1] if c.value}
        rt, tid, st = head.get("RETURN ORDER NUMBER"), head.get("TICKET ID"), head.get("STATUS")
        if rt and tid:
            for row in sheet.iter_rows(min_row=2):
                ticket = created.get(claims.txt(row[rt - 1].value))
                if ticket:
                    row[tid - 1].value = ticket
                    if st:
                        row[st - 1].value = "CREATED (bot)"
    wb.save(path)


# ------------------------------------------------------------------ report

def summary(rows, raised, replies, soon, ctx, dry, today) -> str:
    esc = telegram.esc
    pending = sum(1 for r in rows if not claims.is_received(r))
    ok_raise = [r for r in raised if r["result"] in ("created", "dry-run")]
    by_type = {}
    for r in ok_raise:
        by_type[r["type"]] = by_type.get(r["type"], 0) + 1
    sent = [r for r in replies if r["result"] in ("sent", "dry-run")]
    verb = "to raise" if dry else "raised"
    lines = [f"<b>AJIO SPF - {today:%d-%m-%Y}</b>" + (" (DRY RUN - nothing posted)" if dry else ""),
             f"Returns {len(rows)} | received {len(rows) - pending} | not received {pending}",
             f"Tickets {verb}: {len(ok_raise)}"
             + (" (" + ", ".join(f"{k.split()[0].title()} {v}" for k, v in sorted(by_type.items())) + ")" if by_type else ""),
             f"Replies {'drafted' if dry else 'sent'}: {len(sent)} (re-open {sum(1 for r in sent if r['reopen'])})",
             f"Next days: dispute {soon['dispute']} | 60-day transit {soon['transit']}"]
    pk = ctx.get("pickup")
    if pk:
        lines.append(f"Courier pickup (not received): picked {pk.get('PICKED UP', 0)} | not picked "
                     f"{pk.get('NOT PICKED UP', 0)} | cancelled {pk.get('PICKUP CANCELLED', 0)} | failed "
                     f"{pk.get('PICKUP FAILED', 0)} | no tracking {pk.get('NO TRACKING', 0)}")
    elif ctx.get("pickup_error"):
        lines.append(f"Courier pickup check failed: {esc(ctx['pickup_error'])[:150]}")
    if soon["missed_dispute"]:
        lines.append(f"Delivered over {num('DISPUTE_MAX_DAYS', 365)} days ago, no ticket: {soon['missed_dispute']}")
    other = [r for r in raised if r["result"] not in ("created", "dry-run")]
    if other:
        lines.append(f"Raise skipped/failed: {len(other)}")
    for r in ok_raise[:10]:
        lines.append(f"- {esc(r['return_order'])} {esc(r['type'].split()[0].title())} Rs{r['amount']:.0f} ({esc(r['reason'])})")
    if len(ok_raise) > 10:
        lines.append(f"... +{len(ok_raise) - 10} more in the file")
    if ctx["errors"]:
        lines.append("<b>Errors</b> (raising and replies paused):")
        lines += [f"- {esc(e)[:200]}" for e in ctx["errors"]]
    drive = ctx.get("master_drive")
    if drive is not None and getattr(drive, "meta", None):
        lines.append(f'<a href="{drive.url}">AJIO MASTER</a>')
    return "\n".join(lines)


def actions_file(raised, replies, today) -> Path:
    path = OUTPUT / f"SPF_ACTIONS_{today:%Y-%m-%d}.xlsx"
    raise_cols = ["return_order", "pob", "type", "reason", "awb", "invoice", "amount", "sku", "ticket", "result"]
    reply_cols = ["ticket", "return_order", "status", "reopen", "why", "by", "reply", "result"]
    return write_xlsx(path, {"Raise": (raise_cols, raised), "Replies": (reply_cols, replies)})


# ------------------------------------------------------------------ main

def run(args) -> int:
    dry = dry_run(args.dry_run)
    now = datetime.now()
    today = now.date()
    if WORK.exists():
        shutil.rmtree(WORK)
    (WORK / "input" / "Manual Sheet").mkdir(parents=True)
    ctx = {"errors": []}
    raised: list[dict] = []
    replies: list[dict] = []
    try:
        fetch_drive(ctx)
        bot_log = read_bot_log(WORK / MASTER_NAME)
        fetch_all(today, ctx)
        build_master()
        try:
            # Second pass when OMSGuru changed anything: receipts, and forward AWBs
            # the master script matches the ERP dispatch register on.
            if oms_lookup(today, ctx):
                build_master()
        except Exception as exc:  # noqa: BLE001
            log.error("OMSGuru failed: %s", exc)
            ctx["errors"].append(f"OMSGuru: {exc}")
        try:
            from . import pickup
            ctx["pickup"] = pickup.check(WORK / MASTER_NAME)
        except Exception as exc:  # noqa: BLE001 - informational columns only
            log.error("Courier pickup check failed: %s", exc)
            ctx["pickup_error"] = str(exc)
        rows = claims.read_sheet(WORK / MASTER_NAME)
        rules = claims.Rules(num("DISPUTE_AFTER_DAYS", 5), num("DISPUTE_MAX_DAYS", 365),
                             num("TRANSIT_AFTER_DAYS", 60), num("TRANSIT_MAX_DAYS", 90))
        # A missing lookup source could make a received parcel look lost, so
        # nothing is posted to AJIO on a run where any source failed.
        act_dry = dry or bool(ctx["errors"])
        raised = raise_tickets(claims.candidates(rows, today, rules), ctx, act_dry)
        replies = follow_up(rows, ctx, act_dry, now)
        soon = claims.upcoming(rows, today, rules)
        write_bot_log(WORK / MASTER_NAME, bot_log, raised, replies, act_dry, today)
        upload_failed = False
        try:
            if ctx.get("master_drive") is not None and not args.no_upload:
                ctx["master_drive"].upload(WORK / MASTER_NAME)
            elif env("SPF_MASTER_XLSX_PATH"):
                shutil.copy2(WORK / MASTER_NAME, OUTPUT / MASTER_NAME)
        except Exception as exc:  # noqa: BLE001 - the report of what was posted must still go out
            log.error("Master upload failed: %s", exc)
            ctx["errors"].append(f"Master not saved to Drive: {exc}")
            upload_failed = True
        path = actions_file(raised, replies, today)
        text = summary(rows, raised, replies, soon, ctx, act_dry, today)
        if args.no_telegram:
            print(text)
        else:
            telegram.send(text)
            if raised or replies:
                telegram.send_document(path, f"AJIO SPF actions {today:%d-%m-%Y}")
        return 1 if upload_failed else 0
    except Exception as exc:
        # Tickets / replies already posted are never lost: their list still goes out.
        if any(r.get("result") in ("created", "sent") for r in raised + replies):
            try:
                path = actions_file(raised, replies, today)
                if not args.no_telegram:
                    telegram.send_document(path, f"AJIO SPF actions {today:%d-%m-%Y} (run failed afterwards)")
            except Exception as inner:  # noqa: BLE001
                log.error("Could not send the actions file: %s", inner)
        exc.notes = list(ctx["errors"])
        raise


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="run.py ajio-spf", description="Daily AJIO SPF claim run")
    p.add_argument("--dry-run", choices=("0", "1"), default=None, help="override SPF_DRY_RUN")
    p.add_argument("--no-upload", action="store_true", help="build locally, do not write the master to Drive")
    p.add_argument("--no-telegram", action="store_true", help="do not send the Telegram message")
    return run_bot("AJIO SPF run", p, run, argv)
