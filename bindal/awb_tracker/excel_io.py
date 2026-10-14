"""Reading the AWB list and writing the tracking report."""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TrackResult
from .runner import Job

AWB_HEADERS = {"awb", "awbno", "awbnumber", "waybill", "trackingid",
               "trackingno", "tracking_id", "shipmentid", "lrn"}
PARTNER_HEADERS = {"partner", "partnername", "courier", "couriername",
                   "carrier", "logistics", "partner_name"}

STATUS_FILLS = {
    "DELIVERED": "C6EFCE",
    "OUT_FOR_DELIVERY": "BDD7EE",
    "OUT_FOR_PICKUP": "CCE0F0",
    "IN_TRANSIT": "DDEBF7",
    "PICKED_UP": "E2EFDA",
    "PENDING": "FFF2CC",
    "UNDELIVERED": "FCE4D6",
    "PICKUP_FAILED": "FFE0CC",
    "RTO": "FFD9D9",
    "CANCELLED": "F2F2F2",
    "LOST": "FFC7CE",
    "NOT_FOUND": "FFC7CE",
    "UNKNOWN": "FFFFFF",
    # Return leg: same hues, shifted warmer so the direction reads at a glance.
    "RETURN_DELIVERED": "D9EAD3",
    "RETURN_OUT_FOR_DELIVERY": "D9E2F3",
    "RETURN_IN_TRANSIT": "EAF1FA",
    "RETURN_PICKED_UP": "EDF3E7",
    "RETURN_PENDING": "FFF7E0",
}

FLOW_FILLS = {
    "RETURN": "FFF2CC",
    "FORWARD": "FFFFFF",
}

PARTY_FILLS = {
    "SELLER": "FFE699",
    "CUSTOMER": "B4C7E7",
    "COURIER": "F8CBAD",
    "QC_FAIL": "E4DFEC",
    "UNKNOWN": "F2F2F2",
}


def _norm(value: Any) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum() or ch == "_")


def _pick_columns(header: Sequence[Any]) -> tuple[int, int]:
    """Locate the AWB and partner columns, falling back to positions 0 and 1."""
    awb_col, partner_col = -1, -1
    for i, cell in enumerate(header):
        key = _norm(cell)
        if awb_col < 0 and key in AWB_HEADERS:
            awb_col = i
        elif partner_col < 0 and key in PARTNER_HEADERS:
            partner_col = i
    if awb_col < 0:
        awb_col = 0
    if partner_col < 0:
        partner_col = 1 if len(header) > 1 else -1
    return awb_col, partner_col


def read_jobs(path: str | Path, sheet: str | None = None) -> list[Job]:
    """Read AWB/partner pairs from .xlsx, .xlsm, .csv or .txt."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"input file not found: {path}")

    if path.suffix.lower() in (".csv", ".txt"):
        with path.open(newline="", encoding="utf-8-sig") as fh:
            rows = [r for r in csv.reader(fh) if any(str(c).strip() for c in r)]
    else:
        book = load_workbook(path, read_only=True, data_only=True)
        ws = book[sheet] if sheet else book.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)
                if any(c is not None and str(c).strip() for c in r)]
        book.close()

    if not rows:
        return []

    awb_col, partner_col = _pick_columns(rows[0])
    header_looks_like_data = _norm(rows[0][awb_col]) not in AWB_HEADERS
    body = rows if header_looks_like_data else rows[1:]

    jobs: list[Job] = []
    seen: set[tuple[str, str]] = set()
    for row in body:
        if awb_col >= len(row):
            continue
        awb = row[awb_col]
        if awb is None or not str(awb).strip():
            continue
        # Excel stores long numeric AWBs as floats; keep them as plain digits.
        if isinstance(awb, float) and awb.is_integer():
            awb = int(awb)
        awb = str(awb).strip()
        partner = ""
        if 0 <= partner_col < len(row) and row[partner_col] is not None:
            partner = str(row[partner_col]).strip()
        key = (awb, partner.upper())
        if key in seen:
            continue
        seen.add(key)
        jobs.append(Job(awb, partner))
    return jobs


def _autosize(ws, max_width: int = 55) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(10, longest + 2), max_width)


def _write_sheet(ws, rows: list[dict[str, Any]],
                 colour_columns: dict[str, dict[str, str]] | None = None) -> None:
    if not rows:
        ws.append(["(no rows)"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column, palette in (colour_columns or {}).items():
        if column not in headers:
            continue
        idx = headers.index(column) + 1
        for row_i in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_i, column=idx)
            colour = palette.get(str(cell.value))
            if colour:
                cell.fill = PatternFill("solid", fgColor=colour)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


def write_report(results: Sequence[TrackResult], path: str | Path,
                 stats: dict[str, Any] | None = None) -> Path:
    """Write a three-sheet workbook: Summary, Events, Run Stats."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    book = Workbook()
    _write_sheet(book.active, [r.summary_row() for r in results],
                 colour_columns={"STATUS": STATUS_FILLS, "FLOW": FLOW_FILLS,
                                 "CANCELLED_BY": PARTY_FILLS})
    book.active.title = "Summary"

    events: list[dict[str, Any]] = []
    for res in results:
        events.extend(res.event_rows())
    _write_sheet(book.create_sheet("Events"), events)

    if stats:
        rows = [{"METRIC": "status", "KEY": k, "COUNT": v}
                for k, v in stats.get("by_status", {}).items()]
        rows += [{"METRIC": "partner", "KEY": f"{p} ({k})", "COUNT": v}
                 for p, counts in stats.get("by_partner", {}).items()
                 for k, v in counts.items()]
        rows += [{"METRIC": "cancelled_by", "KEY": k, "COUNT": v}
                 for k, v in stats.get("by_party", {}).items()]
        _write_sheet(book.create_sheet("Run Stats"), rows)

    book.save(path)
    return path


def write_json(results: Iterable[TrackResult], path: str | Path,
               include_raw: bool = False) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump([r.to_dict(include_raw) for r in results], fh,
                  indent=2, ensure_ascii=False, default=str)
    return path
