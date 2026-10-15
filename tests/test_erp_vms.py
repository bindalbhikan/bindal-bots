from datetime import date, datetime

from openpyxl import Workbook

from bindal.vms.erp_vms import OTHER, compare, erp_parcels
from bindal.vms.workbook import VmsWorkbook


class FakeErp:
    def __init__(self, lines):
        self.lines = lines

    def packed(self, portal, since):
        return self.lines


def line(awb, day, sku, who="PRAKASH"):
    return {"Order ID": awb, "Voucher Date": datetime(2026, 10, day), "SKU Code": sku,
            "Created By": who, "Voucher No.": f"V{day}"}


def test_a_parcel_counts_once_on_its_first_pack_day():
    parcels = erp_parcels(FakeErp([line("MYSP1", 3, "A"), line("mysp1", 2, "B"), line("MYSP2", 2, "C"),
                                   {"Order ID": "", "Voucher Date": datetime(2026, 10, 2)}]), date(2026, 10, 1))
    assert set(parcels) == {"MYSP1", "MYSP2"}
    assert parcels["MYSP1"]["date"] == date(2026, 10, 2)
    assert parcels["MYSP1"]["skus"] == ["A", "B"]


def test_compare_counts_vms_and_lists_the_rest():
    parcels = {"MYSP1": {"date": date(2026, 10, 2), "packed_by": "P", "voucher": "V", "skus": ["A"]},
               "MYSP2": {"date": date(2026, 10, 2), "packed_by": "K", "voucher": "V", "skus": ["B"]},
               "MYSP3": {"date": date(2026, 10, 3), "packed_by": "K", "voucher": "V", "skus": ["C"]}}
    per_day, missing = compare(parcels, {"MYSP1": "TLPL", "MYSP2": "MMRJ"}, {"MYSP1", "MYSP3"}, ["TLPL", "MMRJ"])
    assert per_day[date(2026, 10, 2)]["TLPL"] == [1, 1]
    assert per_day[date(2026, 10, 2)]["MMRJ"] == [1, 0]
    assert per_day[date(2026, 10, 3)][OTHER] == [1, 1]
    assert [m["awb"] for m in missing] == ["MYSP2"]


def month_book(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "OCT-26"
    ws["A1"] = date(2026, 10, 1)
    ws.cell(3, 1, "DATE")
    for d in range(1, 32):
        ws.cell(3 + d, 1, date(2026, 10, d))
    ws.cell(35, 1, "TOTAL")
    wb.create_sheet("NO VMS")
    path = tmp_path / "vms.xlsx"
    wb.save(path)
    return VmsWorkbook(path)


def test_erp_sheets(tmp_path):
    book = month_book(tmp_path)
    per_day = {date(2026, 10, 2): {"TLPL": [4, 3], "MMRJ": [2, 2]}}
    book.write_erp_month(2026, 10, per_day, ["TLPL", "MMRJ"], upto=date(2026, 10, 3))
    ws = book.wb["ERP VS VMS OCT-26"]
    assert book.wb.sheetnames.index("ERP VS VMS OCT-26") == book.wb.sheetnames.index("NO VMS") + 1
    assert [c.value for c in ws[2] if c.value] == ["TLPL", "MMRJ", "TOTAL"]
    assert [ws.cell(5, c).value for c in range(1, 11)][:4] == [date(2026, 10, 2), 4, 3, 0.75]
    assert ws.cell(5, 8).value == 6 and ws.cell(5, 9).value == 5          # TOTAL group of the day
    assert ws.cell(7, 1).value == "TOTAL" and ws.cell(7, 2).value == 4

    book.write_erp_no_vms(date(2026, 10, 1), date(2026, 10, 3),
                          [{"date": date(2026, 10, 2), "code": "TLPL", "awb": "MYSP9", "packed_by": "P", "voucher": "V", "skus": "A"}])
    book.write_erp_no_vms(date(2026, 10, 1), date(2026, 10, 3), [])     # next run: list rebuilt, not appended
    assert book.wb["ERP NO VMS"].max_row == 1


def test_erp_text_shows_yesterday_and_today():
    from bindal.vms.pipeline import erp_text
    zero = {"TLPL": [0, 0], OTHER: [0, 0]}
    e = {"from": date(2026, 10, 1), "to": date(2026, 10, 15), "codes": ["TLPL", OTHER],
         "total": {"TLPL": [10, 9], OTHER: [0, 0]}, "missing": [], "unknown_channel": 0,
         "recent": [{"date": date(2026, 10, 14), "codes": {"TLPL": [4, 3], OTHER: [0, 0]}},
                    {"date": date(2026, 10, 15), "codes": zero}]}
    text = erp_text(e)
    assert "Yesterday · 14 Oct" in text and "Today · 15 Oct" in text
    assert "TLPL      4     3  75.0%" in text
