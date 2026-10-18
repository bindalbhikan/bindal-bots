from datetime import date

from openpyxl import Workbook

from bindal.sale import sjit
from bindal.sale.workbook import SaleWorkbook

BLOCKS = ["AZ/FX/FBA", "SJIT"]


def month_sheet(wb, title, year, month, dated_days, rows=31):
    ws = wb.create_sheet(title)
    ws["A1"] = date(year, month, 1)
    col = 2
    for block in BLOCKS:
        ws.cell(2, col, block)
        for i, metric in enumerate(("GROSS", "RETURN", "CANCEL", "NET")):
            ws.cell(3, col + i, metric)
        col += 4
    for i in range(rows):
        ws.cell(4 + i, 1, date(year, month, i + 1) if i < dated_days else None)
    ws.cell(4 + rows, 1, "TOTAL")
    return ws


def test_march_sheet_from_february_has_31_days(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    month_sheet(wb, "JAN-26", 2026, 1, 31)
    month_sheet(wb, "FEB-26", 2026, 2, 28)   # 31 rows, last 3 without a date
    path = tmp_path / "sale.xlsx"
    wb.save(path)
    book = SaleWorkbook(path)
    assert book.write(date(2026, 3, 31), "SJIT", "GROSS", 5)
    assert book.read(date(2026, 3, 31), "SJIT", "GROSS") == 5


def test_sjit_parse_any_header_case():
    csv_text = ("Order Line Id,Created On,PO_Type,Order Status,Cancelled On\n"
                "1,2026-10-01 10:00:00,SJIT,C,\n"
                "2,2026-10-01 11:00:00,PPMP,C,\n"
                "3,2026-10-01 12:00:00,FBM,F,2026-10-02 09:00:00\n")
    gross, cancel, covered = sjit.parse([("r.csv", csv_text.encode())])
    assert gross[date(2026, 10, 1)] == 2
    assert cancel[date(2026, 10, 2)] == 1
    assert date(2026, 10, 1) in covered


def test_cocoblu_has_its_own_block_and_snapdeal_is_ignored():
    from bindal.sale.oms_sales import IGNORE, OMS_BLOCKS, block_for
    assert block_for("BINDAL EXPORTS LIMITED - COCOBLU") == "COCOBLU"
    assert block_for("BINDAL EXPORTS LIMITED - Amazon") == "AZ/FX/FBA"
    assert block_for("BINDAL EXPORTS LIMITED - Snapdeal") == IGNORE
    assert "COCOBLU" in OMS_BLOCKS and IGNORE not in OMS_BLOCKS


def test_move_blocks_adds_cocoblu_after_amazon_and_drops_snapdeal(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("OCT-26")
    ws["A1"] = date(2026, 10, 1)
    blocks = ["AZ/FX/FBA", "MEESHO", "SNAPDEAL", "TOTAL"]
    for i, block in enumerate(blocks):
        ws.cell(2, 2 + 4 * i, block)
        for k, m in enumerate(("GROSS", "RETURN", "Cancel", "NET")):
            ws.cell(3, 2 + 4 * i + k, m)
    ws.cell(3, 18, "Target")
    for d in range(31):
        ws.cell(4 + d, 1, date(2026, 10, d + 1))
        ws.cell(4 + d, 2, 5)          # AZ gross
        ws.cell(4 + d, 6, 7)          # MEESHO gross
        ws.cell(4 + d, 7, 1)          # MEESHO return (by hand)
    ws.cell(35, 1, "TOTAL")
    path = tmp_path / "sale.xlsx"
    wb.save(path)
    book = SaleWorkbook(path)
    book.move_blocks(book.wb["OCT-26"], "COCOBLU", "AZ/FX/FBA", "SNAPDEAL")
    layout, _, _ = book.layout(book.wb["OCT-26"])
    assert list(layout) == ["AZ/FX/FBA", "COCOBLU", "MEESHO", "TOTAL"]
    assert layout["COCOBLU"]["GROSS"] == 6 and layout["MEESHO"]["GROSS"] == 10
    assert book.read(date(2026, 10, 3), "MEESHO", "GROSS") == 7
    assert book.read(date(2026, 10, 3), "MEESHO", "RETURN") == 1
    assert book.wb["OCT-26"].cell(3, 18).value == "Target"
    assert book.has_block(date(2026, 10, 3), "COCOBLU")


# ---- Myntra partner portal (saved session)

from datetime import datetime  # noqa: E402

from bindal.portals import myntra as mp  # noqa: E402


def test_session_expiry_ignores_bot_manager_cookies():
    s = mp.parse_session('{"savedAt": "2026-10-14T20:00:00", "cookies": ['
                         '{"name": "ak_bmsc", "value": "x", "domain": ".myntrainfo.com", "expires": 1, "httpOnly": true},'
                         '{"name": "sid", "value": "y", "domain": ".myntrainfo.com", "expires": 1800000000, "httpOnly": true},'
                         '{"name": "other", "value": "z", "domain": ".example.com", "expires": 2, "httpOnly": true}]}')
    assert [c["name"] for c in s.cookies] == ["ak_bmsc", "sid"]
    assert s.expires() == datetime.fromtimestamp(1800000000)
    assert mp.expiry_warning(s, 3, now=datetime.fromtimestamp(1800000000 - 86400)).endswith("in 1 day(s))")
    assert mp.expiry_warning(s, 3, now=datetime.fromtimestamp(1800000000 - 10 * 86400)) == ""


def test_session_from_cookie_header_and_curl():
    assert [c["name"] for c in mp.parse_session("a=1; b=2=3").cookies] == ["a", "b"]
    curl = "curl 'https://partnersapi.myntrainfo.com/x' -H 'accept: */*' -b 'sid=abc; at=def' -H 'x: y'"
    cookies = mp.parse_session(curl).cookies
    assert [(c["name"], c["value"]) for c in cookies] == [("sid", "abc"), ("at", "def")]
    assert mp.parse_session('{"cookie": "sid=1"}').expires() is None


def _job(id, start, end, created, status="COMPLETED", report=4, store="2297"):
    return {"id": id, "reportId": report, "status": status, "createdOn": int(created.timestamp() * 1000),
            "queryParams": {"storeId": store, "partnerType": "PPMP",
                            "reportStartDate": start, "reportEndDate": end}}


def test_plan_reuses_existing_reports_and_names_missing_days():
    d = lambda n: date(2026, 10, n)  # noqa: E731
    jobs = [_job(1, "2026-10-01", "2026-10-11", datetime(2026, 10, 12, 17)),
            _job(2, "2026-10-12", "2026-10-13", datetime(2026, 10, 14, 22)),
            _job(3, "2026-10-13", "2026-10-13", datetime(2026, 10, 13, 9)),             # end day not over
            _job(4, "2026-10-11", "2026-10-12", datetime(2026, 10, 14, 10), "FAILED"),
            _job(5, "2026-10-12", "2026-10-13", datetime(2026, 10, 14, 23), report=7),  # another report
            _job(6, "2026-10-14", "2026-10-14", datetime(2026, 10, 15, 9), "CREATED")]
    for j in jobs:
        j["successFile"] = "x" if j["status"] == "COMPLETED" else None
    done, running, missing = mp.plan(jobs, [d(11), d(12), d(13), d(14), d(15)])
    assert [j["id"] for j in done] == [2, 1]
    assert [j["id"] for j in running] == [6]
    assert missing == [d(15)]
    assert mp.plan(jobs, [d(12), d(13)])[0][0]["id"] == 2


def test_session_without_expiry_is_taken_as_15_days():
    from datetime import timedelta
    s = mp.parse_session('{"savedAt": "2026-10-14T23:30:00", "cookies": '
                         '[{"name": "session", "value": "1", "domain": ".myntrainfo.com"}]}')
    assert "about 15 days" in mp.expiry_warning(s, 3, now=s.saved_at + timedelta(days=13))
    assert mp.expiry_warning(s, 3, now=s.saved_at + timedelta(days=5)) == ""


def test_session_saved_with_utc_time():
    s = mp.parse_session('{"savedAt": "2026-10-14T18:00:00.000Z", "expiresAt": null, "cookies": ['
                         '{"name": "sid", "value": "1", "domain": "partners.myntrainfo.com", "expires": 1800000000,'
                         ' "httpOnly": true}]}')
    assert s.saved_at is not None and s.saved_at.tzinfo is None
    assert s.expires() == datetime.fromtimestamp(1800000000)


def test_session_from_devtools_request_headers():
    dump = ("set-cookie\nerp.at=old; Max-Age=2592000; Path=/; HttpOnly\n"
            ":authority\npartnersapi.myntrainfo.com\n"
            "cookie\n_ga=1; erp.at=A; erp.rt=R; session=S\n"
            "origin\nhttps://partners.myntrainfo.com\n")
    cookies = {c["name"]: c["value"] for c in mp.parse_session(dump).cookies}
    assert cookies == {"_ga": "1", "erp.at": "A", "erp.rt": "R", "session": "S"}


def test_interrupted_job_is_not_reused():
    jobs = [_job(9, "2026-03-01", "2026-03-31", datetime(2026, 10, 1, 9), "INTERRUPTED")]
    assert mp.exact_job(jobs, 4, date(2026, 3, 1), date(2026, 3, 31)) is None
