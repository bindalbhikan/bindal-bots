import html
import re

import pytest

from bindal.core import telegram
from bindal.core.config import flag
from bindal.core.drive import DriveError, DriveFile


def balanced(part: str) -> bool:
    stack = []
    for m in telegram.TAG.finditer(part):
        name = m.group(2).lower()
        if m.group(1):
            if not stack or stack[-1] != name:
                return False
            stack.pop()
        else:
            stack.append(name)
    return not stack


def test_short_message_is_one_part():
    assert telegram.split_html("<b>hi</b>") == ["<b>hi</b>"]


def test_long_pre_block_is_split_with_balanced_tags():
    rows = "\n".join(f"AWB{i:012d} Shadowfax 14-Oct Over &amp; more" for i in range(400))
    text = "<b>Pickup Pending</b>\n<pre>" + rows + "</pre>\n<i>legend</i>"
    parts = telegram.split_html(text, 1000)
    assert len(parts) > 5
    for part in parts:
        assert len(part) <= 1000
        assert balanced(part), part[:80]
        assert not re.search(r"&[a-z]*$", part.split("</")[0])  # no entity cut in half
    # nothing lost: the rows all come back once the added tags are removed
    joined = "".join(parts).replace("</pre><pre>", "")
    assert joined.count("AWB") == 400


def test_one_huge_line_is_cut():
    parts = telegram.split_html("<pre>" + "x" * 9000 + "</pre>", 4000)
    assert all(len(p) <= 4000 and balanced(p) for p in parts)
    assert sum(p.count("x") for p in parts) == 9000


@pytest.mark.parametrize("value,default,expected", [
    ("1", False, True), ("true", False, True), ("YES", False, True),
    ("0", True, False), ("false", True, False), ("off", True, False),
    ("", True, True), ("maybe", True, True), ("", False, False),
])
def test_flag(monkeypatch, value, default, expected):
    monkeypatch.setenv("X_FLAG", value)
    assert flag("X_FLAG", default) is expected


def test_upload_needs_a_download_first():
    drive = DriveFile("file-id", http=object())
    with pytest.raises(DriveError):
        drive.upload(__file__)


def test_blank_setting_falls_back_to_default(monkeypatch):
    from bindal.core.config import env
    monkeypatch.setenv("X_SETTING", "")          # GitHub: unset repository variable
    assert env("X_SETTING", "45111:TLPL") == "45111:TLPL"
    monkeypatch.setenv("X_SETTING", "  1 ")
    assert env("X_SETTING", "9") == "1"


def test_md_turns_model_markdown_into_phone_html():
    text = telegram.md("## Plan\n| Channel | Units |\n|---|---:|\n| Ajio | 12 |\n- **Push** `SKU-1` <now>")
    assert text == "<b>Plan</b>\nChannel · Units\nAjio · 12\n• <b>Push</b> <code>SKU-1</code> &lt;now&gt;"
    assert telegram.md("Sale 10 - 12 units") == "Sale 10 - 12 units"


def test_report_inline_when_it_fits(monkeypatch):
    sent, docs = [], []
    monkeypatch.setattr(telegram, "send", lambda t: sent.append(t) or True)
    monkeypatch.setattr(telegram, "send_document", lambda p, c="": docs.append(p) or True)
    report = ["<b>Head</b>", telegram.Listing("<b>AWBs: 2</b>", ["• A1", "• A2"], [{"AWB": "A1"}, {"AWB": "A2"}], "AWB")]
    assert telegram.send_report(report, "Test")
    assert sent == ["<b>Head</b>\n<b>AWBs: 2</b>\n• A1\n• A2"] and docs == []


def test_report_too_long_goes_to_excel(monkeypatch):
    from openpyxl import load_workbook
    sent, docs = [], []
    monkeypatch.setattr(telegram, "send", lambda t: sent.append(t) or True)
    monkeypatch.setattr(telegram, "send_document", lambda p, c="": docs.append((p, c)) or True)
    awbs = [f"AWB{i:012d}" for i in range(400)]
    report = ["<b>Pickup pending</b>",
              telegram.Listing("", [f"• {a} · Shadowfax · Over" for a in awbs],
                               [{"AWB": a, "Courier": "Shadowfax"} for a in awbs], "Pickup pending"),
              telegram.Listing("<b>Other: 1</b>", ["• X"], [{"SKU": "X"}], "Pickup pending"),
              "Regards"]
    assert not telegram.fits(report) and telegram.send_report(report, "AJIO_test")
    assert len(sent) == 1 and "AWB000000000001" not in sent[0]
    assert "📎 all 400 in the Excel below (sheet Pickup pending)" in sent[0] and sent[0].endswith("Regards")
    path, caption = docs[0]
    assert path.name.startswith("AJIO_test_") and "Pickup pending 400" in caption
    wb = load_workbook(path)
    assert wb.sheetnames == ["Pickup pending", "Pickup pending 2"]          # names kept unique
    ws = wb["Pickup pending"]
    assert [c.value for c in ws[1]] == ["AWB", "Courier"] and ws.max_row == 401 and ws.freeze_panes == "A2"


def test_pre_tables_fit_a_phone():
    """Every <pre> line of the reports stays within PHONE_WIDTH characters."""
    from datetime import date
    from bindal.mis import pipeline as mis
    from bindal.weekly import pipeline as weekly
    rows = [{"channel": "Myntra MAIN", "units": 12345, "units_prev": 9000, "value": 9_876_543, "value_prev": 8_000_000,
             "cancel_pct": 12, "net_sale": 9_000_000, "commission": 2_250_000, "shipping": 450_000, "net": 6_300_000}]
    texts = [mis.message(rows, [date(2026, 9, 1)], [date(2026, 8, 1)], {}),
             weekly.message([("Myntra MAIN", 4321, 3000, 1_234_567.0, 1_000_000.0, 12)], [], [], [],
                            [date(2026, 10, 18)], [date(2026, 10, 11)])]
    for text in texts:
        for block in re.findall(r"<pre>(.*?)</pre>", text, re.S):
            for line in block.splitlines():
                assert len(html.unescape(line)) <= telegram.PHONE_WIDTH, line
